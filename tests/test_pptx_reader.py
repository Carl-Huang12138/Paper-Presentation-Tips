"""Object-reader regressions, including the two supplied, immutable decks."""

from __future__ import annotations

import base64
import copy
import hashlib
import importlib.util
import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
if importlib.util.find_spec("pptx_reader"):
    from pptx_reader import extract
else:  # The first red run exercises the original implementation.
    from analyze_pptx import extract

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
}


def xp(node, expression):
    return etree.XPath(expression, namespaces=NS)(node)


def walk(items):
    for item in items:
        yield item
        yield from walk(item.get("children", []))


def patch_part(path, part, update):
    with zipfile.ZipFile(path) as package:
        contents = {name: package.read(name) for name in package.namelist()}
    root = etree.fromstring(contents[part])
    update(root)
    contents[part] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    replacement = io.BytesIO()
    with zipfile.ZipFile(replacement, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in contents.items():
            package.writestr(name, data)
    path.write_bytes(replacement.getvalue())


class ReaderBoundaries(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(prefix="reader-", dir=ROOT / "tests")
        self.addCleanup(self.cleanup_fixture)
        self.path = Path(self.folder.name) / "fixture.pptx"
        self.pres = Presentation()
        self.slide = self.pres.slides.add_slide(self.pres.slide_layouts[6])

    def cleanup_fixture(self):
        Path(self.folder.name).resolve().relative_to((ROOT / "tests").resolve())
        self.folder.cleanup()

    def save(self):
        self.pres.save(self.path)
        return self.path

    def test_native_table_has_no_unknown_graphic_data_warning(self):
        self.slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(2), Inches(1)).table.cell(0, 0).text = 'READABLE TABLE'
        item = extract(self.save())['slides'][0]['objects'][0]
        self.assertNotIn('graphic_data', item.get('missing_fields', []))
        self.assertEqual('complete', item['parse_status'])

    def text(self, value, top=1, left=1, width=4, height=1):
        box = self.slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        box.text = value
        return box

    def alternate(self, required="a14", math=False):
        box = self.text("SELECTED BODY")
        sp = box._element
        parent = sp.getparent()
        index = parent.index(sp)
        ac = etree.Element("{%s}AlternateContent" % NS["mc"], nsmap={"mc": NS["mc"], "a14": NS["a14"], "x": "urn:unsupported"})
        choice = etree.SubElement(ac, "{%s}Choice" % NS["mc"], Requires=required)
        choice.append(copy.deepcopy(sp))
        fallback = etree.SubElement(ac, "{%s}Fallback" % NS["mc"])
        fallback_shape = copy.deepcopy(sp)
        xp(fallback_shape, ".//a:t")[0].text = "FALLBACK BODY"
        fallback.append(fallback_shape)
        if math:
            para = xp(choice, ".//a:p")[0]
            wrapper = etree.SubElement(para, "{%s}m" % NS["a14"])
            omath = etree.SubElement(wrapper, "{%s}oMath" % NS["m"])
            fraction = etree.SubElement(omath, "{%s}f" % NS["m"])
            for side, value in [("num", "x"), ("den", "2")]:
                branch = etree.SubElement(fraction, "{%s}%s" % (NS["m"], side))
                run = etree.SubElement(branch, "{%s}r" % NS["m"])
                etree.SubElement(run, "{%s}t" % NS["m"]).text = value
        parent.remove(sp)
        parent.insert(index, ac)
        return self.save()

    def test_alternate_content_selects_one_branch_and_retains_provenance(self):
        result = extract(self.alternate())["slides"][0]
        self.assertIn("SELECTED BODY", result["visible_text"])
        self.assertNotIn("FALLBACK BODY", result["visible_text"])
        obj = next(x for x in walk(result["objects"]) if x.get("text_frame", {}).get("text") == "SELECTED BODY")
        source = obj.get("xml_source", {})
        self.assertEqual("/ppt/slides/slide1.xml", source.get("part"))
        self.assertEqual("2", source.get("id"))
        self.assertIn("Choice", source.get("node_path", ""))
        self.assertEqual("Choice", source.get("alternate_content", {}).get("selected_branch"))

    def test_unsupported_choice_selects_fallback_and_reports_requirement(self):
        result = extract(self.alternate(required="x"))["slides"][0]
        self.assertIn("FALLBACK BODY", result["visible_text"])
        self.assertNotIn("SELECTED BODY", result["visible_text"])
        self.assertTrue(result.get("alternate_content"))
        self.assertIn("urn:unsupported", str(result["alternate_content"]))

    def test_omml_keeps_fraction_structure_and_does_not_invent_flat_math(self):
        result = extract(self.alternate(math=True))["slides"][0]
        formulae = [m for obj in walk(result["objects"]) for m in obj.get("math", [])]
        self.assertTrue(formulae)
        self.assertEqual("unknown", formulae[0].get("semantic_status"))
        self.assertTrue(xp(etree.fromstring(formulae[0].get("raw_xml", "").encode()), ".//m:f"))
        self.assertEqual("oMath", formulae[0].get("structure", {}).get("tag"))
        self.assertNotIn("x2", result["visible_text"])
        self.assertTrue(result["parse_limitations"])

    def test_alternate_text_runs_inside_an_ordinary_shape_are_selected(self):
        box = self.text("HEAD ")
        para = xp(box._element, ".//a:p")[0]
        ac = etree.SubElement(para, "{%s}AlternateContent" % NS["mc"], nsmap={"mc": NS["mc"], "a14": NS["a14"]})
        for branch, text in [("Choice", "SELECTED BODY"), ("Fallback", "FALLBACK BODY")]:
            attrs = {"Requires": "a14"} if branch == "Choice" else {}
            choice = etree.SubElement(ac, "{%s}%s" % (NS["mc"], branch), **attrs)
            run = etree.SubElement(choice, "{%s}r" % NS["a"])
            etree.SubElement(run, "{%s}rPr" % NS["a"], sz="2400", b="1")
            etree.SubElement(run, "{%s}t" % NS["a"]).text = text
        self.save()
        original_digest = hashlib.sha256(self.path.read_bytes()).hexdigest()
        result = extract(self.path)["slides"][0]
        self.assertIn("HEAD SELECTED BODY", result["visible_text"])
        self.assertNotIn("FALLBACK BODY", " ".join(result["visible_text"]))
        frame = result["objects"][0]["text_frame"]
        self.assertEqual("Choice", frame["alternate_content_sources"][0]["selected_branch"])
        selected_run = frame["paragraphs"][0]["runs"][1]
        self.assertEqual("SELECTED BODY", selected_run["text"])
        self.assertEqual(24.0, selected_run["font"]["size_pt"])
        self.assertTrue(selected_run["font"]["bold"])
        self.assertEqual(original_digest, hashlib.sha256(self.path.read_bytes()).hexdigest())

    def test_alternate_text_inside_table_cell_keeps_cell_identity(self):
        table = self.slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1))
        table.table.cell(0, 0).text = "FIRST "
        table.table.cell(0, 1).text = "SECOND CELL"
        para = xp(table.table.cell(0, 0)._tc, ".//a:p")[0]
        ac = etree.SubElement(para, "{%s}AlternateContent" % NS["mc"], nsmap={"mc": NS["mc"], "x": "urn:unsupported"})
        for branch, text in [("Choice", "WRONG BRANCH"), ("Fallback", "SELECTED CELL")]:
            attrs = {"Requires": "x"} if branch == "Choice" else {}
            choice = etree.SubElement(ac, "{%s}%s" % (NS["mc"], branch), **attrs)
            run = etree.SubElement(choice, "{%s}r" % NS["a"])
            etree.SubElement(run, "{%s}t" % NS["a"]).text = text
        result = extract(self.save())["slides"][0]
        cell = result["objects"][0]["table"]["rows"][0][0]
        self.assertEqual("FIRST SELECTED CELL", cell["text"])
        self.assertEqual("shape-1:cell-1-1", cell["selector"])
        self.assertEqual("Fallback", cell["text_frame"]["alternate_content_sources"][0]["selected_branch"])
        self.assertIn("FIRST SELECTED CELL", result["visible_text"])
        self.assertNotIn("WRONG BRANCH", " ".join(result["visible_text"]))

    def test_geometry_source_has_a_part_for_explicit_coordinates(self):
        self.text("SOURCE BOUND GEOMETRY")
        obj = extract(self.save())["slides"][0]["objects"][0]
        self.assertEqual("/ppt/slides/slide1.xml", obj.get("geometry_source", {}).get("components", {}).get("x", {}).get("part"))

    def test_unknown_drawing_node_is_a_traceable_unparsed_region(self):
        first = self.text("FIRST KNOWN")
        unknown_node = etree.SubElement(self.slide._element.spTree, "{%s}futureObject" % NS["p"])
        unknown_node.set("marker", "UNREAD OBJECT")
        second = self.text("SECOND KNOWN")
        result = extract(self.save())["slides"][0]
        unknown = [o for o in result["objects"] if o.get("xml_source", {}).get("tag") == "futureObject"]
        self.assertTrue(unknown)
        self.assertEqual("partial", unknown[0]["parse_status"])
        self.assertIn("UNREAD OBJECT", unknown[0].get("unparsed_xml", ""))
        self.assertIsNone(unknown[0]["slide_bounds"])
        self.assertEqual("unparsed-1", unknown[0]["path"])
        known = {o["text_frame"]["text"]: o["path"] for o in result["objects"]
                 if "text_frame" in o and o["xml_source"]["kind"] == "slide"}
        self.assertEqual({"FIRST KNOWN": "shape-1", "SECOND KNOWN": "shape-2"}, known)

    def test_hidden_parent_suppresses_child_text_and_typography(self):
        group = self.slide.shapes.add_group_shape()
        child = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        child.text = "HIDDEN GROUP CHILD"
        xp(group._element, "./p:nvGrpSpPr/p:cNvPr")[0].set("hidden", "1")
        result = extract(self.save())["slides"][0]
        self.assertNotIn("HIDDEN GROUP CHILD", result["visible_text"])
        nested = result["objects"][0]["children"][0]
        self.assertFalse(nested.get("visible_in_slideshow", True))
        self.assertTrue(nested.get("visibility", {}).get("self_visible"))
        self.assertEqual([], result["typography"]["samples"])

    def group_transform(self, rotation=None, flip=False):
        group = self.slide.shapes.add_group_shape()
        child = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        child.text = "TRANSFORMED CHILD"
        xfrm = xp(group._element, "./p:grpSpPr/a:xfrm")[0]
        for tag, values in [("off", {"x": 5, "y": 3}), ("ext", {"cx": 4, "cy": 2}),
                            ("chOff", {"x": 1, "y": 1}), ("chExt", {"cx": 2, "cy": 1})]:
            node = xp(xfrm, "./a:" + tag)[0]
            for name, value in values.items():
                node.set(name, str(Inches(value)))
        if rotation is not None:
            xfrm.set("rot", str(rotation * 60000))
        if flip:
            xfrm.set("flipH", "1")
        return self.save()

    def test_group_translation_and_scale_produce_slide_bounds(self):
        result = extract(self.group_transform())["slides"][0]["objects"][0]["children"][0]
        self.assertEqual({"left_in": 5.0, "top_in": 3.0, "width_in": 4.0, "height_in": 2.0}, result["bounds"])
        self.assertEqual({"left_in": 1.0, "top_in": 1.0, "width_in": 2.0, "height_in": 1.0}, result.get("local_bounds"))
        self.assertEqual(result["bounds"], result.get("slide_bounds"))

    def test_group_rotation_has_correct_axis_aligned_slide_bounds(self):
        result = extract(self.group_transform(rotation=90, flip=True))["slides"][0]["objects"][0]["children"][0]
        self.assertEqual({"left_in": 6.0, "top_in": 2.0, "width_in": 2.0, "height_in": 4.0}, result["bounds"])
        self.assertEqual("known", result.get("transform", {}).get("status"))
        self.assertTrue(result.get("transform", {}).get("ancestor_transforms"))

    def test_invalid_group_transform_is_explicitly_unknown(self):
        self.group_transform()
        patch_part(self.path, "ppt/slides/slide1.xml", lambda root: xp(root, ".//p:grpSp/p:grpSpPr/a:xfrm/a:chExt")[0].set("cx", "0"))
        result = extract(self.path)["slides"][0]
        child = result["objects"][0]["children"][0]
        self.assertEqual("unknown", child.get("transform", {}).get("status"))
        self.assertIsNone(child.get("slide_bounds"))
        self.assertTrue(result["parse_limitations"])

    def test_master_content_is_included_and_respects_show_master_shapes(self):
        box = self.text("MASTER CONTENT")
        master = self.slide.slide_layout.slide_master
        master._element.spTree.append(copy.deepcopy(box._element))
        box._element.getparent().remove(box._element)
        result = extract(self.save())["slides"][0]
        self.assertIn("MASTER CONTENT", result["visible_text"])
        inherited = next(x for x in result["objects"] if x.get("text_frame", {}).get("text") == "MASTER CONTENT")
        self.assertEqual("master", inherited.get("xml_source", {}).get("kind"))
        self.assertTrue(inherited.get("inheritance", {}).get("effective"))
        patch_part(self.path, "ppt/slides/slide1.xml", lambda root: root.set("showMasterSp", "0"))
        disabled = extract(self.path)["slides"][0]
        self.assertNotIn("MASTER CONTENT", disabled["visible_text"])
        retained = next(x for x in disabled["objects"] if x.get("text_frame", {}).get("text") == "MASTER CONTENT")
        self.assertFalse(retained["visible_in_slideshow"])

    def test_notes_and_hidden_slide_are_distinct_from_hidden_objects(self):
        self.text("SLIDE CONTENT")
        self.slide.notes_slide.notes_text_frame.text = "SPEAKER NOTES"
        self.slide._element.set("show", "0")
        result = extract(self.save())["slides"][0]
        self.assertTrue(result["hidden_in_slideshow"])
        self.assertIn("SLIDE CONTENT", result["visible_text"])
        self.assertNotIn("SPEAKER NOTES", result["visible_text"])
        self.assertIn("SPEAKER NOTES", result["notes_not_visible"])
        self.assertTrue(result.get("visibility", {}).get("objects_on_hidden_slide_retained"))

    def test_first_body_is_only_a_title_candidate(self):
        self.text("This long body paragraph describes an experiment and is not a verified title.", top=3)
        result = extract(self.save())["slides"][0]
        self.assertEqual("candidate", result.get("title_status"))
        self.assertEqual("low", result.get("title_confidence"))
        self.assertTrue(result.get("title_candidates"))

    def test_xy_chart_has_complete_x_y_points_and_cache_sources(self):
        data = XyChartData()
        series = data.add_series("XY")
        series.add_data_point(1, 2)
        series.add_data_point(10, 20)
        self.slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(1), Inches(1), Inches(4), Inches(3), data)
        chart = extract(self.save())["slides"][0]["objects"][0]["chart"]
        points = chart["series"][0].get("points", [])
        self.assertEqual([(1.0, 2.0), (10.0, 20.0)], [(p.get("x"), p.get("y")) for p in points])
        self.assertEqual([2.0, 20.0], chart["series"][0]["values"])
        self.assertEqual("scatterChart", chart.get("plots", [{}])[0].get("type"))
        self.assertIn("numCache", str(chart["series"][0].get("data_sources", {})))

    def test_bubble_chart_keeps_size_and_missing_cache_is_partial(self):
        data = BubbleChartData()
        series = data.add_series("Bubbles")
        series.add_data_point(1, 2, 3)
        self.slide.shapes.add_chart(XL_CHART_TYPE.BUBBLE, Inches(1), Inches(1), Inches(4), Inches(3), data)
        chart = extract(self.save())["slides"][0]["objects"][0]["chart"]
        self.assertEqual([3.0], [p.get("size") for p in chart["series"][0].get("points", [])])
        patch_part(self.path, "ppt/charts/chart1.xml", lambda root: xp(root, ".//c:bubbleSize/c:numRef/c:numCache")[0].getparent().remove(xp(root, ".//c:bubbleSize/c:numRef/c:numCache")[0]))
        partial = extract(self.path)["slides"][0]["objects"][0]["chart"]
        self.assertEqual("partial", partial.get("parse_status"))
        self.assertTrue(partial["limitations"])
        self.assertIsNone(partial["series"][0]["points"][0]["size"])

    def test_combination_chart_preserves_each_plot_type(self):
        data = CategoryChartData()
        data.categories = ["A", "B"]
        data.add_series("Bars", [1, 2])
        data.add_series("Line", [3, 4])
        self.slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data)
        self.save()
        def combine(root):
            bar = xp(root, ".//c:barChart")[0]
            line = copy.deepcopy(bar)
            line.tag = "{%s}lineChart" % NS["c"]
            line.remove(xp(line, "./c:ser")[0])
            bar.remove(xp(bar, "./c:ser")[1])
            for node in list(line):
                if etree.QName(node).localname in ("barDir", "gapWidth", "overlap"):
                    line.remove(node)
            bar.getparent().append(line)
        patch_part(self.path, "ppt/charts/chart1.xml", combine)
        chart = extract(self.path)["slides"][0]["objects"][0]["chart"]
        self.assertEqual(["barChart", "lineChart"], [p.get("type") for p in chart.get("plots", [])])
        self.assertEqual(["barChart", "lineChart"], [s.get("plot_type") for s in chart["series"]])
        self.assertEqual([[1.0, 2.0], [3.0, 4.0]], [s["values"] for s in chart["series"]])

    def test_picture_instances_export_originals_and_expose_crop(self):
        image_data = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScLbtAAAAABJRU5ErkJggg==")
        for left in [1, 4]:
            pic = self.slide.shapes.add_picture(io.BytesIO(image_data), Inches(left), Inches(1), Inches(2), Inches(2))
            pic.crop_left = 0.25
            pic.rotation = 90
        self.save()
        first = extract(self.path)["slides"][0]["objects"][0]["picture"]
        self.assertEqual([1, 1], first.get("pixel_size"))
        inventory = extract(self.path, media_dir=Path(self.folder.name) / "media")
        pictures = [x["picture"] for x in inventory["slides"][0]["objects"] if "picture" in x]
        self.assertEqual(0.25, pictures[0]["crop"]["left"])
        self.assertEqual(90.0, pictures[0]["rotation_degrees"])
        self.assertEqual(pictures[0]["media_part"], pictures[1]["media_part"])
        self.assertNotEqual(pictures[0]["export_path"], pictures[1]["export_path"])
        for picture in pictures:
            self.assertEqual(image_data, Path(picture["export_path"]).read_bytes())

    def test_unknown_graphic_data_is_reported_instead_of_silent_success(self):
        table = self.slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(2), Inches(1))
        table.table.cell(0, 0).text = "UNSUPPORTED GRAPHIC"
        self.save()
        patch_part(self.path, "ppt/slides/slide1.xml", lambda root: xp(root, ".//a:graphicData")[0].set("uri", "urn:unsupported:graphic"))
        result = extract(self.path)["slides"][0]
        self.assertEqual("partial", result["objects"][0].get("parse_status"))
        self.assertTrue(result["parse_limitations"])

    def test_corrupt_input_fails_with_input_context(self):
        self.path.write_bytes(b"not a zip")
        with self.assertRaises(Exception) as caught:
            extract(self.path)
        self.assertIsInstance(caught.exception, ValueError)
        self.assertIn("fixture.pptx", str(caught.exception))


@unittest.skipUnless((ROOT / "Paper-Presentation-Tips-editable-22-v2.pptx").is_file() and (ROOT / "组会1-黄誉铭.pptx").is_file(), "local real-deck acceptance inputs are absent")
class RealReaderAcceptance(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.example_path = ROOT / "Paper-Presentation-Tips-editable-22-v2.pptx"
        cls.talk_path = ROOT / "组会1-黄誉铭.pptx"
        cls.before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in [cls.example_path, cls.talk_path]}
        cls.examples = extract(cls.example_path)
        cls.talk = extract(cls.talk_path)

    def test_counts_and_original_paths_remain_compatible(self):
        self.assertEqual(22, self.examples["slide_count"])
        objects = list(walk([o for slide in self.examples["slides"] for o in slide["objects"]]))
        local = [o for o in objects if o.get("xml_source", {}).get("kind", "slide") == "slide"]
        self.assertEqual(6, sum("table" in o for o in local))
        self.assertEqual(3, sum("chart" in o for o in local))
        self.assertEqual(19, sum("picture" in o for o in local))
        self.assertEqual(19, self.examples["slides"][8]["objects"][17]["id"])
        self.assertEqual("shape-18", self.examples["slides"][8]["objects"][17]["path"])
        for path, digest in self.before.items():
            self.assertEqual(digest, hashlib.sha256(path.read_bytes()).hexdigest())

    def test_talk_nine_and_fifteen_have_missing_body_and_formulas(self):
        self.assertEqual(27, self.talk["slide_count"])
        for slide_no, body_quote in [(9, "输入扰动的传播主要可以分为两部分"), (15, "通过将引理")]:
            slide = self.talk["slides"][slide_no - 1]
            self.assertTrue(any(body_quote in text for text in slide["visible_text"]), (slide_no, slide["visible_text"]))
            formulae = [m for obj in walk(slide["objects"]) for m in obj.get("math", [])]
            self.assertTrue(formulae)
            supplements = [o for o in slide["objects"] if o.get("xml_source", {}).get("alternate_content")]
            self.assertTrue(supplements)
            self.assertTrue(all(o.get("xml_source", {}).get("node_path") for o in supplements))
            self.assertTrue(slide["parse_limitations"])

    def test_talk_header_candidates_do_not_use_the_first_body(self):
        for slide_no, title in [(4, "背景与动机"), (18, "实验设置"), (19, "整体评估")]:
            slide = self.talk["slides"][slide_no - 1]
            self.assertEqual(title, slide["title"])
            self.assertIn(slide.get("title_status"), ["confirmed", "candidate"])
            self.assertTrue(slide.get("title_source", {}).get("object"))

    def test_repeated_left_section_heading_beats_equation_number(self):
        slide = self.talk["slides"][14]
        self.assertEqual("线性自注意力的输入鲁棒性", slide["title"])
        self.assertEqual("shape-1", slide["title_source"]["object"])

    def test_known_foreground_picture_overlap_is_only_a_visual_cue(self):
        for slide_no, picture_path, cover_path in [(9, "shape-17", "shape-18"), (10, "shape-23", "shape-24")]:
            slide = self.examples["slides"][slide_no - 1]
            cues = slide.get("occlusion_cues", [])
            matches = [c for c in cues if c.get("picture") == picture_path and c.get("foreground") == cover_path]
            self.assertTrue(matches, (slide_no, cues))
            self.assertEqual("needs_visual_confirmation", matches[0]["status"])
            self.assertFalse(matches[0].get("confirmed_issue", False))


if __name__ == "__main__":
    unittest.main()

"""A failed review import must not mix artifacts or re-render viewed slides."""
import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches
import analyze_pptx as app


class SnapshotCLI(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.deck = self.root / 'input.pptx'
        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = 'CORE CONTENT'
        pres.save(self.deck)
        self.out = self.root / 'output'

    def run_cli(self, *args):
        with patch.object(sys, 'argv', ['analyze_pptx.py', str(self.deck), '--out', str(self.out), *args]):
            return app.main()

    @staticmethod
    def mock_render(deck, out, slides, **kwargs):
        folder = out / 'renders'
        folder.mkdir(parents=True, exist_ok=True)
        for slide in slides:
            image = folder / f"slide-{slide['slide']:02d}.png"
            image.write_bytes(b'known-image-fixture')
            slide['render'] = {'status': 'ok', 'mapping_status': 'verified', 'pdf_page': slide['slide'],
                               'sha256': hashlib.sha256(image.read_bytes()).hexdigest(),
                               'image': str(image), 'source': 'test_renderer', 'pdf_text': 'CORE CONTENT'}
        entries = [{**s['render'], 'original_slide': s['slide'], 'slide_id': s.get('slide_id'),
                    'slide_part': s.get('slide_part'), 'hidden_in_slideshow': s.get('hidden_in_slideshow', False)}
                   for s in slides]
        (out/'render_manifest.json').write_text(json.dumps({'slides': entries}), encoding='utf-8')
        return []

    def test_failed_import_preserves_entire_successful_run(self):
        with patch.object(app, 'render', self.mock_render):
            self.assertEqual(0, self.run_cli())
        before = {p.relative_to(self.out).as_posix(): p.read_bytes() for p in self.out.rglob('*') if p.is_file()}
        # A second input aimed at the same output used to replace inventory
        # before review validation, while leaving the first report behind.
        changed = Presentation()
        changed.slides.add_slide(changed.slide_layouts[6]).shapes.add_textbox(
            Inches(1), Inches(1), Inches(6), Inches(1)).text = 'DIFFERENT INPUT'
        changed.save(self.deck)
        bad = self.root / 'bad.json'
        bad.write_text(json.dumps({'review_schema_version': 4, 'input_sha256': 'wrong', 'slides': []}), encoding='utf-8')
        try:
            with patch.object(app, 'render', self.mock_render):
                result = self.run_cli('--review', str(bad))
        except (ValueError, SystemExit):
            result = 2
        after = {p.relative_to(self.out).as_posix(): p.read_bytes() for p in self.out.rglob('*') if p.is_file()}
        self.assertEqual(2, result)
        self.assertEqual(before, after)

    def test_formal_import_reuses_same_snapshot(self):
        with patch.object(app, 'render', self.mock_render):
            self.assertEqual(0, self.run_cli())
        template = json.loads((self.out / 'review_template.json').read_text(encoding='utf-8'))
        template.update(preparation_only=False, review_completion='complete', quality_status='无需修改',
                        blocking_unknowns=[])
        for entry in template['slides']:
            entry.update(review_completion='complete', quality_status='无需修改', status='无需修改',
                         blocking_unknowns=[], basis='单一大字观点清楚，无额外表达障碍。',
                         role='观点', issues=[], media_reviews=[])
            entry['assessment'] = {'core_message': '说明一个观点', 'content_organization': '单一概念',
                                   'visual_structure': '字号和留白有效', 'reading_path': '先标题后观点'}
            entry['tip_checks'] = [{'tip': i, 'result': 'adequate', 'basis': '本页单一观点结构有效'} for i in range(1, 12)]
            if 'render_fidelity' in entry:
                entry['render_fidelity'].update(status='confirmed', evidence='测试夹具的已知画面；不代表真实视觉验收',
                                                blocking_unknowns=[])
        record = self.out / 'review.json'
        record.write_text(json.dumps(template, ensure_ascii=False), encoding='utf-8')
        snapshot_before = (self.out/'snapshot_manifest.json').read_bytes()
        with patch.object(app, 'render', side_effect=AssertionError('import re-rendered immutable evidence')):
            self.assertEqual(0, self.run_cli('--review', str(record), '--require-evidence'))
        self.assertEqual(snapshot_before, (self.out/'snapshot_manifest.json').read_bytes())
        result = json.loads((self.out/'diagnosis.json').read_text(encoding='utf-8'))
        self.assertEqual('visual_review', result['review_mode'])
        self.assertEqual(template['snapshot_id'], result['snapshot_id'])

    def test_render_failure_publishes_only_incomplete_state(self):
        def fail(deck, out, slides, **kwargs):
            for slide in slides:
                slide['render'] = {'status': 'failed', 'mapping_status': 'unverified', 'reason': 'renderer unavailable'}
            return ['renderer unavailable']
        with patch.object(app, 'render', fail):
            self.assertEqual(2, self.run_cli())
        result = json.loads((self.out/'diagnosis.json').read_text(encoding='utf-8'))
        self.assertEqual('failed', result['review_completion'])
        self.assertNotEqual('visual_review', result['review_mode'])
        self.assertNotIn('无需修改', [entry['status'] for entry in result['slides']])

    def test_all_failed_review_stays_failed_in_every_output(self):
        with patch.object(app, 'render', self.mock_render):
            self.assertEqual(0, self.run_cli())
        review = json.loads((self.out / 'review_template.json').read_text(encoding='utf-8'))
        review.update(preparation_only=False, review_completion='failed', quality_status='未完成',
                      blocking_unknowns=['本次无法核对主体字形'])
        for entry in review['slides']:
            entry.update(review_completion='failed', quality_status='未完成', status='未完成',
                         blocking_unknowns=['本次无法核对主体字形'],
                         basis='导出成功但无法核对中文实际画面，本次判断失败。')
            if 'render_fidelity' in entry:
                entry['render_fidelity'].update(status='failed', evidence='中文主体未显示',
                    blocking_unknowns=['中文主体未显示'])
        record = self.out / 'review.json'
        record.write_text(json.dumps(review, ensure_ascii=False), encoding='utf-8')
        with patch.object(app, 'render', side_effect=AssertionError('must reuse snapshot')):
            self.assertEqual(2, self.run_cli('--review', str(record)))
        accepted = json.loads((self.out / 'accepted_review.json').read_text(encoding='utf-8'))
        result = json.loads((self.out / 'diagnosis.json').read_text(encoding='utf-8'))
        self.assertEqual('failed', accepted['review_completion'])
        self.assertEqual('failed', result['review_completion'])
        self.assertEqual('failed_visual_review', result['review_mode'])
        self.assertIn('整套审阅完成度：failed', (self.out / 'report.md').read_text(encoding='utf-8'))

    def test_mixed_complete_and_failed_review_is_partial(self):
        pres = Presentation(str(self.deck))
        pres.slides.add_slide(pres.slide_layouts[6]).shapes.add_textbox(
            Inches(1), Inches(1), Inches(6), Inches(1)).text = 'CORE CONTENT'
        pres.save(self.deck)
        with patch.object(app, 'render', self.mock_render):
            self.assertEqual(0, self.run_cli())
        review = json.loads((self.out / 'review_template.json').read_text(encoding='utf-8'))
        review.update(preparation_only=False, review_completion='partial', quality_status='未完成',
                      blocking_unknowns=['第二页无法核对'])
        first, second = review['slides']
        first.update(review_completion='complete', quality_status='无需修改', status='无需修改',
                     blocking_unknowns=[], basis='合成夹具已知内容，不是实际视觉验收。')
        first['render_fidelity'].update(status='confirmed', evidence='合成测试夹具的已知内容', blocking_unknowns=[])
        second.update(review_completion='failed', quality_status='未完成', status='未完成',
                      blocking_unknowns=['本次无法核对第二页'])
        record = self.out / 'review.json'
        record.write_text(json.dumps(review, ensure_ascii=False), encoding='utf-8')
        self.assertEqual(2, self.run_cli('--review', str(record)))
        result = json.loads((self.out / 'diagnosis.json').read_text(encoding='utf-8'))
        self.assertEqual('partial', result['review_completion'])
        self.assertEqual('partial_visual_review', result['review_mode'])
        self.assertEqual(['complete', 'failed'], [p['review_completion'] for p in result['slides']])


if __name__ == '__main__':
    unittest.main()

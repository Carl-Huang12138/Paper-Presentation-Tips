"""Render mappings and immutable evidence snapshots; real renderer probes are explicit."""

import copy
import hashlib
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest import mock

import fitz
from pptx import Presentation
from pptx.util import Inches

try:
    import pptx_runtime as runtime
except ModuleNotFoundError:
    runtime = None

ROOT = Path(__file__).resolve().parents[1]
try:
    SOFFICE = runtime.resolve_soffice() if runtime is not None else Path("C:/Program Files/LibreOffice/program/soffice.com")
except runtime.RenderError:
    SOFFICE = ROOT / "tests/not-installed-soffice"


def deck_fixture(folder, hidden=(False, True, False)):
    """Unique page markers make a wrong page assignment observable."""
    source = folder / "mapping.pptx"
    pres = Presentation()
    slides = []
    for number, is_hidden in enumerate(hidden, 1):
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text = f"PAGE {number} MARKER{number:02d}"
        if is_hidden:
            slide._element.set("show", "0")
        slides.append({"slide": number, "slide_id": slide.slide_id,
                       "slide_part": str(slide.part.partname),
                       "hidden_in_slideshow": is_hidden,
                       "title": box.text, "visible_text": [box.text],
                       "objects": [{"path": "shape-1", "id": 2, "type": "TEXT_BOX",
                                    "visible_in_slideshow": True,
                                    "text_frame": {"text": box.text}}],
                       "parse_limitations": []})
    pres.save(source)
    return source, slides


def inventory_fixture(source, slides):
    return {"input": str(source), "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "slide_count": len(slides), "slides": slides, "output_root": "Z:/not-staging"}


def pdf_fixture(folder, markers):
    path = folder / "export.pdf"
    with fitz.open() as document:
        for marker in markers:
            page = document.new_page(width=720, height=540)
            page.insert_text((72, 72), marker)
        document.save(path)
    return path


def rendered_fixture(folder):
    source, slides = deck_fixture(folder, (False,))
    pdf = pdf_fixture(folder, ["PAGE 1 MARKER01"])
    out = folder / "out"
    with mock.patch("pptx_runtime.resolve_soffice", return_value=SOFFICE):
        with mock.patch("pptx_runtime._renderer_preflight", return_value={"version": "test PDF boundary"}):
            with mock.patch("pptx_runtime._convert_to_pdf", return_value=[pdf]):
                issues = runtime.render(source, out, slides)
    if issues:
        raise AssertionError(issues)
    return source, out, inventory_fixture(source, slides)


class OriginalRenderRegression(unittest.TestCase):
    @unittest.skipUnless(SOFFICE.is_file(), "real LibreOffice is unavailable")
    def test_hidden_middle_has_its_own_render_in_all_scope(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT, ignore_cleanup_errors=True) as name:
            folder = Path(name)
            source, slides = deck_fixture(folder)
            if runtime is None:
                from analyze_pptx import render
                with mock.patch.object(tempfile, "tempdir", str(folder)):
                    render(source, folder / "out", slides)
            else:
                runtime.render(source, folder / "out", slides, soffice=SOFFICE)
            self.assertEqual("ok", slides[1]["render"]["status"])
            self.assertIn("MARKER02", slides[1]["render"]["pdf_text"])
            self.assertEqual("ok", slides[2]["render"]["status"])
            self.assertIn("MARKER03", slides[2]["render"]["pdf_text"])

    @unittest.skipUnless(SOFFICE.is_file(), "real LibreOffice is unavailable")
    def test_deep_output_path_does_not_extend_renderer_profile_path(self):
        with tempfile.TemporaryDirectory(prefix='runtime_test_', dir=ROOT, ignore_cleanup_errors=True) as name:
            folder = Path(name) / ('nested-' + 'a'*30) / ('nested-' + 'b'*30)
            folder.mkdir(parents=True)
            source, slides = deck_fixture(folder)
            issues = runtime.render(source, folder / 'out', slides, soffice=SOFFICE)
            self.assertEqual([], issues)
            self.assertEqual(['ok', 'ok', 'ok'], [s['render']['status'] for s in slides])
            self.assertIn('MARKER02', slides[1]['render']['pdf_text'])


class RuntimeContracts(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(runtime, "pptx_runtime module must implement the runtime contract")

    def test_failed_probe_is_retried_instead_of_poisoning_later_renders(self):
        with tempfile.TemporaryDirectory() as name:
            folder = Path(name)
            executable = folder / 'renderer'
            executable.write_bytes(b'known renderer fixture')
            with mock.patch.dict(runtime._CAPABILITY_CACHE, {}, clear=True), mock.patch(
                    'pptx_runtime._convert_to_pdf', side_effect=[
                        runtime.RenderError('first environment failure'), runtime.RenderError('second attempt reached')]):
                first = runtime.probe_hidden_export(executable, {'version': 'test'}, folder / 'first')
                second = runtime.probe_hidden_export(executable, {'version': 'test'}, folder / 'second')
                self.assertIn('first environment failure', first['error'])
                self.assertIn('second attempt reached', second['error'])

    def test_rendered_picture_crop_uses_page_coordinates_and_keeps_occlusion(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            folder = Path(name)
            source, slides = deck_fixture(folder, (False,))
            slides[0]["objects"].append({"path": "group/1", "visible_in_slideshow": True,
                "slide_bounds": {"left_in": 1, "top_in": 1, "width_in": 2, "height_in": 1},
                "picture": {"pixel_size": [1, 1]}})
            slides[0]["objects"].append({"path": "unknown", "visible_in_slideshow": True,
                "slide_bounds": None, "picture": {}})
            pdf = folder / "export.pdf"
            with fitz.open() as doc:
                page = doc.new_page(width=720, height=540)
                page.draw_rect(fitz.Rect(72, 72, 216, 144), color=None, fill=(1, 0, 0))
                page.draw_rect(fitz.Rect(144, 72, 216, 144), color=None, fill=(1, 1, 1))
                doc.save(pdf)
            with mock.patch("pptx_runtime.resolve_soffice", return_value=SOFFICE), \
                 mock.patch("pptx_runtime._renderer_preflight", return_value={"version": "test PDF"}), \
                 mock.patch("pptx_runtime._convert_to_pdf", return_value=[pdf]):
                self.assertEqual([], runtime.render(source, folder / "out", slides))
            crop = slides[0]["objects"][1]["picture"]["render_crop"]
            pix = fitz.Pixmap(str(folder / "out" / crop["image"]))
            self.assertEqual((576, 288), (pix.width, pix.height))
            self.assertEqual((255, 0, 0), pix.pixel(100, 100))
            self.assertEqual((255, 255, 255), pix.pixel(450, 100))
            self.assertEqual("unknown_geometry", slides[0]["objects"][2]["picture"]["render_crop"]["status"])
            self.assertEqual("rendered_page_clip", crop["source"])

    def test_packet_entry_and_reader_media_paths_are_portable(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            source, out, inventory = rendered_fixture(Path(name))
            media = out / "media" / "source.bin"
            media.parent.mkdir()
            media.write_bytes(b"original")
            inventory["media_manifest"] = [{"path": str(media)}]
            snapshot = runtime.create_snapshot(out, inventory)
            self.assertEqual("media/source.bin", inventory["media_manifest"][0]["path"])
            runtime.create_review_packet(out, inventory, snapshot)
            self.assertIn("../review_packet.md", (out / "review_packet" / "README.md").read_text(encoding="utf-8"))

    def test_explicit_missing_executable_does_not_fall_back(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            folder = Path(name)
            source, slides = deck_fixture(folder, (False,))
            issues = runtime.render(source, folder / "out", slides,
                                    soffice=folder / "missing-soffice")
            self.assertEqual("failed", slides[0]["render"]["status"])
            self.assertTrue(any("soffice" in issue.lower() for issue in issues))
            self.assertFalse((folder / "out" / "renders" / "slide-01.png").exists())

    def test_executable_priority_and_path_discovery(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            folder = Path(name)
            explicit, env_path, path_path = [folder / name for name in ("explicit", "env", "path")]
            for path in (explicit, env_path, path_path):
                path.write_text("test executable", encoding="utf-8")
                path.chmod(0o755)
            with mock.patch.dict(os.environ, {"PPTX_SOFFICE": str(env_path)}):
                with mock.patch("pptx_runtime.shutil.which", return_value=str(path_path)):
                    self.assertEqual(explicit.resolve(), runtime.resolve_soffice(explicit))
                    self.assertEqual(env_path.resolve(), runtime.resolve_soffice())
            with mock.patch.dict(os.environ, {"PPTX_SOFFICE": ""}):
                with mock.patch("pptx_runtime.shutil.which", return_value=str(path_path)):
                    self.assertEqual(path_path.resolve(), runtime.resolve_soffice())

    def test_empty_inventory_does_not_match_nonempty_pdf_text(self):
        compared = runtime.compare_text({"objects": [], "visible_text": []}, "MASTER CONTENT")
        self.assertEqual("partially_compared", compared["comparison_status"])
        self.assertTrue(compared["pdf_text_unattributed"])

    def test_text_comparison_checks_both_directions(self):
        slide = {"objects": [{"path": "shape-1", "text_frame": {"text": "ALPHA BETA"}}],
                 "visible_text": ["ALPHA BETA"]}
        matched = runtime.compare_text(slide, "ALPHA BETA")
        self.assertEqual("matched", matched["comparison_status"])
        mismatch = runtime.compare_text(slide, "ALPHA MASTER")
        self.assertEqual("partially_compared", mismatch["comparison_status"])
        self.assertTrue(mismatch["object_text_absent_in_pdf"])
        self.assertTrue(mismatch["pdf_text_unattributed"])

    def test_blank_text_is_not_testable(self):
        compared = runtime.compare_text({"objects": [], "visible_text": []}, "")
        self.assertEqual("not_testable", compared["comparison_status"])

    def test_pdf_count_mismatch_never_marks_shifted_pages_ok(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            folder = Path(name)
            source, slides = deck_fixture(folder, (False, False, False))
            pdf = pdf_fixture(folder, ["PAGE 1 MARKER01", "PAGE 3 MARKER03"])
            with mock.patch("pptx_runtime.resolve_soffice", return_value=SOFFICE):
                with mock.patch("pptx_runtime._renderer_preflight", return_value={"version": "test PDF boundary"}):
                    with mock.patch("pptx_runtime._convert_to_pdf", return_value=[pdf]):
                        issues = runtime.render(source, folder / "out", slides)
            self.assertTrue(issues)
            self.assertEqual(["mapping_unverified"] * 3, [s["render"]["status"] for s in slides])
            self.assertTrue(all("image" not in s["render"] for s in slides))

    def test_matching_count_cannot_replace_hidden_export_capability_probe(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            folder = Path(name)
            source, slides = deck_fixture(folder, (True, True, True))
            pdf = pdf_fixture(folder, ["PAGE 1 MARKER01", "PAGE 2 MARKER02", "PAGE 3 MARKER03"])
            unverified = {"export_hidden_slides": False, "exclude_hidden_slides": False, "cases": []}
            with mock.patch("pptx_runtime.resolve_soffice", return_value=SOFFICE):
                with mock.patch("pptx_runtime._renderer_preflight", return_value={"version": "test PDF boundary"}):
                    with mock.patch("pptx_runtime.probe_hidden_export", return_value=unverified):
                        with mock.patch("pptx_runtime._convert_to_pdf", return_value=[pdf]):
                            issues = runtime.render(source, folder / "out", slides)
            self.assertTrue(issues)
            self.assertTrue(all(s["render"]["status"] == "mapping_unverified" for s in slides))

    def test_slides_scope_all_hidden_is_explicitly_excluded(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            folder = Path(name)
            source, slides = deck_fixture(folder, (True, True))
            issues = runtime.render(source, folder / "out", slides, render_scope="slides")
            self.assertEqual([], issues)
            self.assertEqual(["excluded", "excluded"], [s["render"]["status"] for s in slides])
            self.assertTrue(all("pdf_page" not in s["render"] for s in slides))

    def test_dpi_and_retained_pdf_are_observable(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            source, out, inventory = rendered_fixture(Path(name))
            image = fitz.Pixmap(out / inventory["slides"][0]["render"]["image"])
            self.assertEqual((1440, 1080), (image.width, image.height))
            self.assertTrue((out / "renders" / "presentation.pdf").is_file())
            manifest = json.loads((out / "render_manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(1, manifest["slides"][0]["pdf_page"])
            self.assertEqual("verified", manifest["mapping_status"])


class SnapshotContracts(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(runtime)
        self.assertTrue(hasattr(runtime, "create_snapshot"), "immutable snapshot creator is missing")

    def test_snapshot_reuse_reads_same_bytes_without_rendering(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            source, out, inventory = rendered_fixture(Path(name))
            snapshot = runtime.create_snapshot(out, inventory, run_id="run-test", parser_version="test-parser")
            self.assertEqual(snapshot["snapshot_id"], inventory["snapshot_id"])
            self.assertEqual("run-test", snapshot["run_id"])
            (out / "report.md").write_text("review may change", encoding="utf-8")
            with mock.patch("pptx_runtime._convert_to_pdf", side_effect=AssertionError("snapshot rerendered")):
                checked = runtime.verify_snapshot(out, inventory["sha256"])
            self.assertTrue(checked["verified"])
            self.assertEqual(snapshot["snapshot_id"], checked["snapshot_id"])
            self.assertNotIn("snapshot_manifest.json", [item["path"] for item in snapshot["immutable_files"]])

    def test_image_tampering_invalidates_snapshot(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            source, out, inventory = rendered_fixture(Path(name))
            runtime.create_snapshot(out, inventory)
            image = out / "renders" / "slide-01.png"
            image.write_bytes(image.read_bytes() + b"changed")
            with self.assertRaises(runtime.SnapshotError):
                runtime.verify_snapshot(out, inventory["sha256"])

    def test_snapshot_rejects_wrong_input_hash_and_missing_file_binding(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            source, out, inventory = rendered_fixture(Path(name))
            snapshot = runtime.create_snapshot(out, inventory)
            with self.assertRaises(runtime.SnapshotError):
                runtime.verify_snapshot(out, "0" * 64)
            snapshot["immutable_files"] = [item for item in snapshot["immutable_files"]
                                           if item["path"] != "renders/slide-01.png"]
            (out / "snapshot_manifest.json").write_text(json.dumps(snapshot), encoding="utf-8")
            with self.assertRaises(runtime.SnapshotError):
                runtime.verify_snapshot(out, inventory["sha256"])

    def test_snapshot_refuses_external_image_paths(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            source, out, inventory = rendered_fixture(Path(name))
            inventory["slides"][0]["render"]["image"] = str(source)
            with self.assertRaises(runtime.SnapshotError):
                runtime.create_snapshot(out, inventory)

    def test_packet_is_unreviewed_and_cannot_claim_quality_pass(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            source, out, inventory = rendered_fixture(Path(name))
            snapshot = runtime.create_snapshot(out, inventory)
            template = runtime.create_review_packet(out, inventory, snapshot)
            entry = template["slides"][0]
            self.assertEqual(4, template["review_schema_version"])
            self.assertEqual(snapshot["snapshot_id"], template["snapshot_id"])
            self.assertEqual("partial", entry["review_completion"])
            self.assertEqual("未完成", entry["quality_status"])
            self.assertTrue(entry["blocking_unknowns"])
            self.assertEqual(11, len(entry["tip_checks"]))
            self.assertTrue((out / "review_packet.md").is_file())

    def test_failed_render_still_prepares_explicitly_incomplete_packet(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            folder = Path(name)
            source, slides = deck_fixture(folder, (False,))
            out = folder / "out"
            runtime.render(source, out, slides, soffice=folder / "missing-soffice")
            inventory = inventory_fixture(source, slides)
            snapshot = runtime.create_snapshot(out, inventory)
            template = runtime.create_review_packet(out, inventory, snapshot)
            self.assertEqual("未完成", template["slides"][0]["status"])
            self.assertTrue(template["slides"][0]["blocking_unknowns"])
            self.assertFalse(runtime.verify_snapshot(out, inventory["sha256"])["render_complete"])

    def test_snapshot_binds_original_media_and_packet_keeps_core_unreviewed(self):
        with tempfile.TemporaryDirectory(prefix="runtime_test_", dir=ROOT) as name:
            source, out, inventory = rendered_fixture(Path(name))
            media = out / "media" / "original.png"
            media.parent.mkdir()
            image = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 30, 20), False)
            image.clear_with(255)
            image.save(media)
            inventory["slides"][0]["objects"].append({"path": "shape-2", "id": 3,
                "visible_in_slideshow": True,
                "picture": {"sha256": hashlib.sha256(media.read_bytes()).hexdigest(),
                            "original_media": str(media), "media_path": str(media),
                            "export_path": str(media), "pixel_size": [30, 20],
                            "crop": {"left": 0.1, "top": 0, "right": 0, "bottom": 0},
                            "rotation_degrees": 10}})
            snapshot = runtime.create_snapshot(out, inventory)
            self.assertIn("media/original.png", [item["path"] for item in snapshot["immutable_files"]])
            self.assertEqual("media/original.png", inventory["slides"][0]["objects"][1]["picture"]["original_media"])
            template = runtime.create_review_packet(out, inventory, snapshot)
            record = template["slides"][0]["media_reviews"][0]
            self.assertEqual("core", record["importance"])
            self.assertEqual("not_reviewed", record["status"])
            self.assertTrue(record["blocking_unknowns"])
            media.write_bytes(media.read_bytes() + b"changed")
            with self.assertRaises(runtime.SnapshotError):
                runtime.verify_snapshot(out, inventory["sha256"])


if __name__ == "__main__":
    unittest.main()

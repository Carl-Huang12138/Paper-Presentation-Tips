"""Install the entire skill, then run it outside the source repository."""
from pathlib import Path
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
INSTALLER = ROOT / 'tools/install_skill.py'
SOURCE = ROOT / 'skills/academic-pptx-review'
NAME = 'academic-pptx-review'


class PortableSkill(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='skill-portable-')
        self.root = Path(self.temp.name).resolve()

    def tearDown(self):
        self.temp.cleanup()

    def command(self, *args):
        env = dict(os.environ, PYTHONUTF8='1')
        return subprocess.run([sys.executable, str(INSTALLER), *map(str, args)],
                              cwd=self.root, env=env, capture_output=True, text=True, encoding='utf-8')

    def test_complete_bundle_installs_for_each_client(self):
        for agent, directory in [('codex', '.agents'), ('claude', '.claude'), ('opencode', '.opencode')]:
            project = self.root / agent
            result = self.command('--agent', agent, '--project', project)
            self.assertEqual(0, result.returncode, result.stderr)
            target = project / directory / 'skills' / NAME
            for relative in ['SKILL.md', 'requirements.txt', 'scripts/analyze_pptx.py',
                             'scripts/pptx_reader.py', 'scripts/pptx_runtime.py',
                             'scripts/pptx_report.py', 'scripts/pptx_review_contract.py',
                             'references/usage.md', 'references/presentation-tips.md',
                             'references/frameworks.md']:
                self.assertEqual((SOURCE / relative).read_bytes(), (target / relative).read_bytes())
            self.assertFalse(list(target.rglob('*.pyc')))

    def test_standalone_extraction_needs_no_repository_or_cwd(self):
        target = self.root / 'generic' / NAME
        result = self.command('--destination', target)
        self.assertEqual(0, result.returncode, result.stderr)
        deck = self.root / 'input.pptx'
        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = 'Portable evidence'
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
        table.cell(0, 0).text, table.cell(0, 1).text = 'Base', '42'
        pres.save(deck)
        digest = hashlib.sha256(deck.read_bytes()).hexdigest()
        script = ('import sys,json; from pathlib import Path; '
                  'sys.path.insert(0,sys.argv[1]); import analyze_pptx as app; '
                  'print(json.dumps(app.extract(Path(sys.argv[2]))))')
        result = subprocess.run([sys.executable, '-c', script, str(target / 'scripts'), str(deck)],
                                cwd=self.root, capture_output=True, text=True)
        self.assertEqual(0, result.returncode, result.stderr)
        inventory = json.loads(result.stdout)
        self.assertEqual(1, inventory['slide_count'])
        self.assertIn('Portable evidence', inventory['slides'][0]['visible_text'])
        self.assertEqual(digest, hashlib.sha256(deck.read_bytes()).hexdigest())
        self.assertTrue(any('table' in obj for obj in inventory['slides'][0]['objects']))

    def test_existing_install_is_preserved_without_explicit_overwrite(self):
        target = self.root / NAME
        target.mkdir()
        marker = target / 'my-notes.txt'
        marker.write_text('keep', encoding='utf-8')
        result = self.command('--destination', target)
        self.assertNotEqual(0, result.returncode)
        self.assertIn('already exists', result.stderr)
        self.assertEqual('keep', marker.read_text(encoding='utf-8'))
        self.assertFalse((target / 'SKILL.md').exists())

    def test_overwrite_retains_a_recoverable_previous_install(self):
        target = self.root / NAME
        target.mkdir()
        (target / 'my-notes.txt').write_text('keep', encoding='utf-8')
        result = self.command('--destination', target, '--overwrite')
        self.assertEqual(0, result.returncode, result.stderr)
        backups = list(self.root.glob(NAME + '.previous-*'))
        self.assertEqual(1, len(backups))
        self.assertEqual('keep', (backups[0] / 'my-notes.txt').read_text(encoding='utf-8'))
        self.assertTrue((target / 'scripts/analyze_pptx.py').is_file())

    def test_source_or_its_ancestor_is_not_an_install_target(self):
        before = (ROOT / 'analyze_pptx.py').read_bytes()
        result = self.command('--destination', SOURCE, '--overwrite')
        self.assertNotEqual(0, result.returncode)
        self.assertIn('separate from', result.stderr)
        self.assertEqual(before, (ROOT / 'analyze_pptx.py').read_bytes())

    def test_incomplete_source_fails_before_creating_target(self):
        incomplete = self.root / 'source' / NAME
        incomplete.mkdir(parents=True)
        (incomplete / 'SKILL.md').write_text('---\nname: academic-pptx-review\ndescription: review\n---\n', encoding='utf-8')
        target = self.root / 'target' / NAME
        result = self.command('--source', incomplete, '--destination', target)
        self.assertNotEqual(0, result.returncode)
        self.assertIn('Incomplete skill bundle', result.stderr)
        self.assertFalse(target.exists())

    def test_repository_entries_resolve_to_single_source(self):
        project = self.root / 'repo'
        source = project / 'skills' / NAME
        shutil.copytree(SOURCE, source, ignore=shutil.ignore_patterns('__pycache__', '*.pyc'))
        result = self.command('--sync-entries', '--source', source, '--project', project)
        self.assertEqual(0, result.returncode, result.stderr)
        for client in ['.agents', '.claude']:
            entry = project / client / 'skills' / NAME / 'SKILL.md'
            self.assertIn('../../../skills/academic-pptx-review/SKILL.md', entry.read_text(encoding='utf-8'))
            self.assertNotIn('import fitz', entry.read_text(encoding='utf-8'))
            self.assertTrue((entry.parent / '../../../skills/academic-pptx-review/SKILL.md').resolve().is_file())


if __name__ == '__main__':
    unittest.main()

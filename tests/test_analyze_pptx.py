"""Acceptance checks against the real 22-slide editable deck."""

import copy
import hashlib
import json
import tempfile
import unittest
from collections import Counter
from pathlib import Path

from analyze_pptx import extract, flatten, validate_review, preliminary, markdown
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "Paper-Presentation-Tips-editable-22-v2.pptx"


class SyntheticExtraction(unittest.TestCase):
    def test_table_fonts_and_unknown_sizes_stay_distinct(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "table-fonts.pptx"
            pres = Presentation()
            slide = pres.slides.add_slide(pres.slide_layouts[6])
            table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(5), Inches(1)).table
            table.cell(0, 0).text = "数据"
            table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size = Pt(9)
            table.cell(0, 1).text = "未知继承字号"
            pres.save(path)
            item = extract(path)["slides"][0]
            samples = item["typography"]["samples"]
            self.assertEqual(("shape-1:cell-1-1", 9), (samples[0]["object"], samples[0]["size_pt"]))
            self.assertIsNone(samples[1]["size_pt"])
            self.assertEqual(1, item["typography"]["unknown_size_runs"])

    def test_text_dominance_is_only_a_visual_review_cue(self):
        slide = {"slide": 1, "title": "方法", "objects": [], "visible_text": ["文字" * 150],
                 "render": {"status": "ok"}}
        result = preliminary(slide)
        self.assertEqual([], result["issues"])
        self.assertEqual("待视觉审阅", result["status"])
        self.assertIn("text_only_structure", [x["kind"] for x in result["uncertainties"]])

    def test_small_font_is_located_but_not_a_confirmed_defect(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "fonts.pptx"
            pres = Presentation()
            slide = pres.slides.add_slide(pres.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3))
            box.text = "正文需要从实际放映效果判断"
            box.text_frame.paragraphs[0].font.size = Pt(14)
            slide.notes_slide.notes_text_frame.text = "备注小字不属于放映画面"
            slide.notes_slide.notes_text_frame.paragraphs[0].font.size = Pt(8)
            pres.save(path)
            item = extract(path)["slides"][0]
            fonts = item["typography"]
            self.assertEqual(14, fonts["minimum_explicit_size_pt"])
            self.assertEqual("shape-1", fonts["samples"][0]["object"])
            self.assertEqual("paragraph", fonts["samples"][0]["size_source"])
            item["render"] = {"status": "ok"}
            cues = preliminary(item)
            self.assertEqual([], cues["issues"])
            self.assertIn("small_font", [x["kind"] for x in cues["uncertainties"]])

    def test_native_table_chart_and_notes_are_separated(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "mini.pptx"
            pres = Presentation()
            slide = pres.slides.add_slide(pres.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            box.text = "实验结论"
            table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
            table.cell(0, 0).text = "指标"
            table.cell(0, 1).text = "结果"
            table.cell(1, 0).text = "准确率"
            table.cell(1, 1).text = "84"
            data = CategoryChartData()
            data.categories = ["A", "B"]
            data.add_series("模型", [84, 96])
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(1),
                                   Inches(4), Inches(3), data)
            slide.notes_slide.notes_text_frame.text = "只供讲者看"
            pres.save(path)
            item = extract(path)["slides"][0]
            self.assertIn("84", item["visible_text"])
            self.assertNotIn("只供讲者看", item["visible_text"])
            self.assertIn("只供讲者看", item["notes_not_visible"])
            chart = next(obj["chart"] for obj in item["objects"] if "chart" in obj)
            self.assertEqual([84.0, 96.0], chart["series"][0]["values"])


class PolishReviewContract(unittest.TestCase):
    def setUp(self):
        self.inventory = {"sha256": "deck", "input": "test.pptx", "slide_count": 1,
                          "slides": [{"slide": 1, "title": "方法", "objects": [{"path": "shape-1"}],
                                      "render": {"status": "failed"}, "parse_limitations": []}]}
        self.issue = {"object": "shape-1", "category": "布局", "evidence": "三段等权内容挤在同页",
                      "impact": "核心结论难以辨认", "severity": "中", "confidence": "高",
                      "tips": [2], "direction": "拆成两页", "scope": "visual", "action": "split",
                      "revision_plan": {"goal": "每页一个任务", "layout": "左图右文",
                                        "steps": ["移动机制细节到第二页"],
                                        "typography": "正文 24–28 pt 起试，重新渲染确认",
                                        "check": "两个页面分别有清楚的阅读顺序"},
                      "split_plan": [
                          {"title": "方法解决什么", "message": "明确任务", "content": ["原页任务定义"],
                           "layout": "上标题，下任务图", "components": ["任务框", "输入输出箭头"]},
                          {"title": "方法如何工作", "message": "解释机制", "content": ["原页机制"],
                           "layout": "左流程右说明", "components": ["流程图", "机制注释"]}]}
        self.review = {"review_schema_version": 2, "input_sha256": "deck", "slides": [
            {"slide": 1, "role": "方法", "status": "有确定问题", "basis": "可见内容拥挤",
             "issues": [self.issue]}]}

    def test_split_requires_new_page_content_structure_and_components(self):
        validate_review(self.review, self.inventory)
        invalid = copy.deepcopy(self.review)
        del invalid["slides"][0]["issues"][0]["split_plan"][1]["components"]
        with self.assertRaisesRegex(ValueError, "components"):
            validate_review(invalid, self.inventory)

    def test_brief_direction_cannot_replace_actionable_plan(self):
        invalid = copy.deepcopy(self.review)
        del invalid["slides"][0]["issues"][0]["revision_plan"]
        with self.assertRaisesRegex(ValueError, "revision_plan"):
            validate_review(invalid, self.inventory)

    def test_content_organization_is_polish_not_a_claim_of_scientific_error(self):
        self.issue.update(category="内容叙事", scope="presentation")
        validate_review(self.review, self.inventory)
        self.issue.update(category="事实性错误", scope="presentation")
        with self.assertRaisesRegex(ValueError, "critical_logic|factual_error"):
            validate_review(self.review, self.inventory)
        self.issue["scope"] = "critical_logic"
        validate_review(self.review, self.inventory)

    def test_markdown_includes_actionable_plan_and_split_components(self):
        report = markdown(self.inventory, self.review["slides"], [])
        for detail in ("移动机制细节到第二页", "正文 24–28", "方法如何工作", "输入输出箭头"):
            self.assertIn(detail, report)

    def test_no_change_does_not_request_revisions(self):
        entry = self.review["slides"][0]
        entry.update(status="无需修改", issues=[])
        validate_review(self.review, self.inventory)
        self.assertNotIn("修改方案", markdown(self.inventory, self.review["slides"], []))

    def test_readable_image_is_not_reported_as_unknown(self):
        self.inventory["slides"][0]["objects"][0]["picture"] = {"text_or_data_inside_image": "unknown"}
        self.inventory["slides"][0]["parse_limitations"] = ["shape-1: 图片内部文字、图形或数据未知"]
        entry = self.review["slides"][0]
        entry["media_reviews"] = [{"object": "shape-1", "status": "reviewed", "source": "rendered_png",
                                  "confidence": "高", "recognized": [{"kind": "visual_content", "content": "猫照片，没有印刷文字"}],
                                  "unreadable_items": []}]
        report = markdown(self.inventory, self.review["slides"], [])
        self.assertIn("猫照片，没有印刷文字", report)
        self.assertNotIn("图片内部文字、图形或数据未知", report)

    def test_partial_inspection_names_only_unreadable_items(self):
        self.inventory["slides"][0]["objects"][0]["picture"] = {}
        self.review["slides"][0]["media_reviews"] = [
            {"object": "shape-1", "status": "partial", "source": "rendered_png_detail", "confidence": "中",
             "recognized": [{"kind": "exact_value", "content": "印刷标签 84%"}],
             "unreadable_items": ["右下角来源文字模糊"]}]
        report = markdown(self.inventory, self.review["slides"], [])
        self.assertIn("印刷标签 84%", report)
        self.assertIn("右下角来源文字模糊", report)
        self.assertNotIn("图片内部文字、图形或数据未知", report)

    def test_media_contract_requires_coverage_and_valid_object(self):
        self.inventory["slides"][0]["objects"][0]["picture"] = {}
        self.review["media_review_version"] = 1
        with self.assertRaisesRegex(ValueError, "图片.*覆盖"):
            validate_review(self.review, self.inventory)
        self.review["slides"][0]["media_reviews"] = [
            {"object": "shape-2", "status": "reviewed", "source": "rendered_png", "confidence": "高",
             "recognized": [{"kind": "visual_content", "content": "图片可辨认"}], "unreadable_items": []}]
        with self.assertRaisesRegex(ValueError, "shape-2"):
            validate_review(self.review, self.inventory)

    def test_estimates_need_a_basis_and_unseen_images_cannot_claim_recognition(self):
        self.inventory["slides"][0]["objects"][0]["picture"] = {}
        self.review["media_review_version"] = 1
        record = {"object": "shape-1", "status": "reviewed", "source": "rendered_png_detail",
                  "confidence": "中", "recognized": [{"kind": "approximate_value", "content": "约 60–65%"}],
                  "unreadable_items": []}
        self.review["slides"][0]["media_reviews"] = [record]
        with self.assertRaisesRegex(ValueError, "basis"):
            validate_review(self.review, self.inventory)
        record["recognized"][0]["basis"] = "可读纵轴为 0–100%，点位于 60 与 70 之间"
        validate_review(self.review, self.inventory)
        record.update(status="not_reviewed", reason="工具不可用")
        with self.assertRaisesRegex(ValueError, "不能声称"):
            validate_review(self.review, self.inventory)

    def test_legacy_observations_do_not_become_a_claim_of_failed_parsing(self):
        self.inventory["slides"][0]["objects"][0]["picture"] = {}
        self.inventory["slides"][0]["parse_limitations"] = ["shape-1: 图片内部文字、图形或数据未知"]
        self.review["slides"][0]["visual_observations"] = [
            {"observation": "shape-1：可辨认猫照片", "source": "rendered_png", "confidence": "高"}]
        report = markdown(self.inventory, self.review["slides"], [])
        self.assertNotIn("图片内部文字、图形或数据未知", report)
        self.assertIn("未提供结构化图片检查状态", report)


class HolisticReviewContract(unittest.TestCase):
    def setUp(self):
        PolishReviewContract.setUp(self)
        # A successful contract fixture must bind an actual image, even when
        # this test concerns another field. Render failure has its own tests.
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        image = Path(temp.name) / "slide.png"
        image.write_bytes(b"synthetic-render-binding")
        self.inventory["slides"][0]["render"] = {"status": "ok", "mapping_status": "verified", "image": str(image)}
        self.review["slides"][0]["render_sha256"] = hashlib.sha256(image.read_bytes()).hexdigest()
        self.review.update(review_schema_version=3, media_review_version=1,
                           deck_assessment={"purpose": "讲清方法", "storyline": "问题到机制再到结果",
                                            "visual_system": "先强化结论层级", "priorities": ["先改结构，再统一样式"]})
        self.entry = self.review["slides"][0]
        self.entry["assessment"] = {"core_message": "机制如何解决任务", "content_organization": "任务与机制竞争",
                                     "visual_structure": "三个同权块缺少层级", "reading_path": "先任务后机制"}
        self.entry["tip_checks"] = [{"tip": n, "result": "problem" if n == 2 else "adequate",
                                     "basis": "任务与机制竞争" if n == 2 else "该项结合本页作用检查后可接受"}
                                    for n in range(1, 12)]

    def test_v3_requires_assessment_even_for_no_change(self):
        self.entry.update(status="无需修改", issues=[])
        self.entry["tip_checks"][1]["result"] = "adequate"
        validate_review(self.review, self.inventory)
        del self.entry["assessment"]["core_message"]
        with self.assertRaisesRegex(ValueError, "core_message"):
            validate_review(self.review, self.inventory)

    def test_v3_requires_all_eleven_checks_without_duplicate_tips(self):
        validate_review(self.review, self.inventory)
        self.entry["tip_checks"][-1]["tip"] = 10
        with self.assertRaisesRegex(ValueError, "11|覆盖"):
            validate_review(self.review, self.inventory)

    def test_no_change_cannot_hide_a_confirmed_typical_problem(self):
        self.entry.update(status="无需修改", issues=[])
        with self.assertRaisesRegex(ValueError, "无需修改"):
            validate_review(self.review, self.inventory)

    def test_v3_requires_deck_assessment(self):
        del self.review["deck_assessment"]
        with self.assertRaisesRegex(ValueError, "deck_assessment"):
            validate_review(self.review, self.inventory)

    def test_report_leads_with_structure_and_keeps_measurements_in_details(self):
        self.issue["revision_plan"]["parameters"] = ["L0.92 in，仅作为待验证参数"]
        report = markdown(self.inventory, self.review["slides"], [], self.review["deck_assessment"])
        self.assertLess(report.index("问题到机制再到结果"), report.index("逐页结论"))
        self.assertIn("机制如何解决任务", report)
        self.assertIn("典型问题检查", report)
        self.assertIn("<summary>必要参数", report)
        self.assertLess(report.index("左图右文"), report.index("L0.92 in"))

    def test_visually_read_formula_limit_is_a_structural_boundary_not_a_failed_read(self):
        self.entry["structural_limitations"] = ["OMML主体未列入对象清单，已从PNG读清，原生字号未知"]
        report = markdown(self.inventory, self.review["slides"], [], self.review["deck_assessment"])
        self.assertIn("OMML主体未列入对象清单", report.split("对象解析能力边界")[1])
        self.assertNotIn("OMML主体未列入对象清单", report.split("实际失败或无法辨认项")[1].split("对象解析能力边界")[0])

    def test_conditional_split_cannot_bypass_complete_page_plans(self):
        self.issue["action"] = "revise"
        self.issue["direction"] = "先重组为一页。"
        self.issue["revision_plan"]["steps"] = ["若仍读不清，则按两条带拆分为两页。"]
        del self.issue["split_plan"]
        with self.assertRaisesRegex(ValueError, "split_plan"):
            validate_review(self.review, self.inventory)

    def test_conditional_split_is_allowed_with_complete_alternative(self):
        self.issue.update(action="revise", direction="若重排仍不足，则拆成两页。")
        validate_review(self.review, self.inventory)
        report = markdown(self.inventory, self.review["slides"], [], self.review["deck_assessment"])
        self.assertIn("方法如何工作", report)
        del self.issue["split_plan"][1]["components"]
        with self.assertRaisesRegex(ValueError, "components"):
            validate_review(self.review, self.inventory)

    def test_parameters_at_issue_level_are_rejected_in_v3(self):
        self.issue["parameters"] = ["现有宽度，仅供定位"]
        with self.assertRaisesRegex(ValueError, "revision_plan.parameters"):
            validate_review(self.review, self.inventory)

    def test_existing_two_page_comparison_does_not_force_a_split(self):
        self.issue["action"] = "revise"
        del self.issue["split_plan"]
        self.issue["revision_plan"]["steps"] = ["保留现有一页，调整两栏分组。"]
        for direction in ("不拆页；两页对照使用同样的组件。", "无需再拆分为两页。",
                          "Do not split into two slides; retain the overview."):
            with self.subTest(direction=direction):
                self.issue["direction"] = direction
                validate_review(self.review, self.inventory)


class EvidenceReviewContract(unittest.TestCase):
    def setUp(self):
        HolisticReviewContract.setUp(self)
        self.review["evidence_review_version"] = 1
        self.inventory["slides"][0]["objects"][0]["text_frame"] = {"text": "任务定义\n输入与输出"}
        self.issue["target_evidence"] = {"source": "pptx_object", "quote": "任务定义"}

    def test_existing_object_cannot_anchor_unrelated_content(self):
        validate_review(self.review, self.inventory)
        self.issue["target_evidence"]["quote"] = "正文求和公式"
        with self.assertRaisesRegex(ValueError, "引用.*目标对象"):
            validate_review(self.review, self.inventory)

    def test_enabled_evidence_contract_requires_an_anchor(self):
        del self.issue["target_evidence"]
        with self.assertRaisesRegex(ValueError, "target_evidence"):
            validate_review(self.review, self.inventory)

    def test_missing_body_uses_region_instead_of_header_object(self):
        self.issue["object"] = "region:页面中部求和公式"
        self.issue["target_evidence"] = {"source": "raw_xml+rendered_png",
                                          "description": "输入向量中的部分，XML id=3，清单未读取正文"}
        validate_review(self.review, self.inventory)
        self.issue["object"] = "shape-1"
        with self.assertRaisesRegex(ValueError, "quote"):
            validate_review(self.review, self.inventory)

    def test_visual_object_uses_a_description_without_invented_text(self):
        self.inventory["slides"][0]["objects"][0].pop("text_frame")
        self.issue["target_evidence"] = {"source": "rendered_png", "description": "图内上方训练带"}
        validate_review(self.review, self.inventory)

    def test_recomputed_sum_rejects_wrong_revision_example(self):
        self.issue["numeric_checks"] = [{"operation": "geometric_sum", "basis": "修改后下标的合法特例",
                                         "inputs": {"coefficient": 1, "ratio": 0.5, "start": 0, "end": 1},
                                         "reported": 0.75}]
        with self.assertRaisesRegex(ValueError, "复算.*不一致"):
            validate_review(self.review, self.inventory)
        self.issue["numeric_checks"][0]["reported"] = 1.5
        validate_review(self.review, self.inventory)
        self.assertIn("数值复算", markdown(self.inventory, self.review["slides"], []))

    def test_gain_superlative_keeps_its_comparison_scope(self):
        self.issue["numeric_checks"] = [{"operation": "compare_gains", "basis": "同一模型相对初始行的提升",
                                         "inputs": {"rows": [
                                             {"label": "MMLU-Pro", "before": 42.3, "after": 49.2},
                                             {"label": "GPQA", "before": 37.1, "after": 38.4}]},
                                         "reported": ["GPQA"]}]
        with self.assertRaisesRegex(ValueError, "复算.*不一致"):
            validate_review(self.review, self.inventory)
        self.issue["numeric_checks"][0]["reported"] = ["MMLU-Pro"]
        validate_review(self.review, self.inventory)

    def test_evidence_version_is_not_silently_ignored(self):
        self.review["evidence_review_version"] = 2
        with self.assertRaisesRegex(ValueError, "evidence_review_version"):
            validate_review(self.review, self.inventory)


@unittest.skipUnless(DECK.is_file() and (ROOT / "audit-v2/inventory.json").is_file() and (ROOT / "validation/v2_review.json").is_file(), "local v2 acceptance PPTX or bound inventory/review is absent")
class RealDeckAcceptance(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.inventory = extract(DECK)
        cls.slides = cls.inventory["slides"]

    def test_all_slides_and_object_kinds(self):
        self.assertEqual(22, self.inventory["slide_count"])
        counts = Counter()
        for slide in self.slides:
            for obj in flatten(slide["objects"]):
                for kind in ("table", "chart", "picture"):
                    if kind in obj:
                        counts[kind] += 1
        self.assertEqual({"table": 6, "chart": 3, "picture": 19}, dict(counts))

    def test_table_cells_and_chart_values_are_bound_to_correct_slides(self):
        table = next(obj["table"] for obj in self.slides[17]["objects"] if "table" in obj)
        self.assertEqual((7, 4), (table["row_count"], table["column_count"]))
        self.assertTrue(any(cell["text"] == "237.9" for row in table["rows"] for cell in row))
        charts = [obj["chart"] for obj in self.slides[19]["objects"] if "chart" in obj]
        self.assertEqual(2, len(charts))
        self.assertEqual([51.5, 65.9], charts[0]["series"][0]["values"])
        self.assertEqual([253.0, 263.3], charts[1]["series"][1]["values"])
        closing_chart = next(obj["chart"] for obj in self.slides[21]["objects"] if "chart" in obj)
        self.assertEqual([84.0, 71.0, 96.0], closing_chart["series"][0]["values"])

    def test_visibility_and_review_binding(self):
        self.assertIn("机器翻译", "".join(self.slides[1]["visible_text"]))
        self.assertTrue(all(slide["render"]["status"] == "ok"
                            for slide in json.loads((ROOT / "audit-v2/inventory.json").read_text(encoding="utf-8"))["slides"]))
        review = json.loads((ROOT / "validation/v2_review.json").read_text(encoding="utf-8"))
        rendered_inventory = json.loads((ROOT / "audit-v2/inventory.json").read_text(encoding="utf-8"))
        validate_review(review, rendered_inventory)
        self.assertEqual(2, review["review_schema_version"])
        self.assertEqual(9, sum(x["status"] == "无需修改" for x in review["slides"]))
        self.assertEqual("无需修改", review["slides"][20]["status"])
        self.assertEqual("有确定问题", review["slides"][5]["status"])
        self.assertEqual(9, min(x["size_pt"] for x in self.slides[5]["typography"]["samples"]
                                if x["size_pt"] is not None))
        self.assertTrue(any(slide["status"] == "无需修改" for slide in review["slides"]))
        self.assertTrue(any(slide["status"] == "有确定问题" for slide in review["slides"]))


if __name__ == "__main__":
    unittest.main()

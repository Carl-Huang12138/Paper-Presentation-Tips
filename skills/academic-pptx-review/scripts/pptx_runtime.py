"""Static PPTX evidence rendering, page identity, and immutable review snapshots.

This module prepares evidence. It never claims a person or vision model has
inspected a PNG, and it never edits the supplied presentation.
"""

from __future__ import annotations

from datetime import datetime, timezone
from contextlib import contextmanager
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import unicodedata
import uuid

import fitz
from pptx import Presentation
from pptx.util import Inches

RUNTIME_VERSION = "4.1.0"
MANIFEST_VERSION = 1
_CAPABILITY_CACHE: dict[tuple, dict] = {}


@contextmanager
def _render_work_directory():
    """Short disposable profile paths, independent of the chosen output depth."""
    if os.name != 'nt':
        with tempfile.TemporaryDirectory(prefix='pptx-render-', ignore_cleanup_errors=True) as name:
            yield Path(name)
        return
    parent = Path(tempfile.gettempdir()).resolve()
    work = parent / ('pptx-rw-' + uuid.uuid4().hex[:12])
    # Windows Python 3.13 mode 700 may deny sandbox access to its own files;
    # normal mkdir inherits the user's TEMP ACL instead.
    work.mkdir()
    try:
        yield work
    finally:
        if not work.is_symlink() and work.resolve().parent == parent:
            shutil.rmtree(work, ignore_errors=True)


class RenderError(RuntimeError):
    """The configured renderer cannot establish a trustworthy page mapping."""


class SnapshotError(ValueError):
    """Evidence differs from the immutable snapshot or is incompletely bound."""


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _sha256(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def _write_json(path: Path, data: dict) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _subprocess_run(command: list[str], *, timeout: int = 120) -> subprocess.CompletedProcess:
    # soffice.com gives useful diagnostics on Windows without a visible console.
    kwargs = {"capture_output": True, "text": True, "errors": "replace", "timeout": timeout}
    if os.name == "nt":
        kwargs["creationflags"] = subprocess.CREATE_NO_WINDOW
    return subprocess.run(command, **kwargs)


def resolve_soffice(soffice: str | Path | None = None) -> Path:
    """Resolve explicit option, PPTX_SOFFICE, PATH, then common install paths."""
    configured = soffice if soffice is not None else os.environ.get("PPTX_SOFFICE") or None
    if configured is not None:
        candidate = str(configured).strip().strip('"')
        path = Path(candidate).expanduser()
        if not path.is_file():
            found = shutil.which(candidate)
            path = Path(found) if found else path
        if not path.is_file():
            raise RenderError(f"Configured soffice executable does not exist: {candidate}")
        return path.resolve()
    for name in ("soffice", "libreoffice", "soffice.com", "soffice.exe"):
        if found := shutil.which(name):
            return Path(found).resolve()
    common = [
        Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "LibreOffice/program/soffice.com",
        Path(os.environ.get("ProgramFiles(x86)", "C:/Program Files (x86)")) / "LibreOffice/program/soffice.com",
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
        Path("/usr/bin/soffice"), Path("/usr/bin/libreoffice"),
        Path("/opt/libreoffice/program/soffice"),
    ]
    for candidate in common:
        if candidate.is_file():
            return candidate.resolve()
    raise RenderError("No soffice executable found; use --soffice or set PPTX_SOFFICE")


def _renderer_preflight(executable: Path) -> dict:
    run = _subprocess_run([str(executable), "--version"], timeout=30)
    version = (run.stdout + run.stderr).strip()
    if run.returncode or "LibreOffice" not in version:
        raise RenderError(f"soffice --version failed ({run.returncode}): {version or 'no version output'}")
    return {"name": "LibreOffice", "executable": str(executable), "version": version,
            "runtime_version": RUNTIME_VERSION, "pymupdf_version": fitz.VersionBind}


def _convert_to_pdf(executable: Path, sources: list[Path], destination: Path,
                    profile: Path, *, include_hidden: bool) -> list[Path]:
    destination.mkdir(parents=True, exist_ok=True)
    options = json.dumps({"ExportHiddenSlides": {"type": "boolean", "value": include_hidden}},
                         separators=(",", ":"))
    command = [str(executable), f"-env:UserInstallation={profile.resolve().as_uri()}",
               "--headless", "--convert-to", "pdf:impress_pdf_Export:" + options,
               "--outdir", str(destination.resolve()), *(str(path.resolve()) for path in sources)]
    run = _subprocess_run(command)
    if run.returncode:
        raise RenderError(f"soffice PDF export failed ({run.returncode}): {(run.stdout + run.stderr).strip()}")
    pdfs = [destination / (source.stem + ".pdf") for source in sources]
    missing = [str(path) for path in pdfs if not path.is_file()]
    if missing:
        raise RenderError("soffice did not create PDF: " + ", ".join(missing) +
                          "; " + (run.stdout + run.stderr).strip())
    return pdfs


def probe_hidden_export(executable: Path, renderer: dict, work_dir: Path) -> dict:
    """Actually exercise every hidden-page boundary, including an all-hidden deck.

    The cached result is tied to this executable's version, size and mtime, and
    exists only for the running Python process. Counts alone are insufficient:
    each PDF page must contain its unique original-page marker in export order.
    """
    stat = executable.stat()
    key = (str(executable), renderer["version"], stat.st_size, stat.st_mtime_ns)
    if key in _CAPABILITY_CACHE:
        return json.loads(json.dumps(_CAPABILITY_CACHE[key]))
    cases = {"first": (True, False, False), "middle": (False, True, False),
             "last": (False, False, True), "consecutive": (False, True, True, False),
             "all_hidden": (True, True, True)}
    report = {"tested_at_utc": _utc_now(), "method": "unique_page_markers",
              "export_hidden_slides": False, "exclude_hidden_slides": False, "cases": []}
    try:
        work_dir.mkdir(parents=True, exist_ok=True)
        sources = []
        for name, hidden in cases.items():
            pres = Presentation()
            for index, is_hidden in enumerate(hidden, 1):
                slide = pres.slides.add_slide(pres.slide_layouts[6])
                box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
                box.text = f"PPTXPROBE{name.upper()}PAGE{index}"
                if is_hidden:
                    slide._element.set("show", "0")
            source = work_dir / (name + ".pptx")
            pres.save(source)
            sources.append(source)
        all_pdfs = _convert_to_pdf(executable, sources, work_dir / "all", work_dir / "profile-all",
                                   include_hidden=True)
        # An all-hidden slideshow has no included pages; never infer a mapping
        # from a renderer's possible placeholder page in that case.
        shown_sources = [source for source in sources if source.stem != "all_hidden"]
        shown_pdfs = _convert_to_pdf(executable, shown_sources, work_dir / "slides",
                                     work_dir / "profile-slides", include_hidden=False)
        shown_lookup = {pdf.stem: pdf for pdf in shown_pdfs}
        for source, pdf in zip(sources, all_pdfs):
            hidden = cases[source.stem]
            expected = [f"PPTXPROBE{source.stem.upper()}PAGE{i}" for i in range(1, len(hidden) + 1)]
            with fitz.open(pdf) as document:
                actual = [re.sub(r"\s+", "", page.get_text()) for page in document]
            included = [marker for marker, is_hidden in zip(expected, hidden) if not is_hidden]
            shown_actual = []
            if source.stem in shown_lookup:
                with fitz.open(shown_lookup[source.stem]) as document:
                    shown_actual = [re.sub(r"\s+", "", page.get_text()) for page in document]
            report["cases"].append({"case": source.stem, "hidden": list(hidden),
                                    "all_expected": expected, "all_actual": actual,
                                    "all_verified": actual == expected,
                                    "slides_expected": included, "slides_actual": shown_actual,
                                    "slides_verified": shown_actual == included,
                                    "source_sha256": _sha256(source), "all_pdf_sha256": _sha256(pdf)})
        report["export_hidden_slides"] = all(case["all_verified"] for case in report["cases"])
        report["exclude_hidden_slides"] = all(case["slides_verified"] for case in report["cases"])
    except Exception as error:
        report["error"] = f"Hidden-slide export capability was not verified: {error}"
    if 'error' not in report:
        _CAPABILITY_CACHE[key] = report
    return json.loads(json.dumps(report))


def _canonical_slides(source: Path, slides: list[dict]) -> list[dict]:
    presentation = Presentation(str(source))
    if len(presentation.slides) != len(slides):
        raise RenderError("PPTX source and object inventory have different slide counts")
    result = []
    for index, (page, supplied) in enumerate(zip(presentation.slides, slides), 1):
        identity = {"original_slide": index, "slide_id": page.slide_id,
                    "slide_part": str(page.part.partname),
                    "hidden_in_slideshow": page._element.get("show") in ("0", "false")}
        if supplied.get("slide") != index:
            raise RenderError(f"Inventory original slide order is invalid at page {index}")
        for key in ("slide_id", "slide_part", "hidden_in_slideshow"):
            if key in supplied and supplied[key] != identity[key]:
                raise RenderError(f"Inventory {key} does not match PPTX source at page {index}")
        supplied.update({key: identity[key] for key in ("slide_id", "slide_part", "hidden_in_slideshow")})
        result.append(identity)
    return result


def _visible_objects(items):
    for item in items:
        if not item.get("effective_visible_in_slideshow", item.get("visible_in_slideshow", True)):
            continue
        yield item
        yield from _visible_objects(item.get("children", []))


def _text_units(slide: dict) -> list[dict]:
    units = []
    for obj in _visible_objects(slide.get("objects", [])):
        path = obj.get("path", "unknown")
        frame = obj.get("text_frame", {})
        paragraphs = frame.get("paragraphs")
        if paragraphs:
            units.extend({"source": f"{path}:paragraph-{i}", "text": paragraph.get("text", "")}
                         for i, paragraph in enumerate(paragraphs, 1))
        elif frame.get("text"):
            units.append({"source": path, "text": frame["text"]})
        for row_index, row in enumerate(obj.get("table", {}).get("rows", []), 1):
            for column_index, cell in enumerate(row, 1):
                units.append({"source": f"{path}:cell-{row_index}-{column_index}", "text": cell.get("text", "")})
        chart = obj.get("chart", {})
        if chart.get("title") and chart["title"] != "unknown":
            units.append({"source": f"{path}:chart-title", "text": chart["title"]})
        if chart.get("legend"):
            units.extend({"source": f"{path}:series-{i}", "text": series.get("name", "")}
                         for i, series in enumerate(chart.get("series", []), 1)
                         if series.get("name") != "unknown")
    return [unit for unit in units if isinstance(unit["text"], str) and unit["text"].strip()]


def compare_text(slide: dict, pdf_text: str) -> dict:
    """Compare both directions; this is an extraction cue, never a quality verdict."""
    compact = lambda value: re.sub(r"\s+", "", unicodedata.normalize("NFKC", value))
    units = _text_units(slide)
    remaining = compact(pdf_text)
    missing = []
    for unit in units:
        text = compact(unit["text"])
        index = remaining.find(text)
        if index < 0:
            missing.append(unit)
        else:
            remaining = remaining[:index] + remaining[index + len(text):]
    unmatched_pdf = [remaining] if remaining else []
    status = "not_testable" if not units and not compact(pdf_text) else (
        "partially_compared" if missing or unmatched_pdf or not units else "matched")
    return {"comparison_status": status, "compared_object_units": len(units),
            "object_text_absent_in_pdf": missing,
            "structure_tokens_absent_in_pdf": [item["text"] for item in missing],
            "pdf_text_unattributed": unmatched_pdf,
            "comparison_limitations": "PDF text may omit formulas and image text, or include inherited content; inspect the PNG."}


def render(pptx_path: Path, out_dir: Path, slides: list[dict], *,
           soffice: str | Path | None = None, dpi: int = 144,
           keep_pdf: bool = True, render_scope: str = "all") -> list[str]:
    """Render static pages with an explicit, verified original-slide mapping."""
    source, root = Path(pptx_path).resolve(), Path(out_dir).resolve()
    if (root / "snapshot_manifest.json").exists():
        raise SnapshotError("A published snapshot is immutable; render into a new staging directory")
    root.mkdir(parents=True, exist_ok=True)
    issues = []
    manifest = {"schema_version": MANIFEST_VERSION, "runtime_version": RUNTIME_VERSION,
                "created_at_utc": _utc_now(), "render_scope": render_scope, "dpi": dpi,
                "renderer": None, "capabilities": None, "mapping_status": "unverified",
                "pdf": None, "slides": []}
    for slide in slides:
        slide["render"] = {"status": "failed", "mapping_status": "unverified",
                           "reason": "Render did not complete", "render_scope": render_scope}
    try:
        if render_scope not in ("all", "slides"):
            raise RenderError("render_scope must be all or slides")
        if type(dpi) is not int or dpi <= 0:
            raise RenderError("Render DPI must be a positive integer")
        before_hash = _sha256(source)
        manifest["input_sha256"] = before_hash
        identities = _canonical_slides(source, slides)
        included = [identity for identity in identities
                    if render_scope == "all" or not identity["hidden_in_slideshow"]]
        for identity, slide in zip(identities, slides):
            entry = {**identity, "render_scope": render_scope, "pdf_page": None,
                     "image": None, "sha256": None, "status": "pending", "mapping_status": "unverified"}
            if identity not in included:
                entry.update(status="excluded", mapping_status="excluded",
                             reason="Hidden original slide excluded by slides scope")
                slide["render"] = {key: value for key, value in entry.items() if value is not None}
            manifest["slides"].append(entry)
        if not included:
            manifest["mapping_status"] = "verified"
            manifest["mapping_method"] = "no_included_original_slides"
            manifest["status"] = "complete"
            _write_json(root / "render_manifest.json", manifest)
            return issues
        executable = resolve_soffice(soffice)
        renderer = _renderer_preflight(executable)
        manifest["renderer"] = renderer
        # Renderer profile caches are disposable and can be briefly held by
        # Windows after successful export. Keep them outside the snapshot and
        # do not turn optional cache cleanup into a false evidence failure.
        with _render_work_directory() as work:
            hidden = any(identity["hidden_in_slideshow"] for identity in identities)
            capabilities = probe_hidden_export(executable, renderer, work / "probes") if hidden else {
                "export_hidden_slides": None, "exclude_hidden_slides": None,
                "method": "not_needed_no_hidden_original_slides", "cases": []}
            manifest["capabilities"] = capabilities
            pdf = _convert_to_pdf(executable, [source], work / "pdf", work / "profile",
                                  include_hidden=render_scope == "all")[0]
            pdf_hash = _sha256(pdf)
            retained_pdf = root / "renders" / "presentation.pdf"
            retained_pdf.parent.mkdir(parents=True, exist_ok=True)
            if keep_pdf:
                shutil.copyfile(pdf, retained_pdf)
            with fitz.open(pdf) as document:
                manifest["pdf"] = {"path": "renders/presentation.pdf" if keep_pdf else None,
                                   "sha256": pdf_hash, "page_count": len(document), "retained": keep_pdf}
                count_ok = len(document) == len(included)
                capability_key = "export_hidden_slides" if render_scope == "all" else "exclude_hidden_slides"
                capability_ok = not hidden or capabilities.get(capability_key) is True
                unchanged = before_hash == _sha256(source)
                if not (count_ok and capability_ok and unchanged):
                    reason = (f"Page mapping unverified: PDF has {len(document)} pages for {len(included)} included originals; "
                              f"hidden export capability={capability_ok}; input unchanged={unchanged}")
                    issues.append(reason)
                    for entry, slide in zip(manifest["slides"], slides):
                        if entry["status"] == "excluded":
                            continue
                        entry.update(status="mapping_unverified", mapping_status="unverified", reason=reason)
                        slide["render"] = {key: value for key, value in entry.items() if value is not None}
                    manifest["status"] = "failed"
                else:
                    manifest["mapping_status"] = "verified"
                    manifest["mapping_method"] = "source_slide_order+verified_hidden_filter+pdf_page_count" if hidden else "source_slide_order+pdf_page_count"
                    by_original = {entry["original_slide"]: entry for entry in manifest["slides"]}
                    for pdf_index, identity in enumerate(included):
                        number = identity["original_slide"]
                        page = document[pdf_index]
                        image_relative = f"renders/slide-{number:02d}.png"
                        image_path = root / image_relative
                        page.get_pixmap(dpi=dpi, alpha=False).save(image_path)
                        pdf_text = page.get_text("text")
                        entry = by_original[number]
                        entry.update(status="ok", mapping_status="verified", pdf_page=pdf_index + 1,
                                     image=image_relative, sha256=_sha256(image_path))
                        slides[number - 1]["render"] = {**entry, "pdf_text": pdf_text,
                                                       **compare_text(slides[number - 1], pdf_text),
                                                       "source": "libreoffice_render_pdf", "dpi": dpi,
                                                       "render_sha256": entry["sha256"]}
                        _picture_clips(page, root, slides[number - 1], dpi=max(288, dpi))
                    manifest["status"] = "complete"
    except Exception as error:
        reason = f"Render failed: {error}"
        issues.append(reason)
        manifest["status"] = "failed"
        manifest["error"] = reason
        for slide in slides:
            if slide["render"]["status"] != "excluded":
                slide["render"] = {"status": "failed", "mapping_status": "unverified",
                                   "reason": reason, "render_scope": render_scope}
        for entry in manifest["slides"]:
            if entry["status"] != "excluded":
                entry.update(status="failed", mapping_status="unverified", reason=reason,
                             image=None, sha256=None, pdf_page=None)
    _write_json(root / "render_manifest.json", manifest)
    return issues


def _all_objects(items):
    for item in items:
        yield item
        yield from _all_objects(item.get("children", []))


def _picture_clips(page, root: Path, slide: dict, *, dpi: int) -> None:
    """Use the rendered page, including foreground paint, rather than raw media.

    Axis-aligned projected bounds are intentional: rotated pictures and groups
    remain in their page context. A clip is supplementary evidence, not OCR or
    proof that every pixel inside it belongs to this object.
    """
    for obj in _all_objects(slide.get("objects", [])):
        if "picture" not in obj or not obj.get("visible_in_slideshow", True):
            continue
        picture = obj["picture"]
        bounds = obj.get("slide_bounds")
        if not bounds:
            picture["render_crop"] = {"status": "unknown_geometry", "reason": "No reliable projected page bounds"}
            continue
        try:
            x, y, width, height = [float(bounds[key]) * 72 for key in
                                  ("left_in", "top_in", "width_in", "height_in")]
            if not all(math.isfinite(value) for value in (x, y, width, height)) or width <= 0 or height <= 0:
                raise ValueError("Invalid picture page bounds")
            clip = fitz.Rect(x, y, x + width, y + height) & page.rect
            if clip.is_empty:
                picture["render_crop"] = {"status": "outside_page", "reason": "Projected bounds lie outside rendered page"}
                continue
            safe_name = re.sub(r"[^A-Za-z0-9_-]", "_", obj["path"])
            relative = f"detail/slide-{slide['slide']:02d}-{safe_name}.png"
            target = root / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            page.get_pixmap(dpi=dpi, clip=clip, alpha=False).save(target)
            picture["render_crop"] = {"status": "ok", "image": relative, "sha256": _sha256(target),
                "source": "rendered_page_clip", "dpi": dpi, "clip_pdf_points": list(clip),
                "geometry_basis": "projected_slide_bounds", "includes_foreground_and_neighbours": True}
        except (ValueError, TypeError, KeyError, RuntimeError) as error:
            picture["render_crop"] = {"status": "failed", "reason": str(error)}


def _relative_file(root: Path, value: str | Path, *, must_exist: bool = True) -> str:
    if not isinstance(value, (str, Path)) or not str(value).strip():
        raise SnapshotError("Evidence paths must be nonempty file paths")
    path = Path(value)
    absolute = (path if path.is_absolute() else root / path).resolve()
    if not absolute.is_relative_to(root) or absolute == root:
        raise SnapshotError(f"Evidence path escapes snapshot root: {value}")
    if must_exist and not absolute.is_file():
        raise SnapshotError(f"Evidence file is missing: {value}")
    return absolute.relative_to(root).as_posix()


def _normalize_inventory_paths(root: Path, inventory: dict) -> None:
    for media in inventory.get("media_manifest", []):
        if media.get("path"):
            media["path"] = _relative_file(root, media["path"])
    for slide in inventory["slides"]:
        rendered = slide.get("render", {})
        if rendered.get("image"):
            rendered["image"] = _relative_file(root, rendered["image"])
            image_hash = _sha256(root / rendered["image"])
            for key in ("sha256", "render_sha256"):
                if rendered.get(key) and rendered[key] != image_hash:
                    raise SnapshotError(f"Original slide {slide['slide']} PNG hash differs from render metadata")
                rendered[key] = image_hash
        for obj in _all_objects(slide.get("objects", [])):
            picture = obj.get("picture", {})
            for key in ("original_media", "media_path", "export_path"):
                if picture.get(key):
                    picture[key] = _relative_file(root, picture[key])
            original = picture.get("original_media") or picture.get("media_path") or picture.get("export_path")
            if original:
                original_hash = _sha256(root / original)
                if picture.get("sha256") and picture["sha256"] != original_hash:
                    raise SnapshotError(f"Original media hash differs at slide {slide['slide']} {obj.get('path')}")
                picture["sha256"] = original_hash
            crop = picture.get("render_crop")
            if isinstance(crop, dict) and crop.get("image"):
                crop["image"] = _relative_file(root, crop["image"])
                crop["sha256"] = _sha256(root / crop["image"])


def _media_manifest(inventory: dict) -> dict:
    instances = []
    for slide in inventory["slides"]:
        for obj in _all_objects(slide.get("objects", [])):
            if "picture" not in obj:
                continue
            picture = obj["picture"]
            original = picture.get("original_media") or picture.get("media_path") or picture.get("export_path")
            instances.append({"slide": slide["slide"], "object": obj.get("path"),
                              "instance_id": picture.get("instance_id", f"slide-{slide['slide']}:{obj.get('path')}"),
                              "original_media": original, "sha256": picture.get("sha256"),
                              "export_status": "exported" if original else "not_exported",
                              "media_part": picture.get("media_part"),
                              "relationship_id": picture.get("relationship_id"),
                              "pixel_size": picture.get("pixel_size"), "crop": picture.get("crop"),
                              "rotation_degrees": picture.get("rotation_degrees"),
                              "slide_bounds": obj.get("slide_bounds", obj.get("bounds")),
                              "visible_in_slideshow": obj.get("effective_visible_in_slideshow", obj.get("visible_in_slideshow", True)),
                              "render_crop": picture.get("render_crop"),
                              "limitation": "Original media can include cropped or occluded content; inspect the full slide to establish visibility."})
    return {"schema_version": MANIFEST_VERSION, "instances": instances}


def _fallback_render_manifest(inventory: dict) -> dict:
    entries = []
    for slide in inventory["slides"]:
        rendered = slide.get("render", {})
        entries.append({"original_slide": slide["slide"], "slide_id": slide.get("slide_id"),
                        "slide_part": slide.get("slide_part"),
                        "hidden_in_slideshow": slide.get("hidden_in_slideshow", False),
                        "status": rendered.get("status", "failed"),
                        "mapping_status": rendered.get("mapping_status", "unverified"),
                        "image": rendered.get("image"), "sha256": rendered.get("sha256"),
                        "pdf_page": rendered.get("pdf_page"), "reason": rendered.get("reason")})
    return {"schema_version": MANIFEST_VERSION, "runtime_version": RUNTIME_VERSION,
            "input_sha256": inventory["sha256"], "render_scope": "all", "renderer": None,
            "capabilities": None, "pdf": None, "mapping_status": "unverified",
            "status": "partial", "slides": entries,
            "limitation": "No renderer manifest was supplied; missing page mapping remains unverified."}


def _referenced_files(root: Path, inventory: dict, render_manifest: dict, media_manifest: dict) -> set[str]:
    paths = {"inventory.json", "render_manifest.json", "media_manifest.json"}
    for slide in inventory["slides"]:
        if image := slide.get("render", {}).get("image"):
            paths.add(_relative_file(root, image))
    for entry in render_manifest.get("slides", []):
        if image := entry.get("image"):
            paths.add(_relative_file(root, image))
    if pdf := (render_manifest.get("pdf") or {}).get("path"):
        paths.add(_relative_file(root, pdf))
    for instance in media_manifest.get("instances", []):
        if original := instance.get("original_media"):
            paths.add(_relative_file(root, original))
        if image := (instance.get("render_crop") or {}).get("image"):
            paths.add(_relative_file(root, image))
    return paths


def create_snapshot(out_dir: Path, inventory: dict, *, run_id: str | None = None,
                    parser_version: str | None = None) -> dict:
    """Freeze the prepared input/object/render/media evidence once, before review."""
    root = Path(out_dir).resolve()
    root.mkdir(parents=True, exist_ok=True)
    if (root / "snapshot_manifest.json").exists():
        raise SnapshotError("Snapshot already exists; verify it or create a new staging directory")
    source = Path(inventory["input"])
    if _sha256(source) != inventory["sha256"]:
        raise SnapshotError("Input PPTX changed after object extraction")
    _canonical_slides(source, inventory["slides"])
    _normalize_inventory_paths(root, inventory)
    snapshot_id = inventory.get("snapshot_id") or str(uuid.uuid4())
    run_id = run_id or inventory.get("run_id") or str(uuid.uuid4())
    if not isinstance(snapshot_id, str) or not snapshot_id.strip() or not isinstance(run_id, str) or not run_id.strip():
        raise SnapshotError("run_id and snapshot_id must be nonempty strings")
    inventory.update(run_id=run_id, snapshot_id=snapshot_id, render_fidelity_review_version=1,
                     snapshot={"id": snapshot_id, "snapshot_id": snapshot_id, "run_id": run_id,
                               "manifest": "snapshot_manifest.json"})
    renderer_path = root / "render_manifest.json"
    if renderer_path.is_file():
        render_manifest = json.loads(renderer_path.read_text(encoding="utf-8"))
    else:
        render_manifest = _fallback_render_manifest(inventory)
    if render_manifest.get("input_sha256") not in (None, inventory["sha256"]):
        raise SnapshotError("Renderer manifest belongs to another input PPTX")
    render_manifest["input_sha256"] = inventory["sha256"]
    render_manifest.update(run_id=run_id, snapshot_id=snapshot_id)
    for entry in render_manifest.get("slides", []):
        if entry.get("image"):
            entry["image"] = _relative_file(root, entry["image"])
            entry["sha256"] = _sha256(root / entry["image"])
    if (render_manifest.get("pdf") or {}).get("path"):
        render_manifest["pdf"]["path"] = _relative_file(root, render_manifest["pdf"]["path"])
    media_manifest = _media_manifest(inventory)
    media_manifest.update(run_id=run_id, snapshot_id=snapshot_id)
    _write_json(renderer_path, render_manifest)
    _write_json(root / "media_manifest.json", media_manifest)
    _write_json(root / "inventory.json", inventory)
    paths = _referenced_files(root, inventory, render_manifest, media_manifest)
    for name in ("renders", "media"):
        directory = root / name
        if directory.exists():
            paths.update(_relative_file(root, path) for path in directory.rglob("*") if path.is_file())
    immutable_files = [{"path": path, "sha256": _sha256(root / path), "size": (root / path).stat().st_size}
                       for path in sorted(paths)]
    code_root = Path(__file__).resolve().parent
    parser_sources = {name: _sha256(code_root / name) for name in ("analyze_pptx.py", "pptx_reader.py")
                      if (code_root / name).is_file()}
    snapshot = {"schema_version": MANIFEST_VERSION, "id": snapshot_id, "snapshot_id": snapshot_id,
                "run_id": run_id, "created_at_utc": _utc_now(), "input_sha256": inventory["sha256"],
                "input": {"path": inventory["input"], "sha256": inventory["sha256"],
                          "slide_count": inventory["slide_count"]},
                "parser": {"version": parser_version or inventory.get("parser_version", "unversioned"),
                           "source_hashes": parser_sources},
                "runtime": {"version": RUNTIME_VERSION, "sha256": _sha256(Path(__file__))},
                "renderer": render_manifest.get("renderer"),
                "render_scope": render_manifest.get("render_scope", "all"),
                "immutable_files": immutable_files,
                "preparation_only": True,
                "limitation": "Hash and mapping checks prepare evidence; they do not establish visual or semantic review."}
    _write_json(root / "snapshot_manifest.json", snapshot)
    verify_snapshot(root, inventory["sha256"])
    return snapshot


def verify_snapshot(out_dir: Path, input_sha: str | None = None) -> dict:
    """Read and verify an existing snapshot; never rerender or mutate evidence."""
    root = Path(out_dir).resolve()
    try:
        snapshot = json.loads((root / "snapshot_manifest.json").read_text(encoding="utf-8"))
        if snapshot.get("schema_version") != MANIFEST_VERSION:
            raise SnapshotError("Unsupported snapshot manifest version")
        if not snapshot.get("snapshot_id") or snapshot.get("id") != snapshot["snapshot_id"] or not snapshot.get("run_id"):
            raise SnapshotError("Snapshot identity is missing or inconsistent")
        source_hash = snapshot.get("input_sha256")
        if source_hash != snapshot.get("input", {}).get("sha256") or not re.fullmatch(r"[0-9a-f]{64}", source_hash or ""):
            raise SnapshotError("Snapshot input hash is missing or inconsistent")
        if input_sha is not None and source_hash != input_sha:
            raise SnapshotError("Snapshot belongs to a different input PPTX")
        files = snapshot.get("immutable_files")
        if not isinstance(files, list) or not files:
            raise SnapshotError("Snapshot immutable file manifest is empty")
        bound = set()
        for entry in files:
            if not isinstance(entry, dict):
                raise SnapshotError("Invalid immutable file entry")
            path = _relative_file(root, entry.get("path"))
            if path in bound or path == "snapshot_manifest.json":
                raise SnapshotError("Duplicate or circular immutable file binding")
            bound.add(path)
            if (root / path).stat().st_size != entry.get("size") or _sha256(root / path) != entry.get("sha256"):
                raise SnapshotError(f"Snapshot evidence hash mismatch: {path}")
        inventory = json.loads((root / "inventory.json").read_text(encoding="utf-8"))
        renderer = json.loads((root / "render_manifest.json").read_text(encoding="utf-8"))
        media = json.loads((root / "media_manifest.json").read_text(encoding="utf-8"))
        for data in (inventory, renderer, media):
            if data.get("snapshot_id") != snapshot["snapshot_id"] or data.get("run_id") != snapshot["run_id"]:
                raise SnapshotError("Evidence files belong to different snapshot/run identities")
        if inventory.get("sha256") != source_hash or renderer.get("input_sha256") != source_hash:
            raise SnapshotError("Evidence files belong to different input hashes")
        if inventory.get("slide_count") != snapshot.get("input", {}).get("slide_count"):
            raise SnapshotError("Snapshot slide count is inconsistent")
        expected = _referenced_files(root, inventory, renderer, media)
        if not expected.issubset(bound):
            raise SnapshotError("Snapshot leaves evidence files unbound: " + ", ".join(sorted(expected - bound)))
        entries = renderer.get("slides", [])
        if len(entries) != len(inventory["slides"]):
            raise SnapshotError("Renderer mapping does not cover every original slide")
        render_complete = True
        for index, (slide, entry) in enumerate(zip(inventory["slides"], entries), 1):
            if slide.get("slide") != index or entry.get("original_slide") != index:
                raise SnapshotError("Renderer mapping original slide order is inconsistent")
            for key in ("slide_id", "slide_part", "hidden_in_slideshow"):
                if entry.get(key) != slide.get(key):
                    raise SnapshotError(f"Renderer mapping {key} differs on original slide {index}")
            rendered = slide.get("render", {})
            if entry.get("status") != rendered.get("status"):
                raise SnapshotError(f"Render status differs on original slide {index}")
            if rendered.get("status") == "ok":
                if entry.get("mapping_status") != "verified" or rendered.get("mapping_status") != "verified":
                    render_complete = False
                if entry.get("image") != rendered.get("image") or entry.get("sha256") != rendered.get("sha256"):
                    raise SnapshotError(f"PNG binding differs on original slide {index}")
                if entry.get("pdf_page") != rendered.get("pdf_page"):
                    raise SnapshotError(f"PDF mapping differs on original slide {index}")
            elif rendered.get("status") != "excluded":
                render_complete = False
        return {**snapshot, "verified": True, "render_complete": render_complete,
                "immutable_file_count": len(bound)}
    except SnapshotError:
        raise
    except (OSError, ValueError, TypeError, KeyError) as error:
        raise SnapshotError(f"Cannot verify snapshot: {error}") from error


def _short_text(value: str, limit: int = 160) -> str:
    value = re.sub(r"\s+", " ", value).strip().replace("|", "\\|")
    return value if len(value) <= limit else value[:limit] + "…"


def create_review_packet(out_dir: Path, inventory: dict, snapshot: dict) -> dict:
    """Write concise per-page evidence views and an explicitly unreviewed v4 template."""
    root = Path(out_dir).resolve()
    checked = verify_snapshot(root, inventory["sha256"])
    if checked["snapshot_id"] != snapshot.get("snapshot_id") or inventory.get("snapshot_id") != checked["snapshot_id"]:
        raise SnapshotError("Review packet does not belong to this evidence snapshot")
    template = {"review_schema_version": 4, "media_review_version": 1, "evidence_review_version": 1,
                "render_fidelity_review_version": inventory.get('render_fidelity_review_version', 1),
                "input_sha256": inventory["sha256"], "snapshot_id": checked["snapshot_id"],
                "run_id": checked["run_id"], "review_completion": "partial", "quality_status": "未完成",
                "preparation_only": True, "blocking_unknowns": ["尚未完成逐页视觉与语义审阅"],
                "deck_assessment": {"purpose": "待填写：材料目标与听众", "storyline": "待填写：论述关系",
                                    "visual_system": "待填写：全套视觉表达", "priorities": []}, "slides": []}
    lines = ["# PPTX 逐页审阅资料包", "", f"输入 SHA-256：`{inventory['sha256']}`",
             f"Snapshot：`{checked['snapshot_id']}`；Run：`{checked['run_id']}`", "",
             "这里只准备资料；尚未查看图片或形成质量判断。先看整页，再核对对象、原图与放映局部，最后回到整页确认裁切和遮挡。",
             "先对包含主要语言、公式和主图的代表页实际看PNG，预检中文/字形、数学符号、主图和动态字段；其余页仍须逐页核对。PDF文字 matched 和 render.status=ok 只说明文字层/导出状态，不能证明字形可见。发现主体缺失时记录受影响页并阻塞完整视觉结论。",
             "", "[空白 v4 审阅模板](review_template.json) · [完整对象清单](inventory.json) · [渲染页映射](render_manifest.json)", ""]
    assessment_keys = ("core_message", "content_organization", "visual_structure", "reading_path")
    for slide in inventory["slides"]:
        number, rendered = slide["slide"], slide.get("render", {})
        limitations = slide.get("parse_limitations", [])
        blockers = ["整页可见内容尚未由审阅者检查"]
        if rendered.get("status") != "ok":
            blockers.append(rendered.get("reason", "未获得映射已确认的整页图"))
        records = []
        for obj in _visible_objects(slide.get("objects", [])):
            if "picture" not in obj:
                continue
            picture = obj["picture"]
            original = picture.get("original_media") or picture.get("media_path") or picture.get("export_path")
            source = "；".join(path for path in (rendered.get("image"), original,
                                              (picture.get("render_crop") or {}).get("image")) if path)
            records.append({"object": obj["path"], "importance": "core", "status": "not_reviewed",
                            "source": source or "尚无可查看的图片证据", "confidence": "低", "recognized": [],
                            "unreadable_items": [], "reason": "资料已准备；尚未查看本页可见图片内容",
                            "blocking_unknowns": ["核心/次要归属与可见图片内容待审阅者确认"]})
        entry = {"slide": number, "snapshot_id": checked["snapshot_id"],
                 "render_sha256": rendered.get("sha256"), "review_completion": "partial",
                 "quality_status": "未完成", "status": "未完成", "role": "待填写：本页作用",
                 "basis": "待填写：依据本页PNG、对象与可见媒体作出的判断；资料准备不等于审阅完成",
                 "assessment": {key: "待填写：本页具体判断" for key in assessment_keys},
                 "tip_checks": [{"tip": tip, "result": "uncertain", "basis": "待填写：本页证据或不适用理由"}
                                for tip in range(1, 12)], "structural_limitations": limitations,
                 "blocking_unknowns": blockers, "issues": [], "media_reviews": records,
                 "visual_observations": [], "optional_suggestions": [], "context_questions": []}
        entry['render_fidelity'] = {
            'status': 'not_checked', 'evidence': '待实际看本页PNG并核对清单；文件产出和文字层不等于画面完整',
            'checks': {'text_glyphs': '待核对标题与正文实际字形', 'math_symbols': '待核对公式符号，或说明不适用',
                       'main_media': '待核对主图实际显示，或说明不适用', 'dynamic_fields': '待核对日期/页码等动态字段，或说明不适用'},
            'blocking_unknowns': ['本页实际画面尚未核对']}
        if rendered.get('status') != 'ok':
            entry['review_completion'] = 'failed'
        template["slides"].append(entry)
        lines += [f"## 原第 {number} 页 · {_short_text(slide.get('title') or '标题候选未知')}", "",
                  f"Slide ID：`{slide.get('slide_id')}`；Part：`{slide.get('slide_part')}`；隐藏页：`{slide.get('hidden_in_slideshow')}`",
                  f"渲染：`{rendered.get('status', 'failed')}`；映射：`{rendered.get('mapping_status', 'unverified')}`；PDF页：`{rendered.get('pdf_page')}`", ""]
        if image := rendered.get("image"):
            lines += [f"[整页 PNG]({image}) · SHA-256：`{rendered.get('sha256')}`", ""]
        else:
            lines += ["整页图不可用：" + rendered.get("reason", "渲染尚未完成"), ""]
        lines += ["| 对象 | 类型 | 可读文字/结构 |", "|---|---|---|"]
        for obj in _visible_objects(slide.get("objects", [])):
            text = obj.get("text_frame", {}).get("text", "")
            if "table" in obj:
                text = f"表格 {obj['table'].get('row_count', '?')}×{obj['table'].get('column_count', '?')}；按独立单元格核对完整对象清单"
            if "chart" in obj:
                text = f"原生图表 {obj['chart'].get('chart_type', 'unknown')}；缓存和显示须交叉核对"
            if "picture" in obj:
                text = "图片内部文字/数据待目视读取"
            lines.append(f"| `{obj.get('path')}` | {_short_text(str(obj.get('type', 'unknown')), 60)} | {_short_text(text)} |")
        lines.append("")
        for obj in _visible_objects(slide.get("objects", [])):
            if "picture" not in obj:
                continue
            picture = obj["picture"]
            original = picture.get("original_media") or picture.get("media_path") or picture.get("export_path")
            crop = picture.get("render_crop") or {}
            links = ([f"[原图]({original})"] if original else ["原图未导出"])
            if crop.get("image"):
                links.append(f"[放映局部]({crop['image']})")
            lines += [f"- 图片实例 `{obj['path']}`：{' · '.join(links)}；像素={picture.get('pixel_size')}；裁切={picture.get('crop')}；旋转={picture.get('rotation_degrees')}",
                      "  原图可能包含放映时被裁切或遮挡的区域；须回到整页确认可见范围。"]
        if limitations:
            lines += ["", "读取盲区（不是页面质量结论）："] + [f"- {item}" for item in limitations]
        comparison = rendered.get("comparison_status")
        if comparison:
            lines += ["", f"对象↔PDF文字核对：`{comparison}`（不代表视觉内容已检查）。"]
            if rendered.get("pdf_text_unattributed"):
                lines.append("- PDF中有未归属对象的文字：" + _short_text("；".join(rendered["pdf_text_unattributed"]), 500))
            if rendered.get("object_text_absent_in_pdf"):
                lines.append("- 部分对象文字未在PDF文字层匹配；对照PNG确认公式、继承或渲染差异。")
        lines.append("")
    from pptx_review_contract import aggregate_completion
    template['review_completion'] = aggregate_completion(template['slides'])
    _write_json(root / "review_template.json", template)
    (root / "review_packet.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    packet_dir = root / "review_packet"
    packet_dir.mkdir(exist_ok=True)
    (packet_dir / "README.md").write_text(
        "# 本次审阅资料\n\n[打开逐页资料包](../review_packet.md)\n\n"
        "先用代表页预检实际字形、符号、主图与动态字段，再逐页查看整页与必要局部，填写本快照 review_template.json 副本。\n"
        "render.status=ok / PDF文字 matched 不证明PNG画面忠实；主体缺失时填写render_fidelity阻塞项。\n",
        encoding="utf-8")
    return template

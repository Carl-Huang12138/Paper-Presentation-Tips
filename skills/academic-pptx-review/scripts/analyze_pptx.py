#!/usr/bin/env python3
"""Evidence-first, read-only PPTX inventory, rendering, and review report."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import tempfile
import os
import shutil
import sys
import uuid
from collections import Counter
from contextlib import contextmanager
from datetime import datetime, timezone
from decimal import Decimal, DecimalException
from pathlib import Path
import pptx_review_contract as contract

import fitz
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

fitz.TOOLS.mupdf_display_errors(False)


EMU_PER_INCH = 914400
TIP_NAMES = {
    1: "每页字数不要太多", 2: "一页只承担一个核心任务", 3: "标题要有意义",
    4: "颜色不要太多太乱", 5: "注意排版细节", 6: "用案例引出问题",
    7: "先讲例子，再讲公式", 8: "精炼阐述创新点", 9: "让听众一眼看到重点",
    10: "实验部分多用图，少用表格", 11: "总结页回顾核心内容",
}


def jdump(path: Path, obj: object) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


@contextmanager
def staging_directory(output: Path):
    # Python 3.13 TemporaryDirectory creates owner-only Windows ACLs. Renaming
    # it into the final result also preserves those ACLs, blocking the next
    # local Codex session. A regular mkdir inherits this workspace's access.
    parent = output.parent.resolve()
    path = parent / f'.{output.name}.staging-{uuid.uuid4().hex}'
    path.mkdir()
    try:
        yield path
    finally:
        if path.exists():
            if path.is_symlink() or path.resolve().parent != parent:
                raise ValueError('Refusing cleanup outside the generated staging parent')
            shutil.rmtree(path)


def safe(fn, default=None):
    try:
        return fn()
    except Exception:
        return default


def pos(shape) -> dict:
    return {key + "_in": round(getattr(shape, key) / EMU_PER_INCH, 3)
            for key in ("left", "top", "width", "height")}


def color(color_format):
    if color_format is None:
        return None
    typ = safe(lambda: str(color_format.type))
    if not typ:
        return None
    if "RGB" in typ:
        rgb = safe(lambda: color_format.rgb)
        return {"type": "rgb", "value": str(rgb) if rgb else None}
    if "SCHEME" in typ:
        theme = safe(lambda: color_format.theme_color)
        return {"type": "theme", "value": str(theme) if theme else None}
    return {"type": typ, "value": None}


def text_detail(frame) -> dict:
    paragraphs = []
    for paragraph in frame.paragraphs:
        runs = []
        for run in paragraph.runs:
            runs.append({
                "text": run.text,
                "font": {"name": run.font.name,
                         "size_pt": round(run.font.size.pt, 2) if run.font.size else None,
                         "bold": run.font.bold, "italic": run.font.italic,
                         "color": color(run.font.color)},
            })
        paragraphs.append({"text": paragraph.text, "level": paragraph.level,
                           "font_size_pt": paragraph.font.size.pt if paragraph.font.size else None,
                           "alignment": str(paragraph.alignment) if paragraph.alignment else None,
                           "runs": runs})
    return {"text": frame.text, "paragraphs": paragraphs,
            "margins_in": {side: round(getattr(frame, "margin_" + side) / EMU_PER_INCH, 3)
                           for side in ("left", "right", "top", "bottom")},
            "word_wrap": frame.word_wrap}


def chart_detail(chart) -> dict:
    result = {"chart_type": str(safe(lambda: chart.chart_type, "unknown")),
              "title": "unknown", "categories": "unknown", "series": [],
              "value_axis": "unknown", "category_axis": "unknown", "legend": None,
              "limitations": []}
    if safe(lambda: chart.has_title, False):
        result["title"] = safe(lambda: chart.chart_title.text_frame.text, "unknown")
    if safe(lambda: chart.has_legend, False):
        result["legend"] = {"position": str(safe(lambda: chart.legend.position, "unknown"))}
    result["categories"] = safe(lambda: [str(c.label) for c in chart.plots[0].categories], "unknown")
    for series in safe(lambda: list(chart.series), []):
        values = safe(lambda: list(series.values), "unknown")
        result["series"].append({"name": safe(lambda: series.name, "unknown"), "values": values,
                                 "fill_color": safe(lambda: color(series.format.fill.fore_color)),
                                 "line_color": safe(lambda: color(series.format.line.color))})
    for axis_name in ("value_axis", "category_axis"):
        axis = safe(lambda: getattr(chart, axis_name))
        if axis:
            result[axis_name] = {
                "title": safe(lambda: axis.axis_title.text_frame.text, None)
                if safe(lambda: axis.has_title, False) else None,
                "minimum_scale": safe(lambda: axis.minimum_scale),
                "maximum_scale": safe(lambda: axis.maximum_scale),
                "tick_labels": "unknown",  # Tick labels are renderer-generated, not OOXML data.
            }
    if result["categories"] == "unknown" or any(s["values"] == "unknown" for s in result["series"]):
        result["limitations"].append("图表缓存数据不可读取；不得推测数值或坐标轴标签")
    return result


def shape_detail(shape, path: str) -> dict:
    item = {"id": shape.shape_id, "path": path, "name": shape.name,
            "type": str(shape.shape_type), "bounds": pos(shape),
            "source": "pptx_object", "z_order": int(path.split("/")[0].split("-")[-1]),
            "visible_in_slideshow": True}
    hidden = safe(lambda: shape._element.xpath(".//p:cNvPr")[0].get("hidden"))
    if hidden in ("1", "true"):
        item["visible_in_slideshow"] = False
    item["fill_color"] = safe(lambda: color(shape.fill.fore_color))
    item["line_color"] = safe(lambda: color(shape.line.color))
    if shape.has_text_frame:
        item["text_frame"] = text_detail(shape.text_frame)
    if shape.has_table:
        table = shape.table
        item["table"] = {"rows": [[{"text": cell.text, "text_frame": text_detail(cell.text_frame),
                                     "fill_color": safe(lambda: color(cell.fill.fore_color))}
                                    for cell in row.cells] for row in table.rows],
                         "row_count": len(table.rows), "column_count": len(table.columns)}
    if shape.has_chart:
        item["chart"] = chart_detail(shape.chart)
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        item["picture"] = {"content_type": safe(lambda: shape.image.content_type),
                           "filename": safe(lambda: shape.image.filename),
                           "sha256": safe(lambda: hashlib.sha256(shape.image.blob).hexdigest()),
                           "text_or_data_inside_image": "unknown",
                           "extraction_status": "not_structurally_parsed"}
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        item["children"] = [shape_detail(child, f"{path}/{i}")
                            for i, child in enumerate(shape.shapes, 1)]
    if "OLE" in str(shape.shape_type):
        item["embedded_object"] = {"meaning": "unknown", "render_only": True}
    return item


def flatten(items):
    for item in items:
        yield item
        yield from flatten(item.get("children", []))


def visible_objects(items):
    for item in items:
        if item.get("visible_in_slideshow", True):
            yield item
            yield from visible_objects(item.get("children", []))


def typography(items: list[dict]) -> dict:
    """Locate explicit run/paragraph sizes; theme/layout inheritance stays unknown."""
    samples = []

    def frames(objects):
        for obj in objects:
            if not obj.get("visible_in_slideshow", True):
                continue
            if "text_frame" in obj:
                yield obj["path"], obj["text_frame"]
            if "table" in obj:
                for row_num, row in enumerate(obj["table"]["rows"], 1):
                    for col_num, cell in enumerate(row, 1):
                        yield f"{obj['path']}:cell-{row_num}-{col_num}", cell["text_frame"]
            yield from frames(obj.get("children", []))

    for path, frame in frames(items):
        for paragraph_num, paragraph in enumerate(frame["paragraphs"], 1):
            for run in paragraph["runs"]:
                if not run["text"].strip():
                    continue
                run_size = run["font"]["size_pt"]
                size = run_size if run_size is not None else paragraph.get("font_size_pt")
                samples.append({"object": path, "paragraph": paragraph_num,
                                "text": run["text"][:100], "size_pt": size,
                                "size_source": "run" if run_size is not None else
                                "paragraph" if size is not None else "unknown"})
    sizes = [s["size_pt"] for s in samples if s["size_pt"] is not None]
    return {"samples": samples, "minimum_explicit_size_pt": min(sizes) if sizes else None,
            "unknown_size_runs": sum(s["size_pt"] is None for s in samples),
            "limitations": "主题/母版继承字号、图表标签和图片内字号未可靠解析；需核对渲染。"}


def extract(pptx_path: Path, media_dir: Path | None = None) -> dict:
    from pptx_reader import extract as read_objects
    return read_objects(pptx_path, media_dir=media_dir)


def render(pptx_path: Path, out_dir: Path, slides: list[dict], **options) -> list[str]:
    from pptx_runtime import render as render_snapshot
    return render_snapshot(pptx_path, out_dir, slides, **options)


def role(slide: dict) -> str:
    title = slide.get("title") or ""
    if re.search(r"总结|结论|致谢|结束|Q\s*&\s*A", title, re.I): return "总结/收尾"
    if re.search(r"实验|结果|性能|消融|评估", title): return "实验/证据"
    if re.search(r"方法|模型|创新|GRPO|ViT|Q-learning", title, re.I): return "方法/概念"
    if re.search(r"背景|动机|问题|任务定义|推理", title): return "背景/动机"
    return "待人工判定"


def preliminary(slide: dict) -> dict:
    # Mechanical cues do not become confirmed quality defects without a visual review.
    objects = list(flatten(slide["objects"]))
    text_count = sum(len(x) for x in slide["visible_text"])
    charts = sum("chart" in o for o in objects)
    tables = sum("table" in o for o in objects)
    pictures = sum("picture" in o for o in objects)
    cues = []
    small = [s for s in slide.get("typography", {}).get("samples", [])
             if s["size_pt"] is not None and s["size_pt"] < 20]
    if small:
        cues.append({"kind": "small_font",
                     "evidence": "；".join(f"{s['object']} 段落{s['paragraph']} {s['size_pt']:g} pt"
                                           for s in small[:8]),
                     "action": "区分主体/图表标签/脚注，核对实际可读性；优先扩展区域、减少重复或拆页，尽量放大并协调字号，不能只凭阈值判错"})
    if not (pictures or charts or tables) and text_count > 200:
        cues.append({"kind": "text_only_structure",
                     "evidence": f"主体对象包含约 {text_count} 字符，未读取到独立图片、表格或图表",
                     "action": "核对是否只有同样式文字且缺少阅读锚点；原生形状也可能已形成有效图示。必要时用现有内容做流程、对照或示例，不添加无关装饰"})
    if text_count > 450:
        cues.append({"kind": "dense_text", "evidence": f"可见对象文本约 {text_count} 字符",
                     "action": "核对渲染图中的字号、行距及听众能否先抓住核心句；字多本身不是问题"})
    if charts or tables:
        cues.append({"kind": "data_display", "evidence": f"{tables} 个表格、{charts} 个图表",
                     "action": "检查行列、轴、图例、数据及结论的对应关系；表格本身不是问题"})
    if pictures:
        cues.append({"kind": "picture_visual_review", "evidence": f"{pictures} 张独立图片",
                     "action": "看渲染图并记录图片内可读元素；未确认的文字和数值保持未知"})
    if slide["render"]["status"] != "ok":
        cues.append({"kind": "render_failed", "evidence": slide["render"]["reason"],
                     "action": "该页视觉判断不可完成，需重新渲染或人工打开"})
    return {"slide": slide["slide"], "role": role(slide), "status": "待视觉审阅",
            "basis": "已完成 PPTX 对象读取；需要把渲染图与对象清单逐项核对",
            "issues": [], "uncertainties": cues}


def validate_review(review: dict, inventory: dict) -> None:
    version = review.get("review_schema_version", 1)
    if version not in (1, 2, 3, 4):
        raise ValueError("review_schema_version 只能为 1、2、3 或 4")
    if review.get("media_review_version") not in (None, 1):
        raise ValueError("media_review_version 只能为 1")
    if review.get("evidence_review_version") not in (None, 1):
        raise ValueError("evidence_review_version 只能为 1")
    if review.get('render_fidelity_review_version') not in (None, 1):
        raise ValueError('render_fidelity_review_version 只能为 1')
    require_fidelity = (version == 4 and inventory.get('render_fidelity_review_version') == 1) or review.get('render_fidelity_review_version') == 1
    if require_fidelity and (version != 4 or review.get('render_fidelity_review_version') != 1):
        raise ValueError('新快照须启用 render_fidelity_review_version: 1，逐页核对实际画面')
    if review.get("evidence_review_version") == 1 and version not in (3, 4):
        raise ValueError("evidence_review_version: 1 须搭配 schema v3/v4")
    if review.get("input_sha256") != inventory["sha256"]:
        raise ValueError("审阅记录与输入 PPTX 的 SHA-256 不一致")
    if version == 4:
        if not inventory.get("snapshot_id") or review.get("snapshot_id") != inventory["snapshot_id"]:
            raise ValueError("v4 必须绑定当前 snapshot_id，不能借用另一批次")
        if review.get("evidence_review_version") != 1:
            raise ValueError("v4 必须启用 evidence_review_version: 1")
    entries = review.get("slides", [])
    if not isinstance(entries, list) or any(not isinstance(e, dict) or type(e.get("slide")) is not int for e in entries):
        raise ValueError("review.slides 须为有整数页码的对象列表")
    nums = [x.get("slide") for x in entries]
    if sorted(nums) != list(range(1, inventory["slide_count"] + 1)):
        raise ValueError("review.slides 必须覆盖每一页且页码不能重复")
    if version == 4:
        expected_completion = contract.aggregate_completion(entries)
        expected_quality = ('未完成' if expected_completion != 'complete' else
                            '有确定问题' if any(e.get('status') == '有确定问题' for e in entries) else
                            '需要人工确认' if any(e.get('status') == '需要人工确认' for e in entries) else '无需修改')
        if review.get('preparation_only') not in (None, False):
            raise ValueError('正式导入须清除 preparation_only 资料准备标记')
        if 'review_completion' in review and review['review_completion'] != expected_completion:
            raise ValueError('顶层 review_completion 与逐页完成度不一致')
        if 'quality_status' in review and review['quality_status'] != expected_quality:
            raise ValueError('顶层 quality_status 与逐页结论不一致')
        if 'blocking_unknowns' in review:
            contract.strings(review['blocking_unknowns'], '顶层 blocking_unknowns')
            if expected_completion == 'complete' and review['blocking_unknowns']:
                raise ValueError('存在顶层核心未知，不能声称 complete')
    valid_status = {"无需修改", "有确定问题", "需要人工确认"} | ({"未完成"} if version == 4 else set())
    required_issue = {"object", "evidence", "impact", "severity", "confidence", "tips", "direction", "category"}
    if version >= 3:
        deck = review.get("deck_assessment")
        if not isinstance(deck, dict) or any(not isinstance(deck.get(k), str) or not deck[k].strip()
                                            for k in ("purpose", "storyline", "visual_system")):
            raise ValueError("v3 缺少具体的 deck_assessment：purpose/storyline/visual_system")
        priorities = deck.get("priorities")
        if not isinstance(priorities, list) or any(not isinstance(x, str) or not x.strip() for x in priorities):
            raise ValueError("deck_assessment.priorities 必须为字符串列表；足够好时可以为空")
        if review.get("media_review_version") != 1:
            raise ValueError("v3 必须提供 media_review_version: 1")
    for entry in entries:
        slide = inventory["slides"][entry["slide"] - 1]
        if version == 4 and 'snapshot_id' in entry and entry['snapshot_id'] != review['snapshot_id']:
            raise ValueError(f"第 {entry['slide']} 页 snapshot_id 与当前快照不一致")
        if entry.get("status") not in valid_status:
            raise ValueError(f"第 {entry['slide']} 页 status 无效")
        contract.require_text(entry, ["basis", "role"], f"第 {entry['slide']} 页")
        if not isinstance(entry.get("issues"), list) or any(not isinstance(i, dict) for i in entry["issues"]):
            raise ValueError("issues 须为对象列表")
        contract.validate_observations(entry)
        if entry["status"] == "无需修改" and entry.get("issues"):
            raise ValueError(f"第 {entry['slide']} 页不能同时标为无需修改和有问题")
        if entry["status"] == "有确定问题" and not entry.get("issues"):
            raise ValueError(f"第 {entry['slide']} 页缺少问题")
        if version >= 3:
            validate_assessment(entry)
        objects = {o["path"]: o for o in flatten(inventory["slides"][entry["slide"] - 1]["objects"])}
        object_paths = set(objects)
        validate_media_reviews(entry, inventory["slides"][entry["slide"] - 1],
                               require_coverage=review.get("media_review_version") == 1)
        contract.validate_completion(entry, slide, inventory, version, require_fidelity=require_fidelity)
        if version == 4:
            contract.validate_suggestions(entry)
        for issue in entry.get("issues", []):
            if missing := required_issue - issue.keys():
                raise ValueError(f"第 {entry['slide']} 页问题缺少字段：{sorted(missing)}")
            contract.require_text(issue, ["object", "evidence", "impact", "direction", "category"], "issue")
            if issue["object"] not in object_paths and not issue["object"].startswith("region:"):
                raise ValueError(f"第 {entry['slide']} 页对象 {issue['object']} 不在对象清单中")
            if issue["severity"] not in ("高", "中", "低") or issue["confidence"] not in ("高", "中", "低"):
                raise ValueError(f"第 {entry['slide']} 页严重程度或置信度无效")
            validate_numeric_checks(issue, entry['slide'])
            if not isinstance(issue["tips"], list) or any(type(tip) is not int or tip not in TIP_NAMES for tip in issue["tips"]) or len(set(issue["tips"])) != len(issue["tips"]):
                raise ValueError(f"第 {entry['slide']} 页建议编号无效")
            if not issue["tips"] and not contract.nonempty(issue.get("tip_mapping_reason")):
                raise ValueError("未映射建议条目须有 tip_mapping_reason")
            if version >= 3:
                checks = {c["tip"]: c["result"] for c in entry["tip_checks"]}
                if any(checks[t] != "problem" for t in issue["tips"]):
                    raise ValueError("issue 对应 tip_checks 须为 problem，不能同时 adequate")
            if version == 4:
                if issue.get("scope") not in ("visual", "presentation", "critical_logic", "factual_error"):
                    raise ValueError("issue.scope 须为 presentation/visual/critical_logic/factual_error")
                if issue["category"] in ("学术证据", "事实性错误") and issue["scope"] not in ("critical_logic", "factual_error"):
                    raise ValueError("事实修正须明确 factual_error/critical_logic")
                if issue.get("recommendation_kind") != "necessary":
                    raise ValueError("issues 仅容纳 necessary；可选改善请用 optional_suggestions")
                contract.require_text(issue, ["obstacle", "benefit"], "necessary issue")
                contract.validate_plan_v4(issue, entry["slide"])
                contract.validate_disposition(issue, inventory)
                contract.validate_numeric_sources(issue, inventory)
            elif version >= 2:
                validate_polish_plan(issue, entry["slide"], strict=version >= 3)
            contract.validate_target(issue, objects, entry["slide"],
                                     required=review.get("evidence_review_version") == 1, strict=version == 4)


def validate_target_evidence(issue: dict, objects: dict, page: int, required: bool) -> None:
    contract.validate_target(issue, objects, page, required)


def numeric_result(check: dict):
    """Recompute declared arithmetic only; never evaluate supplied expressions."""
    if not isinstance(check, dict) or not isinstance(check.get("inputs"), dict):
        raise ValueError("numeric_checks 缺少 inputs")
    data = check["inputs"]

    def number(value):
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise ValueError("numeric_checks 数值须为有限 JSON 数字")
        value = Decimal(str(value))
        if not value.is_finite():
            raise ValueError("numeric_checks 数值须为有限 JSON 数字")
        return value

    try:
        if check.get("operation") == "geometric_sum":
            start, end = data.get("start"), data.get("end")
            if (type(start) is not int or type(end) is not int or not 0 <= start <= end <= 10000):
                raise ValueError("geometric_sum 须有 0≤start≤end≤10000 的整数下标")
            coefficient, ratio = number(data.get("coefficient")), number(data.get("ratio"))
            # Decimal treats 0**0 as invalid; an empty product in a finite sum is 1.
            return coefficient * sum((Decimal(1) if k == 0 else ratio ** k for k in range(start, end + 1)), Decimal(0))
        if check.get("operation") == "compare_gains":
            rows = data.get("rows")
            if not isinstance(rows, list) or not rows or len(rows) > 10000:
                raise ValueError("compare_gains 须有非空 rows")
            gains = {}
            for row in rows:
                if not isinstance(row, dict) or not isinstance(row.get("label"), str) or not row["label"].strip() or row["label"] in gains:
                    raise ValueError("compare_gains 标签须非空且唯一")
                gains[row["label"]] = number(row.get("after")) - number(row.get("before"))
            maximum = max(gains.values())
            return [label for label, value in gains.items() if value == maximum]
    except DecimalException as error:
        raise ValueError("numeric_checks 无法有限复算") from error
    raise ValueError("numeric_checks.operation 仅支持 geometric_sum / compare_gains")


def validate_numeric_checks(issue: dict, page: int) -> None:
    checks = issue.get("numeric_checks", [])
    if not isinstance(checks, list):
        raise ValueError(f"第 {page} 页 numeric_checks 须为列表")
    for check in checks:
        result = numeric_result(check)
        if not isinstance(check.get("basis"), str) or not check["basis"].strip():
            raise ValueError(f"第 {page} 页 numeric_checks 须说明来源、代入前提或比较范围")
        reported = check.get("reported")
        if isinstance(result, list):
            valid = (isinstance(reported, list) and all(isinstance(x, str) for x in reported)
                     and len(reported) == len(set(reported)) and set(reported) == set(result))
        else:
            valid = (isinstance(reported, (int, float)) and not isinstance(reported, bool)
                     and Decimal(str(reported)).is_finite()
                     and abs(Decimal(str(reported)) - result) <= Decimal("1e-9") * max(abs(result), Decimal(1)))
        if not valid:
            raise ValueError(f"第 {page} 页数值复算与 reported 不一致：实际为 {result}")


def validate_assessment(entry: dict) -> None:
    prefix = f"第 {entry['slide']} 页"
    limits = entry.get("structural_limitations", [])
    if not isinstance(limits, list) or any(not isinstance(x, str) or not x.strip() for x in limits):
        raise ValueError(f"{prefix} structural_limitations 必须为字符串列表")
    assessment = entry.get("assessment")
    for key in ("core_message", "content_organization", "visual_structure", "reading_path"):
        if not isinstance(assessment, dict) or not isinstance(assessment.get(key), str) or not assessment[key].strip():
            raise ValueError(f"{prefix} assessment.{key} 必须有具体判断")
    checks = entry.get("tip_checks")
    if not isinstance(checks, list) or len(checks) != 11 or any(not isinstance(c, dict) for c in checks):
        raise ValueError(f"{prefix} tip_checks 必须覆盖全部 11 条建议")
    if any(type(c.get("tip")) is not int for c in checks) or Counter(c.get("tip") for c in checks) != Counter(list(TIP_NAMES)):
        raise ValueError(f"{prefix} tip_checks 必须覆盖全部 11 条建议且不能重复")
    for check in checks:
        if check.get("result") not in ("adequate", "problem", "not_applicable", "uncertain"):
            raise ValueError(f"{prefix} tip_checks.result 无效")
        if not isinstance(check.get("basis"), str) or not check["basis"].strip():
            raise ValueError(f"{prefix} tip_checks.basis 须有页面证据或不适用理由")
        if check["result"] == "problem":
            if entry["status"] == "无需修改":
                raise ValueError(f"{prefix} 无需修改不能同时记录已确定的典型问题")
            if not any(check["tip"] in i.get("tips", []) for i in entry.get("issues", [])):
                raise ValueError(f"{prefix} 建议 {check['tip']} 的确定问题缺少对应修改记录")


def validate_media_reviews(entry: dict, slide: dict, require_coverage: bool) -> None:
    pictures = {o["path"] for o in visible_objects(slide["objects"]) if "picture" in o}
    records = entry.get("media_reviews", [])
    if not isinstance(records, list):
        raise ValueError(f"第 {entry['slide']} 页 media_reviews 必须是列表")
    seen = set()
    for record in records:
        if not isinstance(record, dict) or record.get("object") not in pictures:
            raise ValueError(f"第 {entry['slide']} 页图片对象 {record.get('object') if isinstance(record, dict) else record} 无效")
        obj = record["object"]
        if obj in seen:
            raise ValueError(f"第 {entry['slide']} 页图片对象 {obj} 记录重复")
        seen.add(obj)
        status = record.get("status")
        if status not in ("reviewed", "partial", "not_reviewed"):
            raise ValueError(f"第 {entry['slide']} 页 {obj} 图片检查状态无效")
        if record.get("confidence") not in ("高", "中", "低") or not isinstance(record.get("source"), str) or not record["source"].strip():
            raise ValueError(f"第 {entry['slide']} 页 {obj} 缺少图片来源或置信度")
        recognized = record.get("recognized")
        unreadable = record.get("unreadable_items")
        if not isinstance(recognized, list) or not isinstance(unreadable, list) or any(not isinstance(x, str) or not x.strip() for x in unreadable):
            raise ValueError(f"第 {entry['slide']} 页 {obj} 缺少 recognized/unreadable_items 列表")
        for item in recognized:
            if not isinstance(item, dict) or item.get("kind") not in ("visual_content", "exact_value", "approximate_value") or not isinstance(item.get("content"), str) or not item["content"].strip():
                raise ValueError(f"第 {entry['slide']} 页 {obj} 图片观察类型或内容无效")
            if item["kind"] == "approximate_value" and not item.get("basis"):
                raise ValueError(f"第 {entry['slide']} 页 {obj} 近似读图值缺少 basis")
        if status == "reviewed" and (not recognized or unreadable):
            raise ValueError(f"第 {entry['slide']} 页 {obj} reviewed 须有可辨认内容且没有未读项")
        if status == "partial" and not unreadable:
            raise ValueError(f"第 {entry['slide']} 页 {obj} partial 须列明无法辨认项")
        if status == "not_reviewed" and (recognized or not record.get("reason")):
            raise ValueError(f"第 {entry['slide']} 页 {obj} not_reviewed 须说明未查看原因且不能声称已读出内容")
    if require_coverage and seen != pictures:
        raise ValueError(f"第 {entry['slide']} 页图片审阅未完整覆盖：{sorted(pictures - seen)}")


def media_coverage(inventory: dict, diagnoses: list[dict]) -> list[dict]:
    entries = {d["slide"]: d for d in diagnoses}
    result = []
    for slide in inventory["slides"]:
        records = {m["object"]: m for m in entries.get(slide["slide"], {}).get("media_reviews", [])}
        for obj in visible_objects(slide["objects"]):
            if "picture" in obj:
                record = records.get(obj["path"], {"object": obj["path"], "status": "unrecorded",
                                                   "recognized": [], "unreadable_items": []})
                result.append({**record, "slide": slide["slide"],
                               "structural_status": "not_structurally_parsed"})
    return result


def validate_polish_plan(issue: dict, page: int, strict: bool = False) -> None:
    prefix = f"第 {page} 页"

    def fields(obj, names, label):
        if not isinstance(obj, dict):
            raise ValueError(f"{prefix} 缺少 {label}")
        for key in names:
            if not isinstance(obj.get(key), str) or not obj[key].strip():
                raise ValueError(f"{prefix} {label}.{key} 必须有具体内容")

    def strings(obj, key, label):
        value = obj.get(key)
        if not isinstance(value, list) or not value or any(not isinstance(x, str) or not x.strip() for x in value):
            raise ValueError(f"{prefix} {label}.{key} 必须是非空字符串列表")

    if issue.get("scope") not in ("visual", "presentation", "critical_logic", "factual_error"):
        raise ValueError(f"{prefix} scope 必须为 visual、presentation、critical_logic 或 factual_error")
    if issue["category"] in ("学术证据", "事实性错误") and issue["scope"] not in ("critical_logic", "factual_error"):
        raise ValueError(f"{prefix} 学术事实问题仅接受 critical_logic 或 factual_error")
    if issue.get("action") not in ("revise", "split"):
        raise ValueError(f"{prefix} action 必须为 revise 或 split")
    plan = issue.get("revision_plan")
    fields(plan, ("goal", "layout", "typography", "check"), "revision_plan")
    strings(plan, "steps", "revision_plan")
    if strict and "parameters" in issue:
        raise ValueError(f"{prefix} parameters 须放在 revision_plan.parameters，不能放在问题顶层")
    if "parameters" in plan:
        strings(plan, "parameters", "revision_plan")
    # Recognize common explicit proposals, including conditional ones. This is
    # a completeness guard, not a semantic classifier or a reason to split.
    proposal = "\n".join([issue["direction"], plan["layout"], *plan["steps"]])
    split_operation = r"(?:拆页|(?:拆成|拆分为|拆分成|分成)\s*[两二三四五六七八九十多\d]+\s*页)"
    proposal = re.sub(rf"(?:不需要|无需|不要|不必|避免|不)(?:再|继续|建议|考虑|需要|进行|将内容|把内容|将页面|把页面|进行内容)*{split_operation}", "", proposal)
    english_split = r"\bsplit\s+(?:\w+\s+){0,5}(?:slides|pages)\b"
    proposal = re.sub(rf"(?:not|don't|no need to|avoid)\s+{english_split}", "", proposal, flags=re.I)
    explicit_split = bool(re.search(split_operation, proposal) or
                          re.search(english_split, proposal, re.I))
    if issue["action"] == "split" or "split_plan" in issue or (strict and explicit_split):
        pages = issue.get("split_plan")
        if not isinstance(pages, list) or len(pages) < 2:
            raise ValueError(f"{prefix} split_plan 至少包含两个新页")
        for new_page in pages:
            fields(new_page, ("title", "message", "layout"), "split_plan")
            strings(new_page, "content", "split_plan")
            strings(new_page, "components", "split_plan")


def markdown(inventory: dict, diagnoses: list[dict], render_issues: list[str],
             deck_assessment: dict | None = None) -> str:
    statuses = Counter(d["status"] for d in diagnoses)
    out = ["# 学术汇报 PPTX：视觉润色诊断", "",
           f"输入：`{inventory['input']}`  ", f"SHA-256：`{inventory['sha256']}`  ",
           f"页面：{inventory['slide_count']}；已渲染：{sum(s['render']['status']=='ok' for s in inventory['slides'])}  ",
           f"结论：{dict(statuses)}", "", "## 整套主要发现", ""]
    if deck_assessment:
        out += [f"汇报目标：{deck_assessment['purpose']}", "",
                f"论述与结构：{deck_assessment['storyline']}", "",
                f"整体观感：{deck_assessment['visual_system']}", "", "### 优先修改顺序", ""]
        out += [f"{n}. {p}" for n, p in enumerate(deck_assessment["priorities"], 1)] or [
            "- 尚未列出修改顺序；请以逐页问题与完成度为准。" if any(d.get("issues") or d["status"] != "无需修改" for d in diagnoses)
            else "- 已检查的页面均无需必要修改。"]
        out += ["", "### 主要问题", ""]
    ranked = [(d["slide"], i) for d in diagnoses for i in d.get("issues", [])]
    ranked.sort(key=lambda x: {"高": 0, "中": 1, "低": 2}[x[1]["severity"]])
    if ranked:
        for slide_num, issue in ranked[:8]:
            out.append(f"- 第 {slide_num} 页 · {issue['category']}：{issue['impact']}（{issue['severity']}）")
    else:
        out.append("- 已审阅页面无需修改。" if all(d["status"] == "无需修改" for d in diagnoses)
                   else "- 尚无确认的问题；待审阅或疑点见逐页记录。")
    out += ["", "## 逐页结论", ""]
    for slide, d in zip(inventory["slides"], diagnoses):
        render_link = f"renders/slide-{slide['slide']:02d}.png" if slide["render"]["status"] == "ok" else "无"
        title = d.get("confirmed_title") or slide.get("title") or "标题未确认"
        out += [f"### 第 {slide['slide']} 页 · {title}", "",
                f"作用：{d['role']}；结论：**{d['status']}**；[渲染图]({render_link})", "",
                f"判断依据：{d['basis']}", ""]
        if d.get("review_completion"):
            out += [f"检查完成度：{d['review_completion']}；质量状态与完成范围分别记录。", ""]
        if a := d.get("assessment"):
            out += [f"- 核心信息：{a['core_message']}", f"- 内容组织：{a['content_organization']}",
                    f"- 视觉结构：{a['visual_structure']}", f"- 阅读路径：{a['reading_path']}", ""]
        if checks := d.get("tip_checks"):
            labels = {"adequate": "可接受", "problem": "确定问题", "not_applicable": "不适用", "uncertain": "待确认"}
            out += ["<details>", "<summary>典型问题检查：11 项逐项判断（不作机械打分）</summary>", ""]
            out += [f"- {c['tip']} · {TIP_NAMES[c['tip']]}：{labels[c['result']]}；{c['basis']}" for c in checks]
            out += ["", "</details>", ""]
        if d.get("visual_observations") or d.get("media_reviews"):
            out += ["<details>", "<summary>对象与图片读取证据</summary>", ""]
        for observation in d.get("visual_observations", []):
            out.append(f"- 渲染核对：{observation['observation']}（来源：{observation['source']}；置信度：{observation['confidence']}）")
        for media in d.get("media_reviews", []):
            labels = {"reviewed": "已完成相关可见内容检查", "partial": "部分可辨认", "not_reviewed": "未完成图片查看"}
            out.append(f"- 图片检查 · `{media['object']}`：{labels[media['status']]}（来源：{media['source']}；置信度：{media['confidence']}）")
            kinds = {"visual_content": "可辨认内容", "exact_value": "精确数值", "approximate_value": "近似读图值"}
            for item in media["recognized"]:
                out.append(f"  - {kinds[item['kind']]}：{item['content']}" +
                           (f"；读图依据：{item['basis']}" if item.get("basis") else ""))
            for item in media["unreadable_items"]:
                out.append(f"  - 无法辨认：{item}")
            if media.get("reason"):
                out.append(f"  - 未完成原因：{media['reason']}")
        if d.get("visual_observations") or d.get("media_reviews"):
            out += ["", "</details>", ""]
        for issue in d.get("issues", []):
            tips = "、".join(f"{n}（{TIP_NAMES[n]}）" for n in issue["tips"])
            out += [f"- **{issue['category']} / {issue['severity']} / 置信度{issue['confidence']}** · `{issue['object']}`",
                    f"  - 证据：{issue['evidence']}", f"  - 对听众的影响：{issue['impact']}",
                    f"  - 修改方向：{issue['direction']}", f"  - 对应建议：{tips}"]
            if issue.get("target_evidence") or issue.get("numeric_checks"):
                out += ["", "<details>", "<summary>目标归属与数值复算（不代表完整语义验收）</summary>", ""]
                if anchor := issue.get("target_evidence"):
                    out.append(f"- 目标 `{issue['object']}`；来源：{anchor['source']}；"
                               + (f"原文引用：{anchor['quote']}" if anchor.get("quote") else anchor.get("description", "")))
                    for ref in anchor.get("refs", []):
                        out.append(f"  - 来源 `{ref.get('object', issue['object'])}`：{ref.get('quote') or ref.get('description')}")
                for check in issue.get("numeric_checks", []):
                    out.append(f"- 数值复算：{check['operation']}；输入：`{json.dumps(check['inputs'], ensure_ascii=False)}`；"
                               f"结果：{numeric_result(check)}；依据：{check['basis']}")
                out += ["", "</details>", ""]
            plan = issue.get("revision_plan")
            if plan:
                if plan.get("goal"):
                    out.append(f"  - 修改方案目标：{plan['goal']}")
                if plan.get("layout"):
                    out.append(f"  - 页面结构：{plan['layout']}")
                out += [f"    {i}. {step}" for i, step in enumerate(plan["steps"], 1)]
                if plan.get("typography"):
                    out.append(f"  - 字体与协调：{plan['typography']}")
                out.append(f"  - 修改后检查：{plan['check']}")
                if plan.get("parameters"):
                    out += ["", "<details>", "<summary>必要参数（设计提案，须重新渲染验证）</summary>", ""]
                    out += [f"- {p}" for p in plan["parameters"]]
                    out += ["", "</details>", ""]
            split_groups = [("拆页方案（按顺序）", issue.get("split_plan", []))] + [
                (f"备选拆页，启用条件：{a['activation_condition']}", a['pages']) for a in issue.get("alternatives", [])]
            for label, pages in split_groups:
                if not pages:
                    continue
                out += ["", f"  **{label}**", ""]
                for i, new_page in enumerate(pages, 1):
                    out += [f"  - 新页 {i}：**{new_page['title']}**",
                            f"    - 核心信息：{new_page['message']}",
                            f"    - 内容：{'；'.join(new_page['content'])}",
                            f"    - 结构：{new_page['layout']}",
                            f"    - 组件：{'；'.join(new_page['components'])}"]
            for allocation in issue.get("content_disposition", []):
                ref = allocation['source_ref']
                out.append(f"  - 材料去向：第{ref['slide']}页 `{ref['object']}` → {allocation['destination']}；{allocation['disposition']}；{allocation['reason']}")
        for suggestion in d.get("optional_suggestions", []):
            out.append(f"- 可选改善（当前无需必要修改）：{suggestion['direction']}；收益：{suggestion['benefit']}")
        for question in d.get("context_questions", []):
            out.append(f"- 上下文待确认：{question['evidence']}；{question['action']}")
        out += [f"- 阻塞检查：{x}" for x in d.get("blocking_unknowns", [])]
        for doubt in d.get("uncertainties", []):
            if isinstance(doubt, dict):
                out.append(f"- 待核对：{doubt['evidence']}；{doubt['action']}")
            else:
                out.append(f"- 待核对：{doubt}")
        out.append("")
    out += ["## 读取能力与未完成项", "", "### 实际失败或无法辨认项", ""]
    remaining = []
    for issue in render_issues:
        remaining.append(f"- {issue}")
    structural_boundaries = []
    for slide in inventory["slides"]:
        for blind in slide["parse_limitations"]:
            picture_paths = {o["path"] for o in flatten(slide["objects"]) if "picture" in o}
            if blind.split(":", 1)[0] not in picture_paths:
                structural_boundaries.append(f"- 第 {slide['slide']} 页 · {blind}")
    for d in diagnoses:
        structural_boundaries += [f"- 第 {d['slide']} 页 · {limit}" for limit in d.get("structural_limitations", [])]
    coverage = media_coverage(inventory, diagnoses)
    for media in coverage:
        for item in media["unreadable_items"]:
            remaining.append(f"- 第 {media['slide']} 页 · `{media['object']}`：{item}")
        if media["status"] == "not_reviewed":
            remaining.append(f"- 第 {media['slide']} 页 · `{media['object']}`：未完成图片查看；{media['reason']}")
    out += remaining or ["- 未记录具体的渲染失败或无法辨认项；这不代表未记录的检查已经完成。"]
    out += ["", "### 图片检查记录覆盖", ""]
    counts = Counter(m["status"] for m in coverage)
    out.append(f"- 可见图片 {len(coverage)} 个；结构化检查状态：{dict(counts)}。")
    unrecorded = [m for m in coverage if m["status"] == "unrecorded"]
    if unrecorded:
        out.append("- 未提供结构化图片检查状态：" + "；".join(f"第 {m['slide']} 页 {m['object']}" for m in unrecorded) +
                   "。旧记录的逐页 visual_observations 可能已含目视结果；不能将状态缺失解释为图片不可辨认。")
    out += ["", "### 对象解析能力边界", ""] + structural_boundaries + [
            "- 图片未由对象 API 结构化解析，与目视可辨认程度分别记录；视觉阅读不代表取得原始点数据或可编辑结构。",
            "- SmartArt、OLE 内部结构和图表自动刻度未可靠还原时仍需视觉补充；只有具体受影响项才列为实际盲区。",
            "- PPTX 备注和内部媒体元数据没有被当作放映时可见内容；字体替换可能使 LibreOffice 渲染不同于 PowerPoint。", ""]
    return "\n".join(out)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--review", type=Path, help="逐页视觉审阅 JSON；缺省只产出待审阅线索")
    parser.add_argument("--require-evidence", action="store_true", help="最终审阅须使用当前 v4 快照和 evidence_review_version: 1")
    parser.add_argument("--legacy-review", action="store_true", help="显式导入 v1-v3 历史记录；不称新版验收")
    parser.add_argument("--soffice", help="LibreOffice 可执行文件；然后查 PPTX_SOFFICE、PATH、常见路径")
    parser.add_argument("--dpi", type=int, default=144, help="整页 PNG 分辨率，默认144 DPI")
    parser.add_argument("--no-keep-pdf", action="store_true", help="不保留中间PDF")
    parser.add_argument("--render-scope", choices=("all", "slides"), default="all", help="all包含隐藏页；slides仅放映页")
    parser.add_argument("--rerender", action="store_true", help="创建新快照，须指定一个新输出目录")
    args = parser.parse_args()
    if not args.input.is_file() or args.input.suffix.lower() != ".pptx":
        parser.error("输入必须是存在的 .pptx 文件")
    if args.require_evidence and not args.review:
        parser.error("--require-evidence 须搭配 --review")
    if not 72 <= args.dpi <= 600:
        parser.error("--dpi 须在72至600之间")
    if args.legacy_review and args.require_evidence:
        parser.error("历史导入不能同时称新版 --require-evidence 验收")
    output = args.out.resolve()
    if output == Path.cwd().resolve() or args.input.resolve().is_relative_to(output) or args.out.is_symlink():
        parser.error("输出必须为独立结果目录，不能包含输入PPTX或指向符号链接")
    if args.rerender and output.exists():
        parser.error("--rerender 必须指定新输出目录，以保留已审阅快照")
    from pptx_runtime import create_snapshot, verify_snapshot, create_review_packet
    from pptx_report import compact_report
    try:
        review = json.loads(args.review.read_text(encoding="utf-8")) if args.review else None
        source_sha = hashlib.sha256(args.input.read_bytes()).hexdigest()
        if review is not None:
            if not isinstance(review, dict) or review.get("input_sha256") != source_sha:
                raise ValueError("审阅记录与输入PPTX哈希不一致；未更新结果目录")
            version = review.get("review_schema_version", 1)
            if version < 4 and not args.legacy_review:
                raise ValueError("v1-v3历史记录须显式 --legacy-review；新版审阅使用 v4")
            if version == 4 and args.legacy_review:
                raise ValueError("v4不是历史兼容记录")
            if version == 4 and not (output / "snapshot_manifest.json").is_file():
                raise ValueError("请先不带 --review 准备同一输出目录的快照，再看图填写 v4记录")
            if args.require_evidence and review.get("evidence_review_version") != 1:
                raise ValueError("新版审阅须提供 evidence_review_version: 1 及 target_evidence")
        existing = output.exists()
        if existing:
            if not (output / "inventory.json").is_file():
                raise ValueError("现有输出不是本工具的结果目录，请指定新目录")
            if (output / "snapshot_manifest.json").is_file():
                verify_snapshot(output, source_sha)
            elif not args.legacy_review:
                raise ValueError("现有目录属于历史批次，请为新版读取使用新目录")
            inventory = json.loads((output / "inventory.json").read_text(encoding="utf-8"))
            if inventory.get("sha256") != source_sha:
                raise ValueError("结果目录已绑定另一个输入，请指定新目录")
            inventory["output_root"] = str(output)
            if review is None:
                print(f"Reusing snapshot {inventory.get('snapshot_id', 'legacy')}; {output}")
                return 0 if all(s.get("render", {}).get("status") == "ok" for s in inventory["slides"]) else 2
            validate_review(review, inventory)
        output.parent.mkdir(parents=True, exist_ok=True)
        with staging_directory(output) as temp:
            staging = Path(temp).resolve()
            # Both the temporary and backup paths are validated in the chosen
            # result parent before any recursive copy or directory move.
            assert staging.parent == output.parent and staging != output
            if existing:
                shutil.copytree(output, staging, dirs_exist_ok=True, symlinks=True)
            else:
                inventory = extract(args.input, media_dir=staging / "media")
                inventory["output_root"] = str(output)
                render_issues = render(args.input, staging, inventory["slides"], soffice=args.soffice,
                                       dpi=args.dpi, keep_pdf=not args.no_keep_pdf, render_scope=args.render_scope)
                inventory["render_issues"] = render_issues
                if hashlib.sha256(args.input.read_bytes()).hexdigest() != inventory["sha256"]:
                    raise ValueError("输入在读取过程中变化，未发布批次")
                snapshot = create_snapshot(staging, inventory)
                template = create_review_packet(staging, inventory, snapshot)
                if review is not None:
                    # New legacy imports still validate on actual newly rendered evidence.
                    temp_inventory = dict(inventory, output_root=str(staging))
                    validate_review(review, temp_inventory)
            render_issues = inventory.get("render_issues", [])
            if review:
                diagnoses = sorted(review["slides"], key=lambda x: x["slide"])
                completion = contract.aggregate_completion(diagnoses) if not args.legacy_review else "legacy_unknown"
                mode = "legacy_review" if args.legacy_review else {
                    'complete': 'visual_review', 'partial': 'partial_visual_review',
                    'failed': 'failed_visual_review'}[completion]
                jdump(staging / "accepted_review.json", review)
            else:
                diagnoses = template["slides"]
                for slide, diagnosis in zip(inventory["slides"], diagnoses):
                    diagnosis["uncertainties"] = preliminary(slide)["uncertainties"]
                completion = contract.aggregate_completion(diagnoses)
                mode = "failed_preparation" if completion == "failed" else "preliminary_only"
            result = {"input_sha256": inventory["sha256"], "snapshot_id": inventory.get("snapshot_id"),
                      "generated_at": datetime.now(timezone.utc).isoformat(),
                      "review_schema_version": review.get("review_schema_version", 1) if review else 4,
                      "review_mode": mode, "review_completion": completion,
                      "media_coverage": media_coverage(inventory, diagnoses),
                      "slides": diagnoses, "render_issues": render_issues}
            deck = review.get("deck_assessment") if review else None
            if deck:
                result["deck_assessment"] = deck
            if review and review.get("evidence_review_version"):
                result["evidence_review_version"] = review["evidence_review_version"]
            jdump(staging / "diagnosis.json", result)
            (staging / "report.md").write_text(compact_report(inventory, diagnoses, render_issues, deck, mode=mode), encoding="utf-8")
            (staging / "report-details.md").write_text(markdown(inventory, diagnoses, render_issues, deck), encoding="utf-8")
            backup = output.parent / f"{output.name}.previous-{uuid.uuid4().hex[:10]}"
            assert backup.resolve().parent == output.parent and backup != output
            if existing:
                output.rename(backup)
            try:
                staging.rename(output)
            except OSError:
                if existing:
                    backup.rename(output)
                raise
        print(f"{inventory['slide_count']} slides; {sum(s['render']['status']=='ok' for s in inventory['slides'])} renders; {mode}; {output}")
        return 2 if render_issues or (review and completion in ('partial', 'failed')) else 0
    except (ValueError, OSError, RuntimeError) as error:
        failure = {"state": "failed", "input": str(args.input.resolve()), "output": str(output),
                   "error": str(error), "previous_result_preserved": output.exists(),
                   "occurred_at": datetime.now(timezone.utc).isoformat()}
        if output.parent.exists():
            jdump(output.parent / f"{output.name}.last-failure.json", failure)
        print(f"审阅未完成：{error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

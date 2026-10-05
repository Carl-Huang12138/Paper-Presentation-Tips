"""Read PPTX objects with explicit XML, visibility and geometry provenance.

The inventory describes stored, static content.  It does not claim to interpret
equations, image internals, animation states, or the truth/freshness of chart
data.  Ordinary ``shape-*`` paths remain compatible with existing reviews;
selected alternate branches and inherited shapes have separate ledger paths.
"""

from __future__ import annotations

import hashlib
import copy
import math
import re
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.shapes.shapetree import BaseShapeFactory, SlideShapeFactory
from pptx.text.text import TextFrame

EMU = 914400
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
}
SUPPORTED_MC = set(NS.values())
SHAPES = {"sp", "pic", "grpSp", "graphicFrame", "cxnSp", "contentPart"}
TREE_METADATA = {"nvGrpSpPr", "grpSpPr", "extLst"}
PLOT_TYPES = {"areaChart", "area3DChart", "barChart", "bar3DChart", "lineChart",
              "line3DChart", "pieChart", "pie3DChart", "doughnutChart", "ofPieChart",
              "radarChart", "scatterChart", "bubbleChart", "stockChart", "surfaceChart",
              "surface3DChart"}
IDENTITY = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)


def _xp(node, expression):
    # python-pptx's custom xpath wrapper does not accept a namespaces argument.
    return etree.XPath(expression, namespaces=NS)(node)


def _local(node):
    return etree.QName(node).localname


def _xml_path(node):
    return node.getroottree().getpath(node)


def _bool(value, default=True):
    return default if value is None else value.lower() not in ("0", "false", "off")


def _matrix(a, b):
    """Compose affine matrices, applying b and then a."""
    aa, ab, ac, ad, ae, af = a
    ba, bb, bc, bd, be, bf = b
    return (aa * ba + ac * bb, ab * ba + ad * bb,
            aa * bc + ac * bd, ab * bc + ad * bd,
            aa * be + ac * bf + ae, ab * be + ad * bf + af)


def _translate(x, y):
    return (1.0, 0.0, 0.0, 1.0, x, y)


def _rotate_flip(bounds, angle, flip_h, flip_v):
    x, y, w, h = bounds
    cx, cy = x + w / 2, y + h / 2
    radians = math.radians(angle)
    cosine, sine = math.cos(radians), math.sin(radians)
    rotate = (cosine, sine, -sine, cosine, 0.0, 0.0)
    flip = (-1.0 if flip_h else 1.0, 0.0, 0.0,
            -1.0 if flip_v else 1.0, 0.0, 0.0)
    return _matrix(_translate(cx, cy), _matrix(rotate, _matrix(flip, _translate(-cx, -cy))))


def _point(matrix, x, y):
    a, b, c, d, e, f = matrix
    return a * x + c * y + e, b * x + d * y + f


def _bounds(values):
    if values is None:
        return None
    return dict(zip(("left_in", "top_in", "width_in", "height_in"),
                    (round(value, 3) for value in values)))


def _project(values, matrix):
    if values is None or matrix is None:
        return None, None
    x, y, w, h = values
    corners = [_point(matrix, px, py) for px, py in
               [(x, y), (x + w, y), (x + w, y + h), (x, y + h)]]
    xs, ys = [p[0] for p in corners], [p[1] for p in corners]
    return _bounds((min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys))), corners


def _flatten(items):
    for item in items:
        yield item
        yield from _flatten(item.get("children", []))


def _structure(node):
    return {"tag": _local(node), "namespace": etree.QName(node).namespace,
            "attributes": dict(node.attrib), "text": node.text,
            "children": [_structure(child) for child in node if isinstance(child.tag, str)]}


def _ph(node):
    matches = _xp(node, "./p:nvSpPr/p:nvPr/p:ph | ./p:nvPicPr/p:nvPr/p:ph | "
                        "./p:nvGraphicFramePr/p:nvPr/p:ph")
    if not matches:
        return None
    return {"type": matches[0].get("type", "obj"), "idx": matches[0].get("idx", "0")}


def _nonvisual(node):
    matches = _xp(node, "./p:nvSpPr/p:cNvPr | ./p:nvPicPr/p:cNvPr | "
                        "./p:nvGrpSpPr/p:cNvPr | ./p:nvGraphicFramePr/p:cNvPr | "
                        "./p:nvCxnSpPr/p:cNvPr | ./p:nvContentPartPr/p:cNvPr")
    return matches[0] if matches else None


def _raw_color(container):
    if container is None:
        return None
    for node in container:
        tag = _local(node)
        if tag == "srgbClr":
            return {"type": "rgb", "value": node.get("val")}
        if tag == "schemeClr":
            return {"type": "theme", "value": node.get("val")}
        if tag in ("sysClr", "prstClr", "scrgbClr", "hslClr"):
            return {"type": tag, "value": node.get("lastClr", node.get("val")),
                    "attributes": dict(node.attrib)}
    return None


def _fill(node):
    properties = _xp(node, "./p:spPr | ./p:grpSpPr")
    if properties:
        for child in properties[0]:
            tag = _local(child)
            if tag == "noFill":
                return {"type": "none", "opacity": 0.0, "potentially_visible": False}, None
            if tag in ("solidFill", "gradFill", "blipFill", "pattFill", "grpFill"):
                alphas = _xp(child, ".//a:alpha/@val")
                opacity = float(alphas[0]) / 100000 if alphas else 1.0 if tag == "solidFill" else None
                return {"type": tag, "opacity": opacity, "potentially_visible": opacity != 0.0,
                        "node_path": _xml_path(child)}, _raw_color(child)
    references = _xp(node, "./p:style/a:fillRef")
    if references:
        return {"type": "style_reference", "opacity": None, "potentially_visible": True,
                "style_index": references[0].get("idx"), "needs_theme_resolution": True}, _raw_color(references[0])
    return {"type": "unspecified", "opacity": None, "potentially_visible": False}, None


def _line_color(node):
    matches = _xp(node, "./p:spPr/a:ln/a:solidFill | ./p:style/a:lnRef")
    return _raw_color(matches[0]) if matches else None


def _cache(node, field, numeric, part_name):
    """Read a single field cache, retaining point indices and sparse entries."""
    result = {"values": {}, "count": 0, "status": "missing", "limitations": [],
              "source": {"kind": "chart_xml_cache", "part": part_name, "field": field,
                         "node_path": _xml_path(node) if node is not None else None,
                         "formula": None, "cache_type": None}}
    if node is None:
        result["limitations"].append(f"{field}: data field missing")
        return result
    formulae = _xp(node, ".//c:f/text()")
    result["source"]["formula"] = formulae[0] if formulae else None
    caches = _xp(node, ".//c:numCache | .//c:strCache | ./c:numLit | ./c:strLit | .//c:multiLvlStrCache")
    if not caches:
        result["limitations"].append(f"{field}: XML cache missing; workbook was not substituted")
        return result
    cache = caches[0]
    result["source"].update({"cache_type": _local(cache), "cache_node_path": _xml_path(cache)})
    levels = _xp(cache, "./c:lvl")
    container = levels[0] if levels else cache
    points = _xp(container, "./c:pt")
    for point in points:
        try:
            index = int(point.get("idx"))
            if index < 0 or index in result["values"]:
                raise ValueError("negative or duplicate point index")
            values = _xp(point, "./c:v/text()")
            if not values:
                raise ValueError("point value missing")
            value = float(values[0]) if numeric else str(values[0])
            if numeric and not math.isfinite(value):
                raise ValueError("non-finite value")
            result["values"][index] = value
        except (ValueError, TypeError) as error:
            result["limitations"].append(f"{field}: {_xml_path(point)}: {error}")
    declared = _xp(cache, "./c:ptCount/@val")
    try:
        count = int(declared[0]) if declared else max(result["values"], default=-1) + 1
        if count < 0 or count > 100000:
            raise ValueError("invalid or excessive ptCount")
        result["count"] = max(count, max(result["values"], default=-1) + 1)
    except (ValueError, TypeError) as error:
        result["limitations"].append(f"{field}: {error}")
        result["count"] = max(result["values"], default=-1) + 1
    if levels:
        result["levels"] = [[{"index": p.get("idx"), "label": "".join(_xp(p, "./c:v/text()"))}
                              for p in _xp(level, "./c:pt")] for level in levels]
    if result["count"] == 0:
        result["limitations"].append(f"{field}: empty cache; no data points available")
    if len(result["values"]) != result["count"]:
        result["limitations"].append(f"{field}: sparse or incomplete XML cache")
    result["source"]["point_count"] = result["count"]
    result["status"] = "partial" if result["limitations"] else "complete"
    return result


def _chart_xml(root, part_name):
    result = {"chart_type": "unknown", "title": "unknown", "categories": [], "series": [],
              "plots": [], "axes": [], "value_axis": "unknown", "category_axis": "unknown",
              "legend": None, "limitations": [], "parse_status": "complete",
              "data_provenance": {"kind": "stored_chart_xml_cache", "part": part_name,
                                  "workbook_freshness": "unknown", "external_validity": "unknown"}}
    title_nodes = _xp(root, "./c:chart/c:title")
    if title_nodes:
        titles = _xp(title_nodes[0], ".//a:t/text() | .//c:strCache/c:pt/c:v/text()")
        result["title"] = "".join(titles) if titles else "unknown"
        if not titles:
            result["limitations"].append("chart title present but stored text unavailable")
    legend = _xp(root, "./c:chart/c:legend")
    if legend:
        positions = _xp(legend[0], "./c:legendPos/@val")
        result["legend"] = {"position": positions[0] if positions else "unknown"}
    plot_areas = _xp(root, "./c:chart/c:plotArea")
    plots = [node for area in plot_areas for node in area
             if isinstance(node.tag, str) and _local(node).endswith("Chart")]
    for plot_index, plot in enumerate(plots):
        plot_type = _local(plot)
        kind = "xy" if plot_type == "scatterChart" else "bubble" if plot_type == "bubbleChart" else "category"
        plot_result = {"index": plot_index, "type": plot_type, "data_kind": kind,
                       "node_path": _xml_path(plot), "axis_ids": _xp(plot, "./c:axId/@val"),
                       "series": [], "categories": [], "parse_status": "complete"}
        if plot_type not in PLOT_TYPES:
            result["limitations"].append(f"{plot_type}: unsupported chart plot")
            plot_result["parse_status"] = "partial"
        for ser in _xp(plot, "./c:ser"):
            def child(tag):
                matches = _xp(ser, "./c:" + tag)
                return matches[0] if matches else None
            names = _xp(ser, "./c:tx/c:v/text() | ./c:tx/c:strRef/c:strCache/c:pt/c:v/text()")
            indices = _xp(ser, "./c:idx/@val")
            entry = {"name": "".join(names) if names else "unknown", "values": [], "points": [],
                     "index": int(indices[0]) if indices and indices[0].isdigit() else None,
                     "plot_index": plot_index, "plot_type": plot_type, "data_kind": kind,
                     "node_path": _xml_path(ser), "data_sources": {}, "limitations": [],
                     "parse_status": "complete", "fill_color": None, "line_color": None}
            fills = _xp(ser, "./c:spPr/a:solidFill")
            lines = _xp(ser, "./c:spPr/a:ln/a:solidFill")
            entry["fill_color"] = _raw_color(fills[0]) if fills else None
            entry["line_color"] = _raw_color(lines[0]) if lines else None
            fields = {"x": _cache(child("xVal"), "x", True, part_name),
                      "y": _cache(child("yVal"), "y", True, part_name)} if kind != "category" else {
                          "category": _cache(child("cat"), "category", False, part_name),
                          "value": _cache(child("val"), "value", True, part_name)}
            if kind == "bubble":
                fields["size"] = _cache(child("bubbleSize"), "size", True, part_name)
            count = max((cache["count"] for cache in fields.values()), default=0)
            for field, cache in fields.items():
                entry["data_sources"][field] = cache["source"]
                entry["limitations"].extend(cache["limitations"])
                if cache["count"] != count:
                    entry["limitations"].append(f"{field}: point count differs from other series fields")
            for index in range(count):
                point = {"index": index}
                point.update({field: cache["values"].get(index) for field, cache in fields.items()})
                entry["points"].append(point)
            value_field = "value" if kind == "category" else "y"
            entry["values"] = [fields[value_field]["values"].get(index) for index in range(count)]
            if kind == "category":
                entry["categories"] = [fields["category"]["values"].get(index) for index in range(count)]
                if "levels" in fields["category"]:
                    entry["category_levels"] = fields["category"]["levels"]
                if not plot_result["categories"]:
                    plot_result["categories"] = entry["categories"]
            if entry["limitations"]:
                entry["parse_status"] = "partial"
                plot_result["parse_status"] = "partial"
                result["limitations"].extend(f"plot {plot_index}, series {entry['index']}: {s}"
                                             for s in entry["limitations"])
            plot_result["series"].append(entry)
            result["series"].append(entry)
        if not plot_result["series"]:
            plot_result["parse_status"] = "partial"
            result["limitations"].append(f"plot {plot_index}: no readable series")
        result["plots"].append(plot_result)
    if not plots:
        result["limitations"].append("no supported chart plot found")
    result["chart_type"] = plots[0] is not None and _local(plots[0]) if len(plots) == 1 else "combination" if plots else "unknown"
    if result["plots"]:
        result["categories"] = result["plots"][0]["categories"]
    for axis in _xp(root, "./c:chart/c:plotArea/c:valAx | ./c:chart/c:plotArea/c:catAx | "
                          "./c:chart/c:plotArea/c:dateAx | ./c:chart/c:plotArea/c:serAx"):
        axis_type = _local(axis)
        def attribute(expression):
            values = _xp(axis, expression)
            return values[0] if values else None
        def scale(expression):
            value = attribute(expression)
            if value is None:
                return None
            try:
                return float(value)
            except ValueError:
                result["limitations"].append(f"{_xml_path(axis)}: invalid axis scale {value!r}")
                return None
        record = {"type": axis_type, "id": attribute("./c:axId/@val"),
                  "cross_axis_id": attribute("./c:crossAx/@val"), "node_path": _xml_path(axis),
                  "title": "".join(_xp(axis, "./c:title//a:t/text()")) or None,
                  "minimum_scale": scale("./c:scaling/c:min/@val"),
                  "maximum_scale": scale("./c:scaling/c:max/@val"),
                  "tick_labels": "unknown"}
        result["axes"].append(record)
        legacy = "value_axis" if axis_type == "valAx" else "category_axis"
        if result[legacy] == "unknown":
            result[legacy] = record
    external = _xp(root, "./c:externalData/@r:id")
    if external:
        result["data_provenance"]["workbook_relationship_id"] = external[0]
    if result["limitations"]:
        result["parse_status"] = "partial"
    return result


class _Reader:
    def __init__(self, presentation, media_dir):
        # Lazy import lets analyze_pptx expose this reader through its wrapper.
        from analyze_pptx import text_detail, typography
        self.text_detail = text_detail
        self.typography = typography
        self.presentation = presentation
        self.media_dir = Path(media_dir).resolve() if media_dir is not None else None
        if self.media_dir is not None:
            self.media_dir.mkdir(parents=True, exist_ok=True)
        self.media_manifest = []
        self.slide_number = 0
        self.ac_records = []
        self._ac_by_node = {}
        self.limitations = []
        self.layout = self.master = None

    def limitation(self, item, field, message, error=None):
        item["parse_status"] = "partial"
        if field not in item["missing_fields"]:
            item["missing_fields"].append(field)
        record = {"field": field, "message": message}
        if error is not None:
            record["exception"] = type(error).__name__
            record["detail"] = str(error)
        item.setdefault("parse_errors", []).append(record)
        self.limitations.append(f"{item['path']}: {field}: {message}" + (f" ({type(error).__name__}: {error})" if error else ""))

    def select_alternate(self, node, part_name):
        if node in self._ac_by_node:
            return self._ac_by_node[node]
        branches = []
        selected = None
        for child in node:
            if not isinstance(child.tag, str):
                continue
            required = child.get("Requires", "").split()
            uris = [child.nsmap.get(prefix) for prefix in required]
            supported = _local(child) == "Choice" and bool(required) and all(uri in SUPPORTED_MC for uri in uris)
            branches.append({"branch": _local(child), "requires": required, "namespace_uris": uris,
                             "supported_for_structural_reading": supported, "node_path": _xml_path(child)})
            if selected is None and supported:
                selected = child
        if selected is None:
            selected = next((child for child in node if isinstance(child.tag, str) and _local(child) == "Fallback"), None)
        record = {"part": part_name, "node_path": _xml_path(node), "branches": branches,
                  "selected_branch": _local(selected) if selected is not None else None,
                  "selected_node_path": _xml_path(selected) if selected is not None else None,
                  "parse_status": "selected" if selected is not None else "unreadable",
                  "selection_scope": "stored object structure; render equivalence unverified"}
        self.ac_records.append(record)
        self._ac_by_node[node] = selected, record
        if selected is None:
            self.limitations.append(f"{part_name}:{_xml_path(node)}: AlternateContent has no supported branch or fallback")
        return selected, record

    def active_children(self, tree, part_name, inherited_ac=None):
        for child in tree:
            if not isinstance(child.tag, str):
                continue
            if child.tag == "{%s}AlternateContent" % NS["mc"]:
                selected, record = self.select_alternate(child, part_name)
                if selected is not None:
                    yield from self.active_children(selected, part_name, record)
            elif etree.QName(child).namespace == NS["p"] and _local(child) in SHAPES:
                yield child, inherited_ac
            elif _local(child) not in TREE_METADATA:
                # Preserve unfamiliar drawing nodes as regions instead of
                # silently claiming that the whole shape tree was parsed.
                yield child, inherited_ac

    def active_nodes(self, node, part_name):
        """Walk only the selected MC branches, including math nested in text."""
        yield node
        for child in node:
            if not isinstance(child.tag, str):
                continue
            if child.tag == "{%s}AlternateContent" % NS["mc"]:
                selected, _ = self.select_alternate(child, part_name)
                if selected is not None:
                    yield from self.active_nodes(selected, part_name)
            else:
                yield from self.active_nodes(child, part_name)

    def selected_text_detail(self, frame, part_name):
        """Read selected alternate runs with the usual paragraph/font reader.

        Work on a detached copy; retain original branch provenance and never
        flatten OMML into ordinary text or mutate the source presentation.
        """
        body = frame._txBody
        if not _xp(body, ".//mc:AlternateContent"):
            return self.text_detail(frame)

        sources = []

        def selected_copy(node):
            if node.tag == "{%s}AlternateContent" % NS["mc"]:
                branch, source = self.select_alternate(node, part_name)
                sources.append(source)
                return [copied for child in branch for copied in selected_copy(child)] if branch is not None else []
            copied = copy.deepcopy(node)
            for child in list(copied):
                copied.remove(child)
            for child in node:
                if isinstance(child.tag, str):
                    copied.extend(selected_copy(child))
            return [copied]

        selected_body = selected_copy(body)[0]
        result = self.text_detail(TextFrame(selected_body, frame._parent))
        result["alternate_content_sources"] = sources
        result["alternate_text_basis"] = "selected stored XML branch; renderer confirmation required"
        return result

    def geometry(self, node, kind, part_name):
        sources = [(node, kind, part_name)]
        placeholder = _ph(node)
        if placeholder and kind == "slide":
            for part, source_kind in [(self.layout, "layout"), (self.master, "master")]:
                for candidate in part._element.spTree:
                    match = _ph(candidate) if isinstance(candidate.tag, str) else None
                    if match and ((source_kind == "layout" and match["idx"] == placeholder["idx"]) or
                                  (source_kind == "master" and match["type"] == placeholder["type"])):
                        sources.append((candidate, source_kind, str(part.part.partname)))
                        break
        values = {}
        components = {}
        selected_xfrm = None
        for candidate, source_kind, part_name in sources:
            transforms = _xp(candidate, "./p:spPr/a:xfrm | ./p:grpSpPr/a:xfrm | ./p:xfrm")
            if not transforms:
                continue
            transform = transforms[0]
            if selected_xfrm is None:
                selected_xfrm = transform
            for tag, fields in [("off", ["x", "y"]), ("ext", ["cx", "cy"])]:
                nodes = _xp(transform, "./a:" + tag)
                if nodes:
                    for field in fields:
                        if field not in values and nodes[0].get(field) is not None:
                            values[field] = float(nodes[0].get(field)) / EMU
                            components[field] = {"kind": source_kind, "part": part_name,
                                                 "node_path": _xml_path(nodes[0])}
        if all(field in values for field in ["x", "y", "cx", "cy"]):
            raw = tuple(values[field] for field in ["x", "y", "cx", "cy"])
            if not all(math.isfinite(v) for v in raw) or raw[2] < 0 or raw[3] < 0:
                raise ValueError("invalid shape geometry")
        else:
            raw = None
        source = {"method": "xml_xfrm" if len(sources) == 1 else "xml_placeholder_chain",
                  "components": components, "status": "known" if raw is not None else "unknown"}
        return raw, selected_xfrm, source

    def objects(self, tree, parent, kind, layer, parent_matrix=IDENTITY, ancestors=None,
                parent_visible=True, path_prefix=None, inherited_effective=True):
        part_name = str(parent.part.partname)
        ordinary, supplemental = [], []
        ordinary_index = supplemental_index = unparsed_index = 0
        for stack_order, (node, alternate) in enumerate(self.active_children(tree, part_name), 1):
            supported = etree.QName(node).namespace == NS["p"] and _local(node) in SHAPES
            if not supported:
                unparsed_index += 1
                path = f"{path_prefix}/unparsed-{unparsed_index}" if path_prefix else f"unparsed-{unparsed_index}" if kind == "slide" else f"{kind}/unparsed-{unparsed_index}"
            elif alternate is None:
                ordinary_index += 1
                path = f"{path_prefix}/{ordinary_index}" if path_prefix else f"shape-{ordinary_index}" if kind == "slide" else f"{kind}/shape-{ordinary_index}"
            else:
                supplemental_index += 1
                path = f"{path_prefix}/supplement-{supplemental_index}" if path_prefix else f"supplement-{supplemental_index}" if kind == "slide" else f"{kind}/supplement-{supplemental_index}"
            item = self.object(node, parent, kind, layer, path, stack_order, alternate,
                               parent_matrix, ancestors or [], parent_visible, inherited_effective)
            (ordinary if alternate is None else supplemental).append(item)
        return ordinary + supplemental

    def object(self, node, parent, kind, layer, path, stack_order, alternate,
               parent_matrix, ancestors, parent_visible, inherited_effective):
        nonvisual = _nonvisual(node)
        raw_id = nonvisual.get("id") if nonvisual is not None else None
        raw_name = nonvisual.get("name") if nonvisual is not None else None
        own_visible = not (nonvisual is not None and _bool(nonvisual.get("hidden"), False))
        placeholder = _ph(node)
        inherited_placeholder = kind in ("layout", "master") and placeholder is not None
        effective = own_visible and parent_visible and inherited_effective and not inherited_placeholder
        item = {"id": int(raw_id) if raw_id and raw_id.isdigit() else None, "path": path,
                "name": raw_name, "type": _local(node), "source": "pptx_object" if alternate is None and kind == "slide" else "xml_supplement" if alternate else "inherited_pptx_object",
                "z_order": stack_order, "paint_order": [layer] + [a["stack_order"] for a in ancestors] + [stack_order],
                "visible_in_slideshow": effective, "parse_status": "complete", "missing_fields": [],
                "xml_source": {"kind": kind, "part": str(parent.part.partname), "id": raw_id,
                               "name": raw_name, "node_path": _xml_path(node), "tag": _local(node)},
                "visibility": {"self_visible": own_visible, "ancestors_visible": parent_visible,
                               "effective_visible": effective, "ancestor_paths": [a["path"] for a in ancestors],
                               "animation_state": "unknown"},
                "placeholder": placeholder}
        if alternate:
            item["xml_source"]["alternate_content"] = alternate
        if kind != "slide":
            item["inheritance"] = {"effective": effective, "placeholder_definition_only": inherited_placeholder,
                                   "conditions": {"inherited_shapes_enabled": inherited_effective,
                                                  "self_visible": own_visible, "ancestors_visible": parent_visible},
                                   "reason": "placeholder definitions supply geometry/styles to matching slide placeholders" if inherited_placeholder else "static inherited drawing"}
        if raw_id is None:
            self.limitation(item, "id", "XML object has no reliable nonvisual id")
        if etree.QName(node).namespace != NS["p"] or _local(node) not in SHAPES:
            item["unparsed_xml"] = etree.tostring(node, encoding="unicode")
            self.limitation(item, "drawing_node", "unrecognized drawing node; content and visibility require renderer inspection")
        try:
            shape = SlideShapeFactory(node, parent) if kind == "slide" else BaseShapeFactory(node, parent)
            item["type"] = str(shape.shape_type)
        except Exception as error:
            shape = None
            self.limitation(item, "object_api", "object proxy could not be constructed", error)
        try:
            raw_bounds, xfrm, geometry_source = self.geometry(node, kind, str(parent.part.partname))
            angle = float(xfrm.get("rot", "0")) / 60000 if xfrm is not None else 0.0
            flip_h = _bool(xfrm.get("flipH"), False) if xfrm is not None else False
            flip_v = _bool(xfrm.get("flipV"), False) if xfrm is not None else False
            own_matrix = _rotate_flip(raw_bounds, angle, flip_h, flip_v) if raw_bounds is not None else None
            composed = _matrix(parent_matrix, own_matrix) if parent_matrix is not None and own_matrix is not None else None
            slide_bounds, corners = _project(raw_bounds, composed)
            item.update({"local_bounds": _bounds(raw_bounds), "slide_bounds": slide_bounds, "bounds": slide_bounds,
                         "slide_corners_in": [[round(v, 6) for v in p] for p in corners] if corners else None,
                         "geometry_source": geometry_source,
                         "transform": {"status": "known" if slide_bounds is not None else "unknown",
                                       "matrix": list(composed) if composed is not None else None,
                                       "rotation_degrees": angle, "flip_h": flip_h, "flip_v": flip_v,
                                       "ancestor_transforms": [a["transform"] for a in ancestors],
                                       "bounds_kind": "axis_aligned_slide_projection"}})
            if slide_bounds is None and effective:
                self.limitation(item, "slide_bounds", "no reliable page geometry; no local bounds substituted")
        except Exception as error:
            raw_bounds = xfrm = None
            item.update({"local_bounds": None, "slide_bounds": None, "bounds": None,
                         "geometry_source": {"status": "unknown"}, "transform": {"status": "unknown"}})
            self.limitation(item, "geometry", "geometry or transform could not be read", error)
        try:
            item["fill"], item["fill_color"] = _fill(node)
            item["line_color"] = _line_color(node)
        except Exception as error:
            item["fill"] = {"type": "unknown", "potentially_visible": False}
            item["fill_color"] = item["line_color"] = None
            self.limitation(item, "fill", "fill metadata could not be read", error)
        if shape is not None:
            try:
                if shape.has_text_frame:
                    item["text_frame"] = self.selected_text_detail(shape.text_frame, str(parent.part.partname))
                    item["text_frame"]["source"] = {"part": str(parent.part.partname), "node_path": _xml_path(node)}
            except Exception as error:
                self.limitation(item, "text_frame", "text frame could not be read", error)
            try:
                if shape.has_table:
                    table = shape.table
                    rows = []
                    for row_index, row in enumerate(table.rows, 1):
                        cells = []
                        for column_index, cell in enumerate(row.cells, 1):
                            tc = cell._tc
                            fills = _xp(tc, "./a:tcPr/a:solidFill")
                            cell_frame = self.selected_text_detail(cell.text_frame, str(parent.part.partname))
                            cells.append({"text": cell_frame["text"], "text_frame": cell_frame,
                                          "fill_color": _raw_color(fills[0]) if fills else None,
                                          "row": row_index, "column": column_index,
                                          "selector": f"{path}:cell-{row_index}-{column_index}",
                                          "is_merge_origin": cell.is_merge_origin, "is_spanned": cell.is_spanned,
                                          "span_width": cell.span_width, "span_height": cell.span_height,
                                          "xml_source": {"part": str(parent.part.partname), "node_path": _xml_path(tc)}})
                        rows.append(cells)
                    item["table"] = {"rows": rows, "row_count": len(table.rows), "column_count": len(table.columns)}
            except Exception as error:
                self.limitation(item, "table", "table content could not be fully read", error)
        for graphic in _xp(node, "./a:graphic/a:graphicData"):
            uri = graphic.get("uri")
            if uri not in ("http://schemas.openxmlformats.org/drawingml/2006/table", NS["c"]):
                item["graphic_data"] = {"uri": uri, "node_path": _xml_path(graphic),
                                        "raw_xml": etree.tostring(graphic, encoding="unicode"),
                                        "semantic_status": "unknown"}
                self.limitation(item, "graphic_data", "SmartArt or other graphicData has no reliable semantic reader")
        charts = _xp(node, "./a:graphic/a:graphicData/c:chart")
        if charts:
            try:
                relationship_id = charts[0].get("{%s}id" % NS["r"])
                chart_part = parent.part.related_part(relationship_id)
                chart_root = etree.fromstring(chart_part.blob)
                item["chart"] = _chart_xml(chart_root, str(chart_part.partname))
                item["chart"]["relationship_id"] = relationship_id
                external_id = item["chart"]["data_provenance"].get("workbook_relationship_id")
                if external_id:
                    relationship = chart_part.rels[external_id]
                    item["chart"]["data_provenance"]["workbook_target"] = relationship.target_ref
                    item["chart"]["data_provenance"]["workbook_external"] = relationship.is_external
                if item["chart"]["parse_status"] == "partial":
                    self.limitation(item, "chart_data", "; ".join(item["chart"]["limitations"]))
            except Exception as error:
                item["chart"] = {"chart_type": "unknown", "title": "unknown", "categories": "unknown",
                                 "series": [], "plots": [], "value_axis": "unknown", "category_axis": "unknown",
                                 "legend": None, "parse_status": "partial", "limitations": [str(error)]}
                self.limitation(item, "chart", "chart part or data could not be read", error)
        if _local(node) == "pic":
            self.picture(item, node, parent, shape)
        ole = _xp(node, ".//p:oleObj")
        if ole:
            item["embedded_object"] = {"meaning": "unknown", "render_only": True,
                                       "relationships": [o.get("{%s}id" % NS["r"]) for o in ole],
                                       "raw_xml": [etree.tostring(o, encoding="unicode") for o in ole]}
            self.limitation(item, "embedded_object_semantics", "embedded OLE object content is unknown")
        math_nodes = [n for n in self.active_nodes(node, str(parent.part.partname))
                      if etree.QName(n).namespace == NS["m"] and _local(n) in ("oMath", "oMathPara")
                      and not any(etree.QName(a).namespace == NS["m"] and _local(a) in ("oMath", "oMathPara")
                                  for a in n.iterancestors())]
        # A group's descendants are separate ledger objects; do not count their
        # formulas again on the group container.
        if _local(node) != "grpSp" and math_nodes:
            item["math"] = [{"part": str(parent.part.partname), "node_path": _xml_path(m),
                             "raw_xml": etree.tostring(m, encoding="unicode"), "structure": _structure(m),
                             "semantic_status": "unknown", "parse_status": "structure_preserved",
                             "expression": None} for m in math_nodes]
            if "text_frame" in item:
                item["text_frame"]["contains_math"] = True
                item["text_frame"]["math_semantics"] = "unknown"
            self.limitation(item, "formula_semantics", "OMML structure preserved; no mathematical expression inferred")
        if _local(node) == "contentPart":
            item["unparsed_xml"] = etree.tostring(node, encoding="unicode")
            self.limitation(item, "contentPart", "unsupported contentPart drawing")
        if _local(node) == "grpSp":
            child_matrix = None
            try:
                if xfrm is None or raw_bounds is None or parent_matrix is None:
                    raise ValueError("group transform or ancestor transform missing")
                offsets = _xp(xfrm, "./a:chOff")
                extents = _xp(xfrm, "./a:chExt")
                if not offsets or not extents:
                    raise ValueError("group chOff/chExt missing")
                ox, oy = (float(offsets[0].get(field)) / EMU for field in ("x", "y"))
                ew, eh = (float(extents[0].get(field)) / EMU for field in ("cx", "cy"))
                if ew <= 0 or eh <= 0 or not all(math.isfinite(v) for v in (ox, oy, ew, eh)):
                    raise ValueError("group chExt must be positive finite values")
                x, y, w, h = raw_bounds
                scale = (w / ew, 0.0, 0.0, h / eh, 0.0, 0.0)
                child_matrix = _matrix(parent_matrix, _matrix(own_matrix,
                                       _matrix(_translate(x, y), _matrix(scale, _translate(-ox, -oy)))))
                item["transform"]["child_coordinate_matrix"] = list(child_matrix)
            except Exception as error:
                self.limitation(item, "group_transform", "child page coordinates are unknown", error)
                item["transform"]["child_coordinate_matrix"] = None
            ancestor = {"path": path, "stack_order": stack_order,
                        "transform": {"path": path, "matrix": list(child_matrix) if child_matrix else None,
                                      "status": "known" if child_matrix else "unknown"}}
            item["children"] = self.objects(node, parent, kind, layer, child_matrix, ancestors + [ancestor],
                                            effective, path, inherited_effective)
        return item

    def picture(self, item, node, parent, shape):
        picture = {"text_or_data_inside_image": "unknown", "extraction_status": "not_structurally_parsed",
                   "instance_id": f"slide-{self.slide_number}:{item['path']}", "original_media": None,
                   "media_path": None, "export_path": None, "crop": {side: 0.0 for side in ("left", "top", "right", "bottom")},
                   "rotation_degrees": item.get("transform", {}).get("rotation_degrees"),
                   "content_type": None, "filename": None, "sha256": None, "pixel_size": None,
                   "media_part": None, "relationship_id": None}
        item["picture"] = picture
        crop = _xp(node, "./p:blipFill/a:srcRect")
        if crop:
            picture["crop"] = {side: float(crop[0].get(short, "0")) / 100000
                               for side, short in [("left", "l"), ("top", "t"), ("right", "r"), ("bottom", "b")]}
        blips = _xp(node, "./p:blipFill/a:blip")
        try:
            if not blips:
                raise ValueError("picture has no blip relationship")
            relationship_id = blips[0].get("{%s}embed" % NS["r"])
            linked_id = blips[0].get("{%s}link" % NS["r"])
            picture["relationship_id"] = relationship_id or linked_id
            relationship = parent.part.rels[relationship_id or linked_id]
            if relationship.is_external:
                picture["linked_target"] = relationship.target_ref
                raise ValueError("linked external media is not present in the PPTX package")
            media = parent.part.related_part(relationship_id or linked_id)
            blob = media.blob
            picture.update({"content_type": media.content_type, "media_part": str(media.partname),
                            "filename": Path(str(media.partname)).name, "sha256": hashlib.sha256(blob).hexdigest()})
            try:
                from pptx.parts.image import Image
                picture["pixel_size"] = list(Image.from_blob(blob).size)
            except Exception as error:
                self.limitation(item, "image_pixel_size", "pixel dimensions could not be read (vector/unsupported media possible)", error)
            if self.media_dir is not None:
                extension = Path(str(media.partname)).suffix
                if not re.fullmatch(r"\.[A-Za-z0-9]{1,10}", extension):
                    extension = ".bin"
                instance = re.sub(r"[^A-Za-z0-9_-]", "-", item["path"])
                destination = self.media_dir / f"slide-{self.slide_number:03d}-{instance}-{picture['sha256'][:12]}{extension}"
                destination.resolve().relative_to(self.media_dir)
                destination.write_bytes(blob)
                picture.update({"original_media": str(destination), "media_path": str(destination),
                                "export_path": str(destination), "export_relative_path": destination.relative_to(self.media_dir).as_posix()})
                self.media_manifest.append({"slide": self.slide_number, "object": item["path"],
                                            "instance_id": picture["instance_id"], "media_part": picture["media_part"],
                                            "sha256": picture["sha256"], "path": str(destination),
                                            "bytes": len(blob), "content_type": picture["content_type"]})
        except Exception as error:
            self.limitation(item, "picture_media", "original picture media could not be read or exported", error)
        self.limitations.append(f"{item['path']}: 图片内文字/数据未结构化；需核对原图与放映画面的裁切和遮挡")

    def slide(self, slide, number):
        self.slide_number = number
        self.ac_records, self._ac_by_node, self.limitations = [], {}, []
        self.layout = slide.slide_layout
        self.master = self.layout.slide_master
        hidden = not _bool(slide._element.get("show"))
        inherited_enabled = _bool(slide._element.get("showMasterSp"))
        master_enabled = inherited_enabled and _bool(self.layout._element.get("showMasterSp"))
        objects = self.objects(slide._element.spTree, slide.shapes, "slide", 0)
        objects += self.objects(self.layout._element.spTree, self.layout.shapes, "layout", -1,
                                inherited_effective=inherited_enabled)
        objects += self.objects(self.master._element.spTree, self.master.shapes, "master", -2,
                                inherited_effective=master_enabled)
        # Record top-level transitions and other MC constructs too, while only
        # selected object branches contribute content to the ledger.
        for root, owner in [(slide._element, slide.part), (self.layout._element, self.layout.part),
                            (self.master._element, self.master.part)]:
            list(self.active_nodes(root, str(owner.partname)))
        notes = None
        if slide.has_notes_slide:
            try:
                notes = slide.notes_slide.notes_text_frame.text
            except Exception as error:
                self.limitations.append(f"notes: read failed ({type(error).__name__}: {error})")
        texts = []
        for item in _flatten(objects):
            if not item["visible_in_slideshow"]:
                continue
            text = item.get("text_frame", {}).get("text", "")
            if text.strip():
                texts.append(text)
            if "table" in item:
                texts.extend(cell["text"] for row in item["table"]["rows"] for cell in row
                             if cell["text"].strip() and not cell.get("is_spanned"))
            chart = item.get("chart")
            if chart:
                if chart["title"] and chart["title"] != "unknown":
                    texts.append(chart["title"])
                if chart["legend"]:
                    texts.extend(s["name"] for s in chart["series"] if s["name"] != "unknown")
        source_chain = {"slide": str(slide.part.partname), "layout": str(self.layout.part.partname),
                        "master": str(self.master.part.partname), "theme": None,
                        "conditions": {"slide_show_master_shapes": inherited_enabled,
                                       "layout_show_master_shapes": _bool(self.layout._element.get("showMasterSp")),
                                       "placeholder_definitions_are_not_visible_text": True},
                        "style_resolution": "explicit styles preserved; complete theme/font inheritance unknown"}
        themes = [rel.target_ref for rel in self.master.part.rels.values() if rel.reltype.endswith("/theme")]
        if themes:
            source_chain["theme"] = themes[0]
        if _xp(slide._element, "./p:timing"):
            self.limitations.append("animation: static inventory does not determine all time-dependent visibility states")
        result = {"slide": number, "slide_id": slide.slide_id, "slide_part": str(slide.part.partname),
                  "size_in": [round(self.presentation.slide_width / EMU, 3), round(self.presentation.slide_height / EMU, 3)],
                  "hidden_in_slideshow": hidden, "visible_text": texts, "notes_not_visible": notes,
                  "objects": objects, "parse_limitations": list(dict.fromkeys(self.limitations)),
                  "alternate_content": self.ac_records, "source_chain": source_chain,
                  "visibility": {"hidden_slide": hidden, "objects_on_hidden_slide_retained": True,
                                 "static_object_visibility": "self and ancestor hidden flags applied",
                                 "notes_are_slideshow_content": False, "animation_states": "unknown"},
                  "typography": self.typography(objects)}
        result["parse_status"] = "partial" if result["parse_limitations"] else "complete"
        result["occlusion_cues"] = _occlusions(objects)
        return result


def _occlusions(objects):
    visible = [item for item in _flatten(objects) if item.get("visible_in_slideshow") and item.get("slide_bounds")]
    cues = []
    for picture in (o for o in visible if "picture" in o):
        p = picture["slide_bounds"]
        if p["width_in"] <= 0 or p["height_in"] <= 0:
            continue
        for foreground in visible:
            if foreground["path"] == picture["path"] or foreground.get("paint_order", []) <= picture.get("paint_order", []):
                continue
            if not foreground.get("fill", {}).get("potentially_visible") or foreground.get("children"):
                continue
            f = foreground["slide_bounds"]
            left, top = max(p["left_in"], f["left_in"]), max(p["top_in"], f["top_in"])
            right = min(p["left_in"] + p["width_in"], f["left_in"] + f["width_in"])
            bottom = min(p["top_in"] + p["height_in"], f["top_in"] + f["height_in"])
            if right <= left or bottom <= top:
                continue
            cues.append({"picture": picture["path"], "foreground": foreground["path"],
                         "intersection_bounds": _bounds((left, top, right - left, bottom - top)),
                         "picture_area_fraction": round((right - left) * (bottom - top) / (p["width_in"] * p["height_in"]), 6),
                         "foreground_fill": foreground["fill"], "status": "needs_visual_confirmation",
                         "confirmed_issue": False, "geometry_basis": "axis_aligned page bounds",
                         "limitation": "fill theme, geometry, transparency and presentation intent require rendered-slide inspection"})
    return cues


def _titles(slides):
    frequency = Counter()
    for slide in slides:
        frequency.update(set(o.get("text_frame", {}).get("text", "").strip()
                             for o in _flatten(slide["objects"]) if o.get("visible_in_slideshow")
                             and o.get("text_frame", {}).get("text", "").strip()))
    for slide in slides:
        candidates = []
        width, height = slide["size_in"]
        for item in _flatten(slide["objects"]):
            text = item.get("text_frame", {}).get("text", "").strip()
            if not item.get("visible_in_slideshow") or not text:
                continue
            placeholder = item.get("placeholder") or {}
            role = placeholder.get("type")
            if role in ("dt", "sldNum", "ftr", "hdr"):
                continue
            confirmed = role in ("title", "ctrTitle") and item["xml_source"]["kind"] == "slide"
            bounds = item.get("slide_bounds")
            reasons, score = [], 0
            if confirmed:
                reasons.append("title_placeholder_role")
                score = 100
            if bounds and bounds["top_in"] < min(1.25, height * .22) and bounds["height_in"] < height * .25:
                reasons.append("title_region_geometry")
                score += 3
                if bounds["left_in"] < width * .55:
                    reasons.append("left_title_region")
                    score += 2
            if len(text) <= 120:
                score += 1
            if len(text.splitlines()) <= 3:
                score += 1
            if frequency[text] > 2 and not confirmed:
                reasons.append("repeated_across_slides")
                # A recurring section heading can still be the slide title;
                # repetition alone must not promote an equation number over it.
                score -= 2 if "left_title_region" in reasons else 6
            if not confirmed and re.fullmatch(r"[（(]?\d+[A-Za-z]?[)）]?", text):
                reasons.append("number_only_label_is_not_a_reliable_title")
                score -= 5
            if not confirmed and (len(text) > 120 or len(text.splitlines()) > 3):
                score -= 5
            confidence = "high" if confirmed else "medium" if score >= 4 else "low"
            candidates.append({"text": text, "object": item["path"], "score": score,
                               "confidence": confidence, "status": "confirmed" if confirmed else "candidate",
                               "kind": "title_placeholder" if confirmed else "title_region_candidate" if score >= 4 else "text_candidate",
                               "reasons": reasons or ["first_text_is_only_a_candidate"],
                               "source": item["xml_source"]})
        candidates.sort(key=lambda c: c["score"], reverse=True)
        selected = candidates[0] if candidates else None
        slide["title_candidates"] = candidates
        slide["title"] = selected["text"] if selected else None
        slide["title_status"] = selected["status"] if selected else "unknown"
        slide["title_confidence"] = selected["confidence"] if selected else "unknown"
        slide["title_source"] = ({"object": selected["object"], "kind": selected["kind"],
                                  "part": selected["source"]["part"], "node_path": selected["source"]["node_path"],
                                  "render_confirmation": "not_performed"} if selected else None)


def extract(pptx_path: Path, media_dir: Path | None = None) -> dict:
    """Extract source-bound slide objects, optionally exporting each image instance."""
    pptx_path = Path(pptx_path).resolve()
    try:
        blob = pptx_path.read_bytes()
        presentation = Presentation(str(pptx_path))
    except Exception as error:
        raise ValueError(f"Cannot read PPTX {pptx_path}: {type(error).__name__}: {error}") from error
    try:
        reader = _Reader(presentation, media_dir)
        slides = [reader.slide(slide, index) for index, slide in enumerate(presentation.slides, 1)]
        _titles(slides)
    except Exception as error:
        raise ValueError(f"PPTX extraction failed for {pptx_path}: {type(error).__name__}: {error}") from error
    return {"input": str(pptx_path), "sha256": hashlib.sha256(blob).hexdigest(),
            "slide_count": len(slides), "slides": slides, "reader_schema_version": 2, "parser_version": "2.0",
            "media_manifest": reader.media_manifest,
            "source_separation": "notes and media bytes are separate from slide-visible content; hidden slides retained with flags",
            "limitations": ["Stored static XML is not a renderer, image reader, equation interpreter, or scientific-data validation.",
                            "Theme/font inheritance and all animation states require rendered-slide confirmation."]}

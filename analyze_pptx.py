#!/usr/bin/env python3
"""Evidence-first, read-only PPTX inventory, rendering, and review report."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import tempfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

fitz.TOOLS.mupdf_display_errors(False)


EMU_PER_INCH = 914400
TIP_NAMES = {
    1: "每页字数不要太多", 2: "一页只承担一个核心任务", 3: "标题要有意义",
    4: "颜色不要太多太乱", 5: "注意排版细节", 6: "用案例引出问题",
    7: "先讲例子，再讲公式", 8: "精炼阐述创新点", 9: "让听众一眼看到重点",
    10: "实验部分多用图，少用表格", 11: "总结页回顾核心内容",
}


def jdump(path: Path, obj: object) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def safe(fn, default=None):
    try:
        return fn()
    except Exception:
        return default


def pos(shape) -> dict:
    return {key + "_in": round(getattr(shape, key) / EMU_PER_INCH, 3)
            for key in ("left", "top", "width", "height")}


def color(color_format):
    if color_format is None:
        return None
    typ = safe(lambda: str(color_format.type))
    if not typ:
        return None
    if "RGB" in typ:
        rgb = safe(lambda: color_format.rgb)
        return {"type": "rgb", "value": str(rgb) if rgb else None}
    if "SCHEME" in typ:
        theme = safe(lambda: color_format.theme_color)
        return {"type": "theme", "value": str(theme) if theme else None}
    return {"type": typ, "value": None}


def text_detail(frame) -> dict:
    paragraphs = []
    for paragraph in frame.paragraphs:
        runs = []
        for run in paragraph.runs:
            runs.append({
                "text": run.text,
                "font": {"name": run.font.name,
                         "size_pt": round(run.font.size.pt, 2) if run.font.size else None,
                         "bold": run.font.bold, "italic": run.font.italic,
                         "color": color(run.font.color)},
            })
        paragraphs.append({"text": paragraph.text, "level": paragraph.level,
                           "alignment": str(paragraph.alignment) if paragraph.alignment else None,
                           "runs": runs})
    return {"text": frame.text, "paragraphs": paragraphs,
            "margins_in": {side: round(getattr(frame, "margin_" + side) / EMU_PER_INCH, 3)
                           for side in ("left", "right", "top", "bottom")},
            "word_wrap": frame.word_wrap}


def chart_detail(chart) -> dict:
    result = {"chart_type": str(safe(lambda: chart.chart_type, "unknown")),
              "title": "unknown", "categories": "unknown", "series": [],
              "value_axis": "unknown", "category_axis": "unknown", "legend": None,
              "limitations": []}
    if safe(lambda: chart.has_title, False):
        result["title"] = safe(lambda: chart.chart_title.text_frame.text, "unknown")
    if safe(lambda: chart.has_legend, False):
        result["legend"] = {"position": str(safe(lambda: chart.legend.position, "unknown"))}
    result["categories"] = safe(lambda: [str(c.label) for c in chart.plots[0].categories], "unknown")
    for series in safe(lambda: list(chart.series), []):
        values = safe(lambda: list(series.values), "unknown")
        result["series"].append({"name": safe(lambda: series.name, "unknown"), "values": values,
                                 "fill_color": safe(lambda: color(series.format.fill.fore_color)),
                                 "line_color": safe(lambda: color(series.format.line.color))})
    for axis_name in ("value_axis", "category_axis"):
        axis = safe(lambda: getattr(chart, axis_name))
        if axis:
            result[axis_name] = {
                "title": safe(lambda: axis.axis_title.text_frame.text, None)
                if safe(lambda: axis.has_title, False) else None,
                "minimum_scale": safe(lambda: axis.minimum_scale),
                "maximum_scale": safe(lambda: axis.maximum_scale),
                "tick_labels": "unknown",  # Tick labels are renderer-generated, not OOXML data.
            }
    if result["categories"] == "unknown" or any(s["values"] == "unknown" for s in result["series"]):
        result["limitations"].append("图表缓存数据不可读取；不得推测数值或坐标轴标签")
    return result


def shape_detail(shape, path: str) -> dict:
    item = {"id": shape.shape_id, "path": path, "name": shape.name,
            "type": str(shape.shape_type), "bounds": pos(shape),
            "source": "pptx_object", "z_order": int(path.split("/")[0].split("-")[-1]),
            "visible_in_slideshow": True}
    hidden = safe(lambda: shape._element.xpath(".//p:cNvPr")[0].get("hidden"))
    if hidden in ("1", "true"):
        item["visible_in_slideshow"] = False
    item["fill_color"] = safe(lambda: color(shape.fill.fore_color))
    item["line_color"] = safe(lambda: color(shape.line.color))
    if shape.has_text_frame:
        item["text_frame"] = text_detail(shape.text_frame)
    if shape.has_table:
        table = shape.table
        item["table"] = {"rows": [[{"text": cell.text, "text_frame": text_detail(cell.text_frame),
                                     "fill_color": safe(lambda: color(cell.fill.fore_color))}
                                    for cell in row.cells] for row in table.rows],
                         "row_count": len(table.rows), "column_count": len(table.columns)}
    if shape.has_chart:
        item["chart"] = chart_detail(shape.chart)
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        item["picture"] = {"content_type": safe(lambda: shape.image.content_type),
                           "filename": safe(lambda: shape.image.filename),
                           "sha256": safe(lambda: hashlib.sha256(shape.image.blob).hexdigest()),
                           "text_or_data_inside_image": "unknown"}
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        item["children"] = [shape_detail(child, f"{path}/{i}")
                            for i, child in enumerate(shape.shapes, 1)]
    if "OLE" in str(shape.shape_type):
        item["embedded_object"] = {"meaning": "unknown", "render_only": True}
    return item


def flatten(items):
    for item in items:
        yield item
        yield from flatten(item.get("children", []))


def extract(pptx_path: Path) -> dict:
    pres = Presentation(str(pptx_path))
    slides = []
    for idx, slide in enumerate(pres.slides, 1):
        objects = [shape_detail(shape, f"shape-{i}") for i, shape in enumerate(slide.shapes, 1)]
        slide_xml = slide._element
        hidden = slide_xml.get("show") in ("0", "false")
        notes = safe(lambda: slide.notes_slide.notes_text_frame.text, None) if slide.has_notes_slide else None
        texts = []
        for obj in flatten(objects):
            if not obj.get("visible_in_slideshow"):
                continue
            if obj.get("text_frame", {}).get("text", "").strip():
                texts.append(obj["text_frame"]["text"])
            if "table" in obj:
                texts.extend(cell["text"] for row in obj["table"]["rows"] for cell in row if cell["text"].strip())
            if "chart" in obj:
                chart = obj["chart"]
                if chart["title"] != "unknown" and chart["title"]:
                    texts.append(chart["title"])
                if chart["legend"]:
                    texts.extend(series["name"] for series in chart["series"] if series["name"] != "unknown")
        title = safe(lambda: slide.shapes.title.text, None)
        if not title:
            title = texts[0] if texts else None
        blind = []
        for o in flatten(objects):
            if "picture" in o:
                blind.append(f"{o['path']}: 图片内部文字、图形或数据未知")
            if "embedded_object" in o:
                blind.append(f"{o['path']}: 嵌入对象内容未知")
        slides.append({"slide": idx, "size_in": [round(pres.slide_width / EMU_PER_INCH, 3),
                                                     round(pres.slide_height / EMU_PER_INCH, 3)],
                       "hidden_in_slideshow": hidden, "title": title,
                       "visible_text": texts, "notes_not_visible": notes,
                       "objects": objects, "parse_limitations": blind})
    return {"input": str(pptx_path.resolve()), "sha256": hashlib.sha256(pptx_path.read_bytes()).hexdigest(),
            "slide_count": len(slides), "slides": slides,
            "source_separation": "notes and internal media metadata are not slideshow content"}


def render(pptx_path: Path, out_dir: Path, slides: list[dict]) -> list[str]:
    issues = []
    render_dir = out_dir / "renders"
    render_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="pptx_audit_") as temp:
        temp_path = Path(temp)
        profile = (temp_path / "lo_profile").as_uri()
        cmd = ["C:/Program Files/LibreOffice/program/soffice.com", f"-env:UserInstallation={profile}",
               "--headless", "--convert-to", "pdf", "--outdir", str(temp_path), str(pptx_path.resolve())]
        try:
            run = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
            pdf = temp_path / (pptx_path.stem + ".pdf")
            if run.returncode or not pdf.exists():
                raise RuntimeError((run.stdout + run.stderr).strip() or "PDF was not created")
            doc = fitz.open(pdf)
            if len(doc) != len(slides):
                issues.append(f"渲染页数 {len(doc)} 与 PPTX 页数 {len(slides)} 不一致；页码对应关系待人工确认")
            for i, slide in enumerate(slides):
                if i >= len(doc):
                    slide["render"] = {"status": "failed", "reason": "missing PDF page"}
                    continue
                page = doc[i]
                image_path = render_dir / f"slide-{i+1:02d}.png"
                page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False).save(image_path)
                pdf_text = page.get_text("text")
                structural = "".join(slide["visible_text"])
                # A mismatch is a review prompt, not evidence that one source is wrong.
                tokens = [t for t in re.split(r"[\s，。,:;：；（）()]+", structural) if len(t) >= 3]
                missing = [t for t in tokens if t not in pdf_text]
                slide["render"] = {"status": "ok", "image": str(image_path.resolve()),
                                   "pdf_text": pdf_text, "structure_tokens_absent_in_pdf": missing[:30],
                                   "comparison_status": "needs_visual_review" if missing else "text_matches_or_not_testable",
                                   "source": "libreoffice_render_pdf"}
            doc.close()
        except Exception as exc:
            issues.append(f"渲染失败：{exc}")
            for slide in slides:
                slide["render"] = {"status": "failed", "reason": str(exc)}
    return issues


def role(slide: dict) -> str:
    title = slide.get("title") or ""
    if re.search(r"总结|结论|致谢|结束|Q\s*&\s*A", title, re.I): return "总结/收尾"
    if re.search(r"实验|结果|性能|消融|评估", title): return "实验/证据"
    if re.search(r"方法|模型|创新|GRPO|ViT|Q-learning", title, re.I): return "方法/概念"
    if re.search(r"背景|动机|问题|任务定义|推理", title): return "背景/动机"
    return "待人工判定"


def preliminary(slide: dict) -> dict:
    # Mechanical cues do not become confirmed quality defects without a visual review.
    objects = list(flatten(slide["objects"]))
    text_count = sum(len(x) for x in slide["visible_text"])
    charts = sum("chart" in o for o in objects)
    tables = sum("table" in o for o in objects)
    pictures = sum("picture" in o for o in objects)
    cues = []
    if text_count > 450:
        cues.append({"kind": "dense_text", "evidence": f"可见对象文本约 {text_count} 字符",
                     "action": "核对渲染图中的字号、行距及听众能否先抓住核心句；字多本身不是问题"})
    if charts or tables:
        cues.append({"kind": "data_display", "evidence": f"{tables} 个表格、{charts} 个图表",
                     "action": "检查行列、轴、图例、数据及结论的对应关系；表格本身不是问题"})
    if pictures:
        cues.append({"kind": "picture_unknown", "evidence": f"{pictures} 张独立图片",
                     "action": "看渲染图并记录图片内可读元素；未确认的文字和数值保持未知"})
    if slide["render"]["status"] != "ok":
        cues.append({"kind": "render_failed", "evidence": slide["render"]["reason"],
                     "action": "该页视觉判断不可完成，需重新渲染或人工打开"})
    return {"slide": slide["slide"], "role": role(slide), "status": "待视觉审阅",
            "basis": "已完成 PPTX 对象读取；需要把渲染图与对象清单逐项核对",
            "issues": [], "uncertainties": cues}


def validate_review(review: dict, inventory: dict) -> None:
    if review.get("input_sha256") != inventory["sha256"]:
        raise ValueError("审阅记录与输入 PPTX 的 SHA-256 不一致")
    entries = review.get("slides", [])
    nums = [x.get("slide") for x in entries]
    if sorted(nums) != list(range(1, inventory["slide_count"] + 1)):
        raise ValueError("review.slides 必须覆盖每一页且页码不能重复")
    valid_status = {"无需修改", "有确定问题", "需要人工确认"}
    required_issue = {"object", "evidence", "impact", "severity", "confidence", "tips", "direction", "category"}
    for entry in entries:
        image = inventory["slides"][entry["slide"] - 1]["render"].get("image")
        if image and entry.get("render_sha256") != hashlib.sha256(Path(image).read_bytes()).hexdigest():
            raise ValueError(f"第 {entry['slide']} 页审阅图像哈希不匹配")
        if entry.get("status") not in valid_status:
            raise ValueError(f"第 {entry['slide']} 页 status 无效")
        if not entry.get("basis"):
            raise ValueError(f"第 {entry['slide']} 页缺少判断依据")
        if entry["status"] == "无需修改" and entry.get("issues"):
            raise ValueError(f"第 {entry['slide']} 页不能同时标为无需修改和有问题")
        if entry["status"] == "有确定问题" and not entry.get("issues"):
            raise ValueError(f"第 {entry['slide']} 页缺少问题")
        object_paths = {o["path"] for o in flatten(inventory["slides"][entry["slide"] - 1]["objects"])}
        for issue in entry.get("issues", []):
            if missing := required_issue - issue.keys():
                raise ValueError(f"第 {entry['slide']} 页问题缺少字段：{sorted(missing)}")
            if issue["object"] not in object_paths and not issue["object"].startswith("region:"):
                raise ValueError(f"第 {entry['slide']} 页对象 {issue['object']} 不在对象清单中")
            if issue["severity"] not in ("高", "中", "低") or issue["confidence"] not in ("高", "中", "低"):
                raise ValueError(f"第 {entry['slide']} 页严重程度或置信度无效")
            if any(tip not in TIP_NAMES for tip in issue["tips"]):
                raise ValueError(f"第 {entry['slide']} 页建议编号无效")


def markdown(inventory: dict, diagnoses: list[dict], render_issues: list[str]) -> str:
    statuses = Counter(d["status"] for d in diagnoses)
    out = ["# 学术汇报 PPTX：先读对，再判断", "",
           f"输入：`{inventory['input']}`  ", f"SHA-256：`{inventory['sha256']}`  ",
           f"页面：{inventory['slide_count']}；已渲染：{sum(s['render']['status']=='ok' for s in inventory['slides'])}  ",
           f"结论：{dict(statuses)}", "", "## 整套主要发现", ""]
    ranked = [(d["slide"], i) for d in diagnoses for i in d.get("issues", [])]
    ranked.sort(key=lambda x: {"高": 0, "中": 1, "低": 2}[x[1]["severity"]])
    if ranked:
        for slide_num, issue in ranked[:8]:
            out.append(f"- 第 {slide_num} 页 · {issue['category']}：{issue['impact']}（{issue['severity']}）")
    else:
        out.append("- 尚无经过视觉审阅确认的问题；请逐页完成渲染图核对。")
    out += ["", "## 逐页结论", ""]
    for slide, d in zip(inventory["slides"], diagnoses):
        render_link = f"renders/slide-{slide['slide']:02d}.png" if slide["render"]["status"] == "ok" else "无"
        out += [f"### 第 {slide['slide']} 页 · {slide['title'] or '无标题'}", "",
                f"作用：{d['role']}；结论：**{d['status']}**；[渲染图]({render_link})", "",
                f"判断依据：{d['basis']}", ""]
        for observation in d.get("visual_observations", []):
            out.append(f"- 渲染核对：{observation['observation']}（来源：{observation['source']}；置信度：{observation['confidence']}）")
        for issue in d.get("issues", []):
            tips = "、".join(f"{n}（{TIP_NAMES[n]}）" for n in issue["tips"])
            out += [f"- **{issue['category']} / {issue['severity']} / 置信度{issue['confidence']}** · `{issue['object']}`",
                    f"  - 证据：{issue['evidence']}", f"  - 对听众的影响：{issue['impact']}",
                    f"  - 修改方向：{issue['direction']}", f"  - 对应建议：{tips}"]
        for doubt in d.get("uncertainties", []):
            if isinstance(doubt, dict):
                out.append(f"- 待核对：{doubt['evidence']}；{doubt['action']}")
            else:
                out.append(f"- 待核对：{doubt}")
        out.append("")
    out += ["## 解析盲区", ""]
    for issue in render_issues:
        out.append(f"- {issue}")
    for slide in inventory["slides"]:
        for blind in slide["parse_limitations"]:
            out.append(f"- 第 {slide['slide']} 页 · {blind}")
    out += ["- SmartArt、OLE 和图片内的数据不能由对象清单可靠还原；图表轴刻度可能由渲染器生成。",
            "- PPTX 备注和内部媒体元数据没有被当作放映时可见内容；字体替换可能使 LibreOffice 渲染不同于 PowerPoint。", ""]
    return "\n".join(out)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--review", type=Path, help="逐页视觉审阅 JSON；缺省只产出待审阅线索")
    args = parser.parse_args()
    if not args.input.is_file() or args.input.suffix.lower() != ".pptx":
        parser.error("输入必须是存在的 .pptx 文件")
    args.out.mkdir(parents=True, exist_ok=True)
    inventory = extract(args.input)
    render_issues = render(args.input, args.out, inventory["slides"])
    jdump(args.out / "inventory.json", inventory)
    if args.review:
        review = json.loads(args.review.read_text(encoding="utf-8"))
        validate_review(review, inventory)
        diagnoses = sorted(review["slides"], key=lambda x: x["slide"])
    else:
        diagnoses = [preliminary(slide) for slide in inventory["slides"]]
    result = {"input_sha256": inventory["sha256"], "generated_at": datetime.now(timezone.utc).isoformat(),
              "review_mode": "visual_review" if args.review else "preliminary_only",
              "slides": diagnoses, "render_issues": render_issues}
    jdump(args.out / "diagnosis.json", result)
    (args.out / "report.md").write_text(markdown(inventory, diagnoses, render_issues), encoding="utf-8")
    print(f"{inventory['slide_count']} slides; {sum(s['render']['status']=='ok' for s in inventory['slides'])} renders; {args.out.resolve()}")
    return 0 if not render_issues else 2


if __name__ == "__main__":
    raise SystemExit(main())
